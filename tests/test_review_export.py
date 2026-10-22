import io
import gc
import os
import sqlite3
import tempfile
import unittest
import uuid
import zipfile
from pathlib import Path

from openpyxl import Workbook, load_workbook

import server


class ReviewExportTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.root = Path(self.temp.name)
        self.originals = {
            "DATA_DIR": server.DATA_DIR,
            "STAGING_DIR": server.STAGING_DIR,
            "SUBMISSION_DIR": server.SUBMISSION_DIR,
            "DB_PATH": server.DB_PATH,
            "LOG_DIR": server.LOG_DIR,
        }
        server.DATA_DIR = self.root / "data"
        server.STAGING_DIR = server.DATA_DIR / "staging"
        server.SUBMISSION_DIR = server.DATA_DIR / "submissions"
        server.DB_PATH = server.DATA_DIR / "certification.db"
        server.LOG_DIR = self.root / "logs"
        server.DATA_DIR.mkdir()
        server.LOG_DIR.mkdir()

        # Simulate an existing production schema with business data.
        with sqlite3.connect(server.DB_PATH) as conn:
            conn.execute(
                """CREATE TABLE submissions (
                     id TEXT PRIMARY KEY, name TEXT NOT NULL, level TEXT, category TEXT,
                     applied_at TEXT, source_file TEXT NOT NULL, status TEXT NOT NULL,
                     imported_at TEXT NOT NULL, raw_path TEXT NOT NULL)"""
            )
            conn.execute(
                "INSERT INTO submissions VALUES (?,?,?,?,?,?,?,?,?)",
                ("candidate-1", "测试学员", "T4", "平台软件子类", "2026-08-07", "source.xlsx", "材料已导入", server.now_text(), "submissions/candidate-1/source.xlsx"),
            )
        with sqlite3.connect(server.DB_PATH) as conn:
            old_count = conn.execute("SELECT COUNT(*) FROM submissions").fetchone()[0]
        self.assertEqual(old_count, 1)

        server.init_storage()
        with server.db() as conn:
            columns = {row["name"] for row in conn.execute("PRAGMA table_info(submissions)")}
            self.assertNotIn("current_position", columns)
            self.assertEqual(conn.execute("SELECT COUNT(*) FROM submissions").fetchone()[0], 1)

            modules = [
                ("m1", "standard", "行为标准", "需求分析", 0.6, 0),
                ("m2", "standard", "行为标准", "方案设计", 0.4, 1),
                ("m3", "bonus", "知识成果贡献", "成果输出", None, 2),
                ("m4", "bonus", "团队成长贡献", "标准制定", None, 3),
            ]
            for module_id, section, item, name, weight, order in modules:
                conn.execute(
                    "INSERT INTO modules VALUES (?,?,?,?,?,?,?,?,?,?,?)",
                    (module_id, "candidate-1", section, item, name, weight, "标准", "举证", "", "举证", order),
                )
            for index, (reviewer_name, reviewer_weight, raw_scores, opinion) in enumerate([
                ("评委甲", 0.4, (4, 5, 2, 1), "通过"),
                ("评委乙", 0.6, (5, 4, 0.5, 2), "不通过"),
            ], start=1):
                reviewer_id, assignment_id, review_id = f"reviewer-{index}", f"assignment-{index}", f"review-{index}"
                conn.execute("INSERT INTO reviewers(id,name,created_at) VALUES(?,?,?)", (reviewer_id, reviewer_name, server.now_text()))
                conn.execute(
                    "INSERT INTO review_assignments VALUES(?,?,?,?,?,?,?,?,?)",
                    (assignment_id, "candidate-1", reviewer_id, reviewer_weight, uuid.uuid4().hex, "已完成", None, server.now_text(), server.now_text()),
                )
                conn.execute(
                    """INSERT INTO reviews(id,submission_id,reviewer_name,strengths,improvements,opinion,opinion_note,status,updated_at,assignment_id)
                       VALUES(?,?,?,?,?,?,?,?,?,?)""",
                    (review_id, "candidate-1", reviewer_name, f"{reviewer_name}填写的优势", f"{reviewer_name}填写的不足", opinion, "补充意见", "已提交", server.now_text(), assignment_id),
                )
                for module_id, score in zip(("m1", "m2", "m3", "m4"), raw_scores):
                    conn.execute("INSERT INTO module_scores VALUES(?,?,?,?)", (review_id, module_id, score, "单项评价"))

    def tearDown(self):
        for name, value in self.originals.items():
            setattr(server, name, value)
        gc.collect()
        self.temp.cleanup()

    def test_diagnostic_bundle_contains_logs_but_no_business_data(self):
        (server.LOG_DIR / "server.log").write_text("safe runtime event\n", encoding="utf-8")
        (server.LOG_DIR / "error.log").write_text("safe error event\n", encoding="utf-8")
        content, filename = server.create_diagnostic_bundle()
        self.assertRegex(filename, r"^CertificationOps-Diagnostics-\d{8}-\d{6}\.zip$")
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            names = set(archive.namelist())
            self.assertIn("diagnostic-summary.json", names)
            self.assertIn("logs/server.log", names)
            self.assertNotIn("data/certification.db", names)
            self.assertFalse(any("config" in name.lower() or "submission" in name.lower() for name in names))
            summary = archive.read("diagnostic-summary.json").decode("utf-8")
            self.assertNotIn("clientSecret", summary)

    def test_log_formatter_has_milliseconds_source_line_and_redaction(self):
        record = server.logging.LogRecord("test", server.logging.INFO, "server.py", 321, "event", (), None, "handler")
        formatter = server.APP_LOG.handlers[0].formatter
        rendered = formatter.format(record)
        self.assertRegex(rendered, r"\d{4}-\d{2}-\d{2} \d{2}:\d{2}:\d{2}\.\d{3}")
        self.assertIn("server.py:321", rendered)
        self.assertIn("handler()", rendered)
        redacted = server.safe_log_json({"code": "2608", "sessionToken": "secret", "name": "真实姓名", "scores": [5]})
        self.assertNotIn("2608", redacted)
        self.assertNotIn("secret", redacted)
        self.assertNotIn("真实姓名", redacted)
        self.assertNotIn("5", redacted)

    def test_logged_database_records_operation_without_values(self):
        captured = []

        class Capture(server.logging.Handler):
            def emit(self, record):
                captured.append(record.getMessage())

        handler = Capture()
        server.APP_LOG.addHandler(handler)
        try:
            with server.db() as conn:
                conn.execute("SELECT name FROM submissions WHERE id=?", ("candidate-1",)).fetchone()
        finally:
            server.APP_LOG.removeHandler(handler)
        joined = "\n".join(captured)
        self.assertIn("operation=SELECT", joined)
        self.assertIn("target=submissions", joined)
        self.assertNotIn("candidate-1", joined)

    def test_runtime_backup_is_portable_and_restores_database_and_files(self):
        server.SUBMISSION_DIR.mkdir(parents=True, exist_ok=True)
        marker = server.SUBMISSION_DIR / "candidate-1" / "附件.txt"
        marker.parent.mkdir(parents=True, exist_ok=True)
        marker.write_text("portable attachment", encoding="utf-8")
        backup = server.create_runtime_backup(retention=30)
        with zipfile.ZipFile(backup) as archive:
            self.assertIn("data/certification.db", archive.namelist())
            self.assertIn("data/submissions/candidate-1/附件.txt", archive.namelist())
        with server.db() as conn:
            conn.execute("DELETE FROM submissions WHERE id=?", ("candidate-1",))
        marker.unlink()
        server.restore_runtime_backup(backup)
        with server.db() as conn:
            self.assertEqual(conn.execute("SELECT COUNT(*) FROM submissions WHERE id=?", ("candidate-1",)).fetchone()[0], 1)
        self.assertEqual(marker.read_text(encoding="utf-8"), "portable attachment")

    def test_existing_data_survives_migration_and_exports(self):
        content, filename = server.export_submission_reviews("candidate-1")
        self.assertEqual(filename, "测试学员-评委评分导出.xlsx")
        if fixture_path := os.environ.get("CERT_EXPORT_FIXTURE_PATH"):
            Path(fixture_path).parent.mkdir(parents=True, exist_ok=True)
            Path(fixture_path).write_bytes(content)
        workbook = load_workbook(io.BytesIO(content), data_only=False)
        self.assertEqual(workbook.sheetnames, ["评委评分汇总", "逐项评分明细", "字段与计算说明"])

        summary = workbook["评委评分汇总"]
        self.assertEqual(summary["B5"].value, "测试学员")
        self.assertEqual(summary["D5"].value, "T4")
        self.assertEqual(summary["F5"].value, "平台软件子类")
        headers = [summary.cell(9, column).value for column in range(1, summary.max_column + 1)]
        self.assertIn("需求分析", headers)
        self.assertIn("方案设计", headers)
        self.assertIn("成果输出", headers)
        self.assertIn("标准制定", headers)
        self.assertIn("加权后总分", headers)
        self.assertIn("评估结论", headers)

        row_by_reviewer = {summary.cell(row, 1).value: row for row in range(11, summary.max_row + 1)}
        self.assertEqual(set(row_by_reviewer), {"评委甲", "评委乙"})
        weighted_column = headers.index("加权后总分") + 1
        opinion_column = headers.index("评估结论") + 1
        self.assertAlmostEqual(summary.cell(row_by_reviewer["评委甲"], weighted_column).value, 36.4)
        self.assertAlmostEqual(summary.cell(row_by_reviewer["评委乙"], weighted_column).value, 56.7)
        self.assertEqual(summary.cell(row_by_reviewer["评委甲"], opinion_column).value, "通过")
        self.assertEqual(summary.cell(row_by_reviewer["评委乙"], opinion_column).value, "不通过")

        detail = workbook["逐项评分明细"]
        self.assertEqual(detail.max_row - 4, 8)
        self.assertEqual(detail["B5"].value, "T4")
        workbook.close()

        batch_content, batch_filename = server.export_all_submission_reviews()
        self.assertRegex(batch_filename, r"^全部被评价人-评委评分导出-\d{8}-\d{6}\.zip$")
        with zipfile.ZipFile(io.BytesIO(batch_content)) as archive:
            self.assertEqual(archive.namelist(), [filename])
            batch_workbook = load_workbook(io.BytesIO(archive.read(filename)), data_only=False)
            self.assertEqual(batch_workbook["评委评分汇总"]["B5"].value, "测试学员")
            batch_workbook.close()

    def test_completed_scores_survive_weight_changes_and_dashboard_recalculates(self):
        with server.db() as conn:
            scores_before = [tuple(row) for row in conn.execute(
                "SELECT review_id,module_id,score,comment FROM module_scores ORDER BY review_id,module_id"
            )]
            reviews_before = [tuple(row) for row in conn.execute(
                "SELECT id,strengths,improvements,opinion,opinion_note,status FROM reviews ORDER BY id"
            )]

        server.update_assignment("assignment-1", {"reviewerName": "评委甲", "weight": 0.5})
        server.update_assignment("assignment-2", {"reviewerName": "评委乙", "weight": 0.5})

        with server.db() as conn:
            scores_after = [tuple(row) for row in conn.execute(
                "SELECT review_id,module_id,score,comment FROM module_scores ORDER BY review_id,module_id"
            )]
            reviews_after = [tuple(row) for row in conn.execute(
                "SELECT id,strengths,improvements,opinion,opinion_note,status FROM reviews ORDER BY id"
            )]
        self.assertEqual(scores_after, scores_before)
        self.assertEqual(reviews_after, reviews_before)

        dashboard_item = server.dashboard_data()["items"][0]
        self.assertEqual(dashboard_item["completedCount"], 2)
        self.assertAlmostEqual(dashboard_item["weightTotal"], 1.0)
        self.assertAlmostEqual(dashboard_item["weightedTotal"], 92.75)
        self.assertEqual({row["weight"] for row in dashboard_item["reviewers"]}, {0.5})

    def test_assignment_template_import_updates_weight_without_overwriting_scores(self):
        template, filename = server.assignment_import_template()
        self.assertEqual(filename, "评价关系批量导入模板.xlsx")
        template_book = load_workbook(io.BytesIO(template))
        self.assertEqual(template_book["评价关系"]["A1"].value, "被评价人")
        self.assertEqual(template_book["评价关系"]["C1"].value, "权重（%）")
        template_book.close()

        with server.db() as conn:
            scores_before = [tuple(row) for row in conn.execute(
                "SELECT review_id,module_id,score,comment FROM module_scores ORDER BY review_id,module_id"
            )]
        upload = Workbook()
        sheet = upload.active
        sheet.title = "评价关系"
        sheet.append(["被评价人", "评价人", "权重（%）"])
        sheet.append(["测试学员", "评委甲", 70])
        buffer = io.BytesIO()
        upload.save(buffer)
        upload.close()
        result = server.import_assignment_workbook(buffer.getvalue())
        self.assertEqual(result, {"created": 0, "updated": 1, "errors": []})
        with server.db() as conn:
            self.assertAlmostEqual(conn.execute("SELECT weight FROM review_assignments WHERE id='assignment-1'").fetchone()[0], 0.7)
            scores_after = [tuple(row) for row in conn.execute(
                "SELECT review_id,module_id,score,comment FROM module_scores ORDER BY review_id,module_id"
            )]
        self.assertEqual(scores_after, scores_before)

    def test_reviewer_access_code_is_hashed_and_required_for_tasks(self):
        server.set_reviewer_access_code("reviewer-1", "2608")
        with server.db() as conn:
            row = conn.execute("SELECT access_code_salt,access_code_hash FROM reviewers WHERE id='reviewer-1'").fetchone()
            self.assertTrue(row["access_code_salt"])
            self.assertNotEqual(row["access_code_hash"], "2608")
        with self.assertRaisesRegex(ValueError, "访问码错误"):
            server.verify_reviewer_access("reviewer-1", "1111")
        login = server.verify_reviewer_access("reviewer-1", "2608")
        tasks = server.authorized_reviewer_tasks("reviewer-1", login["token"])
        self.assertEqual(len(tasks), 1)
        with self.assertRaisesRegex(ValueError, "身份验证已失效"):
            server.authorized_reviewer_tasks("reviewer-1", "invalid")

    def test_online_application_reuses_review_data_model_and_locks_after_submit(self):
        admin = {"userId": "mock-admin", "name": "模拟管理员", "role": "管理员"}
        student = {"userId": "mock-student", "name": "模拟学员", "role": "学员"}
        family = server.save_job_taxonomy({"name": "研发职族", "kind": "family"})
        category = server.save_job_taxonomy({"name": "软件职类", "kind": "category", "parentId": family["id"]})
        sub = server.save_job_taxonomy({"name": "平台软件子类", "kind": "subcategory", "parentId": category["id"]})
        standard = server.save_qualification_standard({
            "name": "T4 在线认证标准",
            "level": "T4",
            "category": "平台软件子类",
            "subcategoryId": sub["id"],
            "status": "已存档",
            "modules": [
                {"sectionType": "standard", "evaluationItem": "行为标准", "moduleName": "需求分析", "weight": .6, "standard": "能独立完成复杂需求分析"},
                {"sectionType": "standard", "evaluationItem": "行为标准", "moduleName": "方案设计", "weight": .4, "standard": "能输出可落地方案"},
                {"sectionType": "bonus", "evaluationItem": "团队成长贡献", "moduleName": "标准制定", "standard": "有标准制定成果"},
            ],
        }, admin)
        application = server.create_online_application(standard["standard"]["id"], student)
        application_id = application["submission"]["id"]
        self.assertEqual(application["submission"]["status"], "填报中")
        self.assertEqual(application["submission"]["entry_mode"], "online")
        self.assertTrue(all(not item["evidence"] for item in application["modules"]))

        with self.assertRaisesRegex(ValueError, "请先完成"):
            server.save_online_application(application_id, {"submit": True, "modules": []}, student)

        module_payload = [
            {"moduleId": item["id"], "evidence": f"{item['module_name']}举证", "selfScore": "4"}
            for item in application["modules"]
        ]
        attachment = server.store_application_attachment(
            application_id, "举证.txt", io.BytesIO(b"evidence"), 8, student,
            application["modules"][0]["id"], "需求分析举证",
        )
        self.assertTrue(attachment["id"])
        linked = server.store_application_link(application_id, {
            "moduleId": application["modules"][1]["id"],
            "materialName": "方案文档",
            "linkUrl": "https://docs.example.test/design",
        }, student)
        self.assertTrue(linked["id"])
        submitted = server.save_online_application(application_id, {
            "name": "模拟学员", "appliedAt": "2026-09-01", "submit": True, "modules": module_payload,
        }, student)
        self.assertEqual(submitted["submission"]["status"], "待评审")
        self.assertEqual(len(submitted["attachments"]), 2)
        with self.assertRaisesRegex(ValueError, "已提交"):
            server.save_online_application(application_id, {"submit": False, "modules": module_payload}, student)
        self.assertIn(application_id, {item["id"] for item in server.list_assignments()["submissions"]})


if __name__ == "__main__":
    unittest.main()
