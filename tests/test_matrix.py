import copy
import json
import unittest

import server
import qualification_matrix as matrix
import test_standard_framework as fixtures


class MatrixTests(unittest.TestCase):
    setUp = fixtures.StandardFrameworkTests.setUp
    tearDown = fixtures.StandardFrameworkTests.tearDown

    def matrix(self):
        value = matrix.default_matrix()
        for dim in value['dimensions']:
            limit = 7 if dim['basis'] == 'grade' else 4
            dim['rows'] = [{'id': dim['id']+'row', 'name': dim['name']+'子项',
                            'cells': {str(i): f'{dim["name"]}新增{i}' for i in range(1, limit+1)}}]
        return value

    def save(self, value=None, archive=True, sid=None, revision=0):
        return server.save_qualification_standard({'name':'矩阵标准','subcategoryId':self.sub['id'],
            'status':'已存档' if archive else '草稿', 'matrix':value or self.matrix(), 'revision':revision},self.admin,sid)

    def test_grade_tier_mapping_and_inheritance(self):
        value = matrix.normalize_matrix(self.matrix(), archive=True)
        for grade, tier in [(1,1),(2,1),(3,2),(4,2),(5,3),(6,3),(7,4)]:
            resolved = matrix.resolve_requirements(value,grade)
            self.assertEqual(len(resolved[0]['rows'][0]['requirements']),grade)
            self.assertEqual(len(resolved[1]['rows'][0]['requirements']),tier)
        value['dimensions'][0]['rows'][0]['cells']['2']=matrix.NO_ADDITION
        self.assertNotIn(matrix.NO_ADDITION,matrix.resolve_requirements(value,3)[0]['rows'][0]['text'])
        for invalid in (0,8,12,True,1.5,None,'T4'):
            with self.assertRaises(ValueError):
                matrix.resolve_requirements(value,invalid)

    def test_draft_and_archive_allow_pending_levels(self):
        value = self.matrix()
        value['dimensions'][0]['rows'][0]['cells']['7']=''
        draft = self.save(value,False)['standard']
        self.assertEqual(draft['revision'],1)
        archived = self.save(value,True,draft['id'],1)['standard']
        self.assertEqual(archived['status'],'已存档')
        resolved = matrix.resolve_requirements(archived['matrix'],7)[0]['rows'][0]
        self.assertTrue(resolved['pending'])
        self.assertNotIn('7', [part['level'] for part in resolved['requirements']])
        with self.assertRaisesRegex(ValueError,'不可修改'):
            self.save(value,False,draft['id'],2)

    def test_archive_keeps_selected_dimensions_only_and_marks_pending_levels(self):
        value = self.matrix()
        value['dimensions'] = [value['dimensions'][0], value['dimensions'][1]]
        value['dimensions'][1]['rows'].append({'id':'knowledge-b','name':'知识子项B','cells':{}})
        standard = self.save(value,True)['standard']
        self.assertEqual([dim['name'] for dim in standard['matrix']['dimensions']],['基本条件','知识标准'])
        application = server.create_online_application(standard['id'],self.student,7)
        self.assertEqual(len(application['modules']),3)
        pending = next(m for m in application['modules'] if m['module_name']=='知识子项B')
        self.assertIn('待补充',pending['standard'])
        resolved = json.loads(application['submission']['standard_snapshot_json'])['resolved'][1]['rows'][1]
        self.assertTrue(resolved['pending'])
        self.assertEqual(resolved['text'],'')

    def test_dimension_accepts_subset_and_column_style_is_validated(self):
        value = self.matrix()
        value['dimensions'] = value['dimensions'][:1]
        value['dimensions'][0]['columnStyle'] = 'number'
        normalized = matrix.normalize_matrix(value,archive=True)
        self.assertEqual(len(normalized['dimensions']),1)
        self.assertEqual(normalized['dimensions'][0]['columnStyle'],'number')
        self.assertEqual(matrix.normalize_matrix(self.matrix())['dimensions'][1]['columnStyle'],'name')
        broken = self.matrix()
        broken['dimensions'][1]['columnStyle'] = 'number'
        with self.assertRaisesRegex(ValueError,'列口径'):
            matrix.normalize_matrix(broken)
        broken = self.matrix()
        broken['dimensions'][0]['columnStyle'] = 'roman'
        with self.assertRaisesRegex(ValueError,'列口径'):
            matrix.normalize_matrix(broken)

    def test_empty_dimension_still_blocks_archive(self):
        value = self.matrix()
        value['dimensions'][0]['rows'] = []
        with self.assertRaisesRegex(ValueError,'添加子维度'):
            self.save(value,True)

    def test_basic_condition_is_informational_and_not_required(self):
        standard = self.save()['standard']
        application = server.create_online_application(standard['id'],self.student,7)
        sid = application['submission']['id']
        modules = [{'moduleId':m['id'],'evidence':'具体举证'} for m in application['modules'] if m['evaluation_item']!='基本条件']
        submitted = server.save_online_application(sid,{'submit':True,'modules':modules},self.student)
        self.assertEqual(submitted['submission']['status'],'待评审')
        reread = server.submission_detail(sid)
        basic = next(m for m in reread['modules'] if m['evaluation_item']=='基本条件')
        self.assertEqual(basic['evidence'],'')

    def test_attachment_link_accepts_internal_addresses_but_blocks_scripts(self):
        standard = self.save()['standard']
        application = server.create_online_application(standard['id'],self.student,7)
        sid = application['submission']['id']
        internal = server.store_application_link(sid,{'moduleId':application['modules'][0]['id'],'materialName':'内部文档','linkUrl':'dms://doc/2026/001','storagePath':'内部知识库：2026-001'},self.student)
        self.assertTrue(internal['id'])
        with self.assertRaisesRegex(ValueError,'脚本'):
            server.store_application_link(sid,{'moduleId':'','materialName':'危险链接','linkUrl':'javascript:alert(1)'},self.student)

    def test_export_subclass_standard_matrix(self):
        import io
        from openpyxl import load_workbook
        standard = self.save()['standard']
        content, filename = server.export_subclass_standard_matrix(self.sub['id'])
        self.assertTrue(content.startswith(b'PK'))
        self.assertIn('平台软件', filename)
        workbook = load_workbook(io.BytesIO(content), data_only=True)
        sheet = workbook['任职资格标准']
        self.assertIn('平台软件', sheet.cell(1,1).value)
        first_column = [sheet.cell(row,1).value for row in range(1, sheet.max_row + 1)]
        self.assertIn('一、基本条件', first_column)
        self.assertIn('二、知识标准', first_column)
        workbook.close()
        empty_sub = server.save_job_taxonomy({'name':'空子类','kind':'subcategory','parentId':self.category['id']})
        with self.assertRaisesRegex(ValueError, '暂无已存档'):
            server.export_subclass_standard_matrix(empty_sub['id'])

    def test_only_unarchived_standards_can_be_deleted(self):
        draft = self.save(archive=False)['standard']
        server.delete_qualification_standard(draft['id'])
        self.assertIsNone(server.qualification_standard_detail(draft['id']))
        with self.assertRaisesRegex(ValueError,'不存在'):
            server.delete_qualification_standard(draft['id'])
        with self.assertRaisesRegex(ValueError,'不可删除'):
            server.delete_qualification_standard(self.save()['standard']['id'])

    def test_dimensions_can_be_added_removed_renamed_and_snapshot_isolated(self):
        value=self.matrix()
        value['dimensions']=value['dimensions'][1:]
        custom=copy.deepcopy(value['dimensions'][0]);custom.update(id='custom',name='自定义能力',basis='tier')
        value['dimensions'].append(custom)
        standard=self.save(value)['standard']
        source=copy.deepcopy(standard['matrix'])
        template=server.get_qualification_template()
        template['matrix']['dimensions'][0]['name']='新的基本条件'
        server.save_qualification_template(template)
        self.assertEqual(server.qualification_standard_detail(standard['id'])['standard']['matrix'],source)
        second=server.iterate_qualification_standard(standard['id'],self.admin,'新增维度')['standard']
        self.assertEqual(second['matrix'],source)
        second['matrix']['dimensions'][0]['name']='修订知识标准'
        self.save(second['matrix'],False,second['id'],second['revision'])
        self.assertEqual(server.qualification_standard_detail(standard['id'])['standard']['matrix'],source)
        self.assertEqual(server.qualification_standard_history(second['id'])[0]['snapshot']['standard']['matrix']['dimensions'][0]['name'],'修订知识标准')

    def test_template_validation_revision_and_blank_cells(self):
        template=server.get_qualification_template()
        template['matrix']=self.matrix()
        saved=server.save_qualification_template(template)
        self.assertEqual(saved['matrix']['dimensions'][0]['rows'][0]['cells']['1'],'')
        with self.assertRaises(ValueError):
            server.save_qualification_template(template)
        for mutate in (lambda v:v['dimensions'].append(copy.deepcopy(v['dimensions'][0])),
                       lambda v:v['dimensions'][0]['rows'].append(copy.deepcopy(v['dimensions'][0]['rows'][0])),
                       lambda v:v['dimensions'][1]['rows'][0]['cells'].update({'5':'非法层级'})):
            value=self.matrix();mutate(value)
            with self.assertRaises(ValueError):matrix.normalize_matrix(value)

    def test_positions_per_subclass_and_immutable_archive_application(self):
        positions={'1':'助理工程师','7':'专家工程师'}
        server.save_job_positions(self.sub['id'],{'positions':positions,'original':{}})
        standard=self.save()['standard']
        original=server.get_job_positions(self.sub['id'])
        server.save_job_positions(self.sub['id'],{'positions':{'7':'新岗位名称'},'original':original})
        self.assertEqual(server.qualification_standard_detail(standard['id'])['standard']['positions']['7'],'专家工程师')
        application=server.create_online_application(standard['id'],self.student,7)
        snapshot=json.loads(application['submission']['standard_snapshot_json'])
        self.assertEqual(snapshot['positionName'],'专家工程师')
        self.assertEqual(snapshot['grade'],7)
        self.assertEqual(len(snapshot['resolved'][0]['rows'][0]['requirements']),7)
        self.assertEqual(len(snapshot['resolved'][1]['rows'][0]['requirements']),4)
        self.assertEqual(application['submission']['level'],'七级')
        with self.assertRaisesRegex(ValueError,'快照绑定'):
            server.update_submission_basic(application['submission']['id'],{'name':'测试人','level':'一级','category':application['submission']['category']})
        self.assertEqual(len(application['modules']),6)
        other=server.save_job_taxonomy({'name':'其他子类','kind':'subcategory','parentId':self.category['id']})
        self.assertEqual(server.get_job_positions(other['id']),{})
        with self.assertRaisesRegex(ValueError,'已被更新'):
            server.save_job_positions(self.sub['id'],{'positions':positions,'original':{}})

    def test_application_reference_gate_and_grade_specific_duplicates(self):
        draft=self.save(archive=False)['standard']
        with self.assertRaises(ValueError):server.create_online_application(draft['id'],self.student,4)
        standard=self.save(sid=draft['id'],revision=1)['standard']
        with self.assertRaises(ValueError):server.create_online_application(standard['id'],self.student)
        a=server.create_online_application(standard['id'],self.student,3)
        b=server.create_online_application(standard['id'],self.student,4)
        self.assertNotEqual(a['submission']['id'],b['submission']['id'])
        self.assertEqual(a['submission']['id'],server.create_online_application(standard['id'],self.student,3)['submission']['id'])

    def test_all_evidence_dimensions_require_weights_and_are_scored(self):
        standard=self.save()['standard']
        application=server.create_online_application(standard['id'],self.student,7)
        sid=application['submission']['id']
        with self.assertRaises(ValueError):server.save_online_application(sid,{'submit':True},self.student)
        submitted=server.save_online_application(sid,{'submit':True,'modules':[{'moduleId':m['id'],'evidence':'举证','selfScore':'符合'} for m in application['modules']]},self.student)
        self.assertEqual(submitted['submission']['status'],'待评审')
        scores={'submit':True,'moduleScores':[{'moduleId':m['id'],'score':2} for m in application['modules']]}
        with self.assertRaisesRegex(ValueError,'权重'):server.save_review(sid,scores)
        with self.assertRaisesRegex(ValueError,'权重'):server.save_review(sid,{**scores,'submit':False})
        behavior=next(m for m in application['modules'] if m['section_type']=='standard')
        with self.assertRaisesRegex(ValueError,'权重'):
            server.save_matrix_review_weights(sid,{'weights':{behavior['id']:1}})
        ordinary=[m for m in application['modules'] if m['section_type']!='bonus']
        weights={m['id']:1/len(ordinary) for m in ordinary}
        server.save_matrix_review_weights(sid,{'weights':weights})
        result=server.save_review(sid,scores)
        self.assertEqual(result['scoreSummary']['total'],42)
        self.assertEqual(len(result['scoreSummary']['items']),6)
        with server.db() as conn:
            self.assertEqual(conn.execute('SELECT COUNT(*) FROM module_scores WHERE review_id=?',(result['reviewId'],)).fetchone()[0],6)
        with self.assertRaises(ValueError):server.save_matrix_review_weights(sid,{'weights':weights})

    def test_reference_scores_persist_missing_items_block_submit_and_export_is_complete(self):
        import io
        from openpyxl import load_workbook
        standard=self.save()['standard']
        detail=server.create_online_application(standard['id'],self.student,7)
        sid=detail['submission']['id']
        server.save_online_application(sid,{'submit':True,'modules':[{'moduleId':m['id'],'evidence':'项目成果'} for m in detail['modules']]},self.student)
        ordinary=[m for m in detail['modules'] if m['section_type']!='bonus']
        server.save_matrix_review_weights(sid,{'weights':{m['id']:1/len(ordinary) for m in ordinary}})
        assigned=server.upsert_assignment({'submissionId':sid,'reviewerName':'逐项测试评委','weight':1})
        scores=[{'moduleId':m['id'],'score':0 if m['section_type']=='bonus' else 4,'comment':'逐项判断依据'} for m in detail['modules']]
        with self.assertRaisesRegex(ValueError,'逐项评分'):
            server.save_assigned_review(assigned['token'],{'submit':True,'moduleScores':scores[1:]})
        with server.db() as conn:
            self.assertEqual(conn.execute('SELECT COUNT(*) FROM reviews WHERE submission_id=?',(sid,)).fetchone()[0],0)
        draft=server.save_assigned_review(assigned['token'],{'moduleScores':scores})
        self.assertEqual(draft['scoreSummary']['total'],80)
        reread=server.assigned_review_detail(assigned['token'])
        self.assertEqual(len(reread['scores']),6)
        self.assertTrue(all(s['comment']=='逐项判断依据' for s in reread['scores']))
        with self.assertRaisesRegex(ValueError,'0 至 5'):
            server.save_assigned_review(assigned['token'],{'moduleScores':[{'moduleId':ordinary[0]['id'],'score':6}]})
        final=server.save_assigned_review(assigned['token'],{'submit':True,'moduleScores':scores})
        self.assertEqual(final['status'],'已提交')
        blob,_=server.export_submission_reviews(sid)
        workbook=load_workbook(io.BytesIO(blob),data_only=True)
        headers=[cell.value for cell in workbook['评委评分汇总'][9]]
        for m in ordinary:self.assertIn(m['module_name'],headers)
        self.assertIn('各维度评分折算总分',headers)
        self.assertEqual(workbook['逐项评分明细'].max_row,10)
        workbook.close()

    def test_attachment_only_evidence_also_requires_an_explicit_score(self):
        import io
        standard=self.save()['standard']
        detail=server.create_online_application(standard['id'],self.student,7)
        sid=detail['submission']['id'];bonus=next(m for m in detail['modules'] if m['section_type']=='bonus')
        server.store_application_attachment(sid,'证明.pdf',io.BytesIO(b'test'),4,self.student,bonus['id'])
        ordinary=[m for m in detail['modules'] if m['section_type']!='bonus']
        server.save_online_application(sid,{'submit':True,'modules':[{'moduleId':m['id'],'evidence':'举证'} for m in ordinary]},self.student)
        server.save_matrix_review_weights(sid,{'weights':{m['id']:1/len(ordinary) for m in ordinary}})
        scores=[{'moduleId':m['id'],'score':5} for m in ordinary]
        with self.assertRaisesRegex(ValueError,'逐项评分'):
            server.save_review(sid,{'submit':True,'moduleScores':scores})
        result=server.save_review(sid,{'submit':True,'moduleScores':scores+[{'moduleId':bonus['id'],'score':0}]})
        self.assertEqual(result['scoreSummary']['total'],100)

    def test_matrix_cannot_be_saved_through_legacy_endpoint_payload(self):
        standard=self.save(archive=False)['standard']
        with self.assertRaisesRegex(ValueError,'矩阵编辑'):
            server.save_qualification_standard({'name':'错误覆盖','status':'草稿'},self.admin,standard['id'])
        self.assertIsNotNone(server.qualification_standard_detail(standard['id'])['standard']['matrix'])
