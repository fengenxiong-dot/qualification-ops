import io
import json
import threading
import unittest
from datetime import datetime,timedelta
from http.client import HTTPConnection
from openpyxl import load_workbook
import server
import test_standard_framework as fixtures


class PersonnelTests(unittest.TestCase):
    setUp=fixtures.StandardFrameworkTests.setUp
    tearDown=fixtures.StandardFrameworkTests.tearDown

    def person(self,name='职级测试人',number='001',level='一级'):
        return server.save_personnel({'name':name,'employeeNo':number,'department':'研发','initialLevel':level,'initialDate':'2020-01-01'},self.admin)['id']

    def manual(self,pid,day,level='二级',outcome='通过',promoted=True,effective=None):
        return server.save_personnel_certification(pid,{'outcome':outcome,'decidedOn':day,'effectiveOn':effective or day,
            'resultLevel':level,'promoted':promoted,'manualScore':88,'note':'历史认证通知','appliedLevel':level,'category':'平台软件'},self.admin)['id']

    def application(self):
        standard=server.save_qualification_standard(fixtures.StandardFrameworkTests.payload(self),self.admin)
        return server.create_online_application(standard['standard']['id'],self.student)

    def complete(self,app):
        sid=app['submission']['id']
        server.save_online_application(sid,{'submit':True,'modules':[{'moduleId':m['id'],'evidence':'项目成果'} for m in app['modules']]},self.student)
        assigned=server.upsert_assignment({'submissionId':sid,'reviewerName':'评委','weight':1})
        server.save_assigned_review(assigned['token'],{'submit':True,'moduleScores':[{'moduleId':m['id'],'score':4 if m['section_type']=='standard' else 0} for m in app['modules']]})
        return assigned

    def test_identity_is_explicit_same_names_are_not_merged_and_link_is_unique(self):
        first=self.application();second=self.application()
        self.assertEqual(server.personnel_data()['items'],[])
        self.assertEqual(len(server.personnel_data()['unlinked']),2)
        a=self.person('同名','001');b=self.person('同名','002')
        server.link_personnel_submission(a,first['submission']['id'],self.admin)
        server.link_personnel_submission(b,second['submission']['id'],self.admin)
        self.assertEqual(server.personnel_data()['unlinked'],[])
        self.assertEqual([p['certificationCount'] for p in server.personnel_data()['items']],[1,1])
        with self.assertRaisesRegex(ValueError,'已关联'):server.link_personnel_submission(b,first['submission']['id'],self.admin)
        with self.assertRaisesRegex(ValueError,'工号'):self.person('另一个人','001')

    def test_multiple_certifications_failures_and_future_effective_grade(self):
        pid=self.person()
        self.manual(pid,'2024-01-01','二级')
        self.manual(pid,'2025-01-01',outcome='未通过',promoted=False)
        self.manual(pid,'2026-01-01','三级')
        self.manual(pid,'2026-02-01','三级',promoted=False)
        future=(datetime.fromisoformat(server.now_text())+timedelta(days=40)).date().isoformat()
        self.manual(pid,server.now_text()[:10],'四级',effective=future)
        p=server.personnel_data(pid)['person']
        self.assertEqual((p['initial_level'],p['currentLevel']),('一级','三级'))
        self.assertEqual((p['certificationCount'],p['passedCount'],p['failedCount'],p['promotionCount']),(5,4,1,2))
        self.assertEqual(p['records'][1]['beforeLevel'],'二级')
        self.assertEqual(p['records'][1]['afterLevel'],'二级')
        self.assertEqual(p['records'][-1]['effectStatus'],'待生效')
        with self.assertRaisesRegex(ValueError,'高于'):self.manual(pid,'2026-03-01','二级')
        self.assertEqual(server.personnel_data(pid)['person']['certificationCount'],5)

    def test_out_of_order_history_replays_and_corrections_are_audited(self):
        pid=self.person();later=self.manual(pid,'2026-01-01','三级')
        self.manual(pid,'2024-01-01','二级')
        p=server.personnel_data(pid)['person'];self.assertEqual(p['records'][1]['beforeLevel'],'二级')
        r=next(r for r in p['records'] if r['id']==later)
        values={'outcome':'通过','decidedOn':'2026-01-01','effectiveOn':'2026-01-01','resultLevel':'四级','promoted':True,'revision':r['revision'],'note':'核对原始通知后更正'}
        server.save_personnel_certification(pid,values,self.admin,later)
        self.assertEqual(server.personnel_data(pid)['person']['currentLevel'],'四级')
        with self.assertRaisesRegex(ValueError,'更新'):server.save_personnel_certification(pid,values,self.admin,later)
        self.assertIn('修订认证结论',[h['action'] for h in server.personnel_data(pid)['history']])

    def test_confirmation_requires_completed_reviews_and_keeps_score_snapshot(self):
        pid=self.person();app=self.application();sid=app['submission']['id']
        rid=server.link_personnel_submission(pid,sid,self.admin)['id']
        values={'outcome':'通过','decidedOn':server.now_text()[:10],'effectiveOn':server.now_text()[:10],'resultLevel':'二级','promoted':True,'revision':0}
        with self.assertRaisesRegex(ValueError,'全部评委'):server.save_personnel_certification(pid,values,self.admin,rid)
        assignment=self.complete(app)
        self.assertTrue(server.personnel_data(pid)['person']['records'][0]['reviewReady'])
        server.save_personnel_certification(pid,values,self.admin,rid)
        p=server.personnel_data(pid)['person'];self.assertEqual(p['currentLevel'],'二级')
        self.assertEqual(p['records'][0]['scoreData']['weightedTotal'],80)
        server.save_assigned_review(assignment['token'],{'submit':True,'moduleScores':[{'moduleId':m['id'],'score':1} for m in app['modules']]})
        self.assertEqual(server.personnel_data(pid)['person']['records'][0]['scoreData']['weightedTotal'],80)
        with server.db() as conn:conn.execute('DELETE FROM submissions WHERE id=?',(sid,))
        r=server.personnel_data(pid)['person']['records'][0]
        self.assertIsNone(r['submission_id']);self.assertEqual(r['source']['id'],sid);self.assertEqual(r['scoreData']['weightedTotal'],80)

    def test_validation_and_initial_grade_lock(self):
        pid=self.person(level='')
        with self.assertRaisesRegex(ValueError,'初始职级'):self.manual(pid,'2024-01-01')
        p=server.personnel_data(pid)['person']
        server.save_personnel({'name':p['name'],'employeeNo':'001','initialLevel':'一级','initialDate':'2020-01-01','revision':0},self.admin,pid)
        self.manual(pid,'2024-01-01')
        with self.assertRaisesRegex(ValueError,'锁定'):
            server.save_personnel({'name':p['name'],'initialLevel':'三级','revision':1},self.admin,pid)
        with self.assertRaises(ValueError):self.manual(pid,'2024-02-30','三级')
        with self.assertRaisesRegex(ValueError,'此前职级'):self.manual(pid,'2025-01-01','三级',promoted=False)
        self.assertEqual(server.personnel_data(pid)['person']['currentLevel'],'二级')

    def test_export_matches_timeline_and_prevents_excel_formula_injection(self):
        pid=self.person('=1+1');self.manual(pid,'2024-01-01')
        app=self.application();server.link_personnel_submission(pid,app['submission']['id'],self.admin)
        other=self.person('其他人员','002')
        data,name=server.export_personnel(query='=1+1')
        book=load_workbook(io.BytesIO(data),data_only=False)
        self.assertTrue(name.endswith('.xlsx'));self.assertEqual(book['人员职级清单'].max_row,2)
        self.assertEqual(book['人员职级清单']['B2'].value,'=1+1');self.assertEqual(book['人员职级清单']['B2'].data_type,'s')
        self.assertEqual(book['人员职级清单']['F2'].value,'二级')
        self.assertEqual(book['认证与职级变迁'].max_row,3)
        self.assertEqual(book['人员职级清单'].freeze_panes,'A2');book.close()
        data,_=server.export_personnel(pid);book=load_workbook(io.BytesIO(data));self.assertEqual(book['人员职级清单'].max_row,2);self.assertNotIn('待关联认证材料',book.sheetnames);book.close()

    def test_admin_only_http_reads_writes_and_exports(self):
        pid=self.person();http=server.CertificationHTTPServer(('127.0.0.1',0),server.Handler)
        thread=threading.Thread(target=http.serve_forever,daemon=True);thread.start()
        def request(method,path,role,body=None):
            conn=HTTPConnection('127.0.0.1',http.server_port,timeout=10)
            conn.request(method,path,json.dumps(body or {}) if method=='POST' else None,{'Cookie':'cert_mock_user='+role,'Content-Type':'application/json'})
            response=conn.getresponse();data=response.read();status=response.status;conn.close();return status,data
        try:
            for path in ('/api/admin/personnel','/api/admin/personnel/'+pid,'/api/admin/personnel/export'):
                self.assertEqual(request('GET',path,'mock-student')[0],403)
                self.assertEqual(request('GET',path,'mock-admin')[0],200)
            self.assertEqual(request('POST','/api/admin/personnel','mock-reviewer',{'name':'禁止建档'})[0],403)
            status,data=request('POST','/api/admin/personnel/'+pid+'/records','mock-admin',{'outcome':'通过'})
            self.assertEqual(status,400)
        finally:http.shutdown();http.server_close();thread.join()


if __name__=='__main__':unittest.main()
