from __future__ import annotations

import base64
import hashlib
import hmac
import io
import json
import logging
from logging.handlers import TimedRotatingFileHandler
import traceback
import mimetypes
import os
import posixpath
import re
import shutil
import secrets
import sqlite3
import subprocess
import sys
import tempfile
import threading
import time
import uuid
import zipfile
import xml.etree.ElementTree as ET
from datetime import datetime
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path
from urllib.parse import quote, unquote, urlparse

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from qualification_matrix import (default_matrix, normalize_matrix, resolve_requirements,
                                  grade_number, GRADE_NAMES, TIERS)
import evidence_workflow as evidence
evidence.APP = sys.modules[__name__]

BASIC_CONDITION_NAME = "基本条件"


ROOT = Path(sys.executable).resolve().parent if getattr(sys, "frozen", False) else Path(__file__).resolve().parent
STATIC_DIR = ROOT / "static"
DATA_DIR = Path(os.environ.get("CERT_DATA_DIR", ROOT / "data"))
STAGING_DIR = DATA_DIR / "staging"
SUBMISSION_DIR = DATA_DIR / "submissions"
DB_PATH = DATA_DIR / "certification.db"
LOG_DIR = Path(os.environ.get("CERT_LOG_DIR", ROOT / "logs"))
MAX_FILE_MB = max(10, int(os.environ.get("CERT_MAX_FILE_MB", "300")))
PARSE_TIMEOUT_SECONDS = max(30, int(os.environ.get("CERT_PARSE_TIMEOUT_SECONDS", "60")))
MAX_BODY = MAX_FILE_MB * 1024 * 1024
UPLOAD_CHUNK = 1024 * 1024
JOB_LOCK = threading.Lock()
TRACE_LOCAL = threading.local()


SENSITIVE_KEYS = {
    "instructions", "note", "code_hash", "token_hash",
    "code", "password", "passwd", "secret", "clientsecret", "token", "sessiontoken",
    "cookie", "authorization", "access_code_hash", "access_code_salt", "content",
    "evidence", "standard", "strengths", "improvements", "opinion_note", "comment",
    "score", "scores", "reviewername", "reviewer_name", "name", "recipient_name", "query", "matrix", "cells", "positions",
    "employeeno", "currentlevel", "department", "initiallevel", "resultlevel", "initialdate", "decidedon", "effectiveon", "outcome", "promoted", "manualscore",
}


def safe_log_value(value, key: str = "", depth: int = 0):
    """Return bounded, privacy-safe metadata suitable for operational logs."""
    normalized = re.sub(r"[^a-z0-9_]", "", str(key).lower())
    if normalized in SENSITIVE_KEYS or any(part in normalized for part in ("password", "secret", "token", "cookie", "accesscode")):
        return "<redacted>"
    if depth > 3:
        return "<max-depth>"
    if value is None or isinstance(value, (bool, int, float)):
        return value
    if isinstance(value, bytes):
        return {"type": "bytes", "size": len(value)}
    if isinstance(value, str):
        if len(value) > 120:
            return {"type": "text", "length": len(value), "preview": value[:32] + "…"}
        return re.sub(r"[a-f0-9]{24,}", "<redacted-id>", value, flags=re.I)
    if isinstance(value, dict):
        return {str(k): safe_log_value(v, str(k), depth + 1) for k, v in list(value.items())[:50]}
    if isinstance(value, (list, tuple, set)):
        items = list(value)
        return {"type": type(value).__name__, "count": len(items), "items": [safe_log_value(v, key, depth + 1) for v in items[:20]]}
    return {"type": type(value).__name__, "value": str(value)[:120]}


def safe_log_json(value) -> str:
    try:
        return json.dumps(safe_log_value(value), ensure_ascii=False, default=str)
    except Exception:
        return json.dumps({"type": type(value).__name__, "serialization": "failed"})


def db_parameter_summary(parameters) -> dict:
    if isinstance(parameters, dict):
        return {"count": len(parameters), "fields": list(parameters.keys())[:50], "types": {str(k): type(v).__name__ for k, v in list(parameters.items())[:50]}}
    if isinstance(parameters, (list, tuple)):
        return {"count": len(parameters), "types": [type(value).__name__ for value in parameters[:50]]}
    return {"type": type(parameters).__name__}


def db_row_summary(row) -> dict:
    if row is None:
        return {"found": False}
    if isinstance(row, sqlite3.Row):
        return {"found": True, "fields": list(row.keys())}
    if isinstance(row, (tuple, list)):
        return {"found": True, "columns": len(row), "types": [type(value).__name__ for value in row[:50]]}
    return {"found": True, "type": type(row).__name__}


def configure_logging() -> tuple[logging.Logger, logging.Logger]:
    """Configure privacy-conscious rotating application and import logs."""
    LOG_DIR.mkdir(parents=True, exist_ok=True)
    formatter = logging.Formatter(
        "%(asctime)s.%(msecs)03d %(levelname)s [%(threadName)s] "
        "%(filename)s:%(lineno)d %(funcName)s() %(message)s",
        datefmt="%Y-%m-%d %H:%M:%S",
    )

    def handler(filename: str, level: int = logging.INFO) -> TimedRotatingFileHandler:
        result = TimedRotatingFileHandler(
            LOG_DIR / filename, when="midnight", interval=1, backupCount=30,
            encoding="utf-8", delay=True,
        )
        result.setLevel(level)
        result.setFormatter(formatter)
        return result

    application = logging.getLogger("certification.server")
    application.setLevel(logging.INFO)
    application.handlers.clear()
    application.addHandler(handler("server.log"))
    application.addHandler(handler("error.log", logging.ERROR))
    application.propagate = False

    imports = logging.getLogger("certification.import")
    imports.setLevel(logging.INFO)
    imports.handlers.clear()
    imports.addHandler(handler("import.log"))
    imports.propagate = False
    return application, imports


APP_LOG, IMPORT_LOG = configure_logging()


def exception_trace(frame, event, arg):
    """Log every raised Python exception at its originating source line."""
    if event != "exception":
        return exception_trace
    exc_type, exc_value, exc_tb = arg
    if exc_type in {StopIteration, StopAsyncIteration, GeneratorExit, KeyboardInterrupt, SystemExit}:
        return exception_trace
    if getattr(TRACE_LOCAL, "logging", False):
        return exception_trace
    if Path(frame.f_code.co_filename).resolve() != (ROOT / "server.py").resolve():
        return exception_trace
    try:
        TRACE_LOCAL.logging = True
        APP_LOG.error(
            "exception_raised type=%s reason=%s source=%s:%s",
            exc_type.__name__, safe_log_value(str(exc_value)), frame.f_code.co_filename, frame.f_lineno,
        )
    finally:
        TRACE_LOCAL.logging = False
    return exception_trace


def enable_exception_tracing() -> None:
    sys.settrace(exception_trace)
    threading.settrace(exception_trace)


def dingtalk_config() -> dict:
    mode = os.environ.get("CERT_DINGTALK_MODE", "mock").strip().lower()
    return {
        "mode": mode if mode in {"mock", "live", "off"} else "mock",
        "corpId": os.environ.get("CERT_DINGTALK_CORP_ID", "").strip(),
        "clientId": os.environ.get("CERT_DINGTALK_CLIENT_ID", "").strip(),
        "clientSecret": os.environ.get("CERT_DINGTALK_CLIENT_SECRET", "").strip(),
        "agentId": os.environ.get("CERT_DINGTALK_AGENT_ID", "").strip(),
        "baseUrl": os.environ.get("CERT_BASE_URL", "http://127.0.0.1:8787").rstrip("/"),
    }

SHEET_APPLICATION = "2-1职级认证申请及成果评估表"
SHEET_EVIDENCE = "2-2员工举证附件材料区"


def now_text() -> str:
    return datetime.now().astimezone().isoformat(timespec="seconds")


def sql_operation(sql: str) -> str:
    match = re.match(r"\s*([A-Za-z]+)", sql or "")
    return match.group(1).upper() if match else "UNKNOWN"


def sql_target(sql: str) -> str:
    match = re.search(r"\b(?:FROM|INTO|UPDATE|TABLE)\s+([A-Za-z_][A-Za-z0-9_]*)", sql or "", re.I)
    return match.group(1) if match else "unknown"


class LoggedCursor(sqlite3.Cursor):
    def execute(self, sql, parameters=()):
        APP_LOG.info("db_execute operation=%s target=%s params=%s", sql_operation(sql), sql_target(sql), safe_log_json(db_parameter_summary(parameters)))
        try:
            result = super().execute(sql, parameters)
            APP_LOG.info("db_execute_complete operation=%s target=%s rowcount=%s", sql_operation(sql), sql_target(sql), self.rowcount)
            return result
        except Exception as exc:
            APP_LOG.exception("db_execute_failed operation=%s target=%s reason=%s", sql_operation(sql), sql_target(sql), safe_log_value(str(exc)))
            raise

    def executemany(self, sql, seq_of_parameters):
        parameters = list(seq_of_parameters)
        APP_LOG.info("db_executemany operation=%s target=%s items=%s parameter_shape=%s", sql_operation(sql), sql_target(sql), len(parameters), safe_log_json(db_parameter_summary(parameters[0] if parameters else ())))
        try:
            return super().executemany(sql, parameters)
        except Exception as exc:
            APP_LOG.exception("db_executemany_failed operation=%s target=%s reason=%s", sql_operation(sql), sql_target(sql), safe_log_value(str(exc)))
            raise

    def fetchone(self):
        row = super().fetchone()
        APP_LOG.info("db_fetchone data=%s", safe_log_json(db_row_summary(row)))
        return row

    def fetchall(self):
        rows = super().fetchall()
        APP_LOG.info("db_fetchall rows=%s row_shape=%s", len(rows), safe_log_json(db_row_summary(rows[0] if rows else None)))
        return rows


class LoggedConnection(sqlite3.Connection):
    def cursor(self, factory=LoggedCursor):
        return super().cursor(factory)

    def execute(self, sql, parameters=()):
        return self.cursor().execute(sql, parameters)

    def executemany(self, sql, seq_of_parameters):
        return self.cursor().executemany(sql, seq_of_parameters)

    def executescript(self, sql_script):
        APP_LOG.info("db_executescript statements=%s", sql_script.count(";"))
        try:
            return super().executescript(sql_script)
        except Exception as exc:
            APP_LOG.exception("db_executescript_failed reason=%s", safe_log_value(str(exc)))
            raise

    def commit(self):
        APP_LOG.info("db_commit")
        return super().commit()

    def rollback(self):
        APP_LOG.warning("db_rollback")
        return super().rollback()


def db() -> sqlite3.Connection:
    conn = sqlite3.connect(DB_PATH, factory=LoggedConnection)
    conn.row_factory = sqlite3.Row
    conn.execute("PRAGMA foreign_keys = ON")
    APP_LOG.info("db_open path=%s", DB_PATH)
    return conn


def init_storage() -> None:
    DATA_DIR.mkdir(exist_ok=True)
    STAGING_DIR.mkdir(exist_ok=True)
    SUBMISSION_DIR.mkdir(exist_ok=True)
    with db() as conn:
        conn.executescript(
            """
            CREATE TABLE IF NOT EXISTS submissions (
              id TEXT PRIMARY KEY,
              name TEXT NOT NULL,
              level TEXT,
              category TEXT,
              applied_at TEXT,
              source_file TEXT NOT NULL,
              status TEXT NOT NULL DEFAULT '材料已导入',
              imported_at TEXT NOT NULL,
              raw_path TEXT NOT NULL
            );
            CREATE TABLE IF NOT EXISTS modules (
              id TEXT PRIMARY KEY,
              submission_id TEXT NOT NULL REFERENCES submissions(id) ON DELETE CASCADE,
              section_type TEXT NOT NULL,
              evaluation_item TEXT,
              module_name TEXT,
              weight REAL,
              standard TEXT,
              evidence TEXT,
              self_score TEXT,
              certification_method TEXT,
              sort_order INTEGER NOT NULL
            );
            CREATE TABLE IF NOT EXISTS attachments (
              id TEXT PRIMARY KEY,
              submission_id TEXT NOT NULL REFERENCES submissions(id) ON DELETE CASCADE,
              serial TEXT,
              project_name TEXT,
              material_name TEXT,
              storage_path TEXT,
              image_path TEXT,
              uploaded_file_path TEXT,
              source_row INTEGER,
              sort_order INTEGER NOT NULL
            );
            CREATE TABLE IF NOT EXISTS module_attachments (
              module_id TEXT NOT NULL REFERENCES modules(id) ON DELETE CASCADE,
              attachment_id TEXT NOT NULL REFERENCES attachments(id) ON DELETE CASCADE,
              source TEXT NOT NULL DEFAULT '自动识别',
              PRIMARY KEY (module_id, attachment_id)
            );
            CREATE TABLE IF NOT EXISTS reviews (
              id TEXT PRIMARY KEY,
              submission_id TEXT NOT NULL REFERENCES submissions(id) ON DELETE CASCADE,
              reviewer_name TEXT NOT NULL,
              strengths TEXT, improvements TEXT, opinion TEXT, opinion_note TEXT,
              status TEXT NOT NULL DEFAULT '草稿', updated_at TEXT NOT NULL,
              UNIQUE(submission_id, reviewer_name)
            );
            CREATE TABLE IF NOT EXISTS module_scores (
              review_id TEXT NOT NULL REFERENCES reviews(id) ON DELETE CASCADE,
              module_id TEXT NOT NULL REFERENCES modules(id) ON DELETE CASCADE,
              score REAL, comment TEXT,
              PRIMARY KEY (review_id, module_id)
            );
            CREATE TABLE IF NOT EXISTS reviewers (
              id TEXT PRIMARY KEY,
              name TEXT NOT NULL UNIQUE,
              created_at TEXT NOT NULL
            );
            CREATE TABLE IF NOT EXISTS review_assignments (
              id TEXT PRIMARY KEY,
              submission_id TEXT NOT NULL REFERENCES submissions(id) ON DELETE CASCADE,
              reviewer_id TEXT NOT NULL REFERENCES reviewers(id) ON DELETE CASCADE,
              weight REAL NOT NULL,
              token TEXT NOT NULL UNIQUE,
              status TEXT NOT NULL DEFAULT '待推送',
              pushed_at TEXT,
              completed_at TEXT,
              created_at TEXT NOT NULL,
              UNIQUE(submission_id, reviewer_id)
            );
            CREATE TABLE IF NOT EXISTS directory_users (
              id TEXT PRIMARY KEY,
              provider TEXT NOT NULL DEFAULT 'dingtalk',
              external_user_id TEXT NOT NULL UNIQUE,
              name TEXT NOT NULL,
              department TEXT,
              mobile_masked TEXT,
              active INTEGER NOT NULL DEFAULT 1,
              synced_at TEXT NOT NULL
            );
            CREATE TABLE IF NOT EXISTS notification_logs (
              id TEXT PRIMARY KEY,
              assignment_id TEXT REFERENCES review_assignments(id) ON DELETE SET NULL,
              channel TEXT NOT NULL,
              recipient_external_id TEXT,
              recipient_name TEXT,
              title TEXT NOT NULL,
              content TEXT,
              target_url TEXT,
              status TEXT NOT NULL,
              provider_message_id TEXT,
              error_message TEXT,
              created_at TEXT NOT NULL
            );
            CREATE TABLE IF NOT EXISTS app_users (
              id TEXT PRIMARY KEY,
              external_user_id TEXT NOT NULL UNIQUE,
              name TEXT NOT NULL,
              role TEXT NOT NULL,
              active INTEGER NOT NULL DEFAULT 1,
              created_at TEXT NOT NULL,
              updated_at TEXT NOT NULL
            );
            CREATE TABLE IF NOT EXISTS audit_logs (
              id TEXT PRIMARY KEY,
              actor_external_id TEXT,
              actor_name TEXT,
              actor_role TEXT,
              action TEXT NOT NULL,
              entity_type TEXT,
              entity_id TEXT,
              summary TEXT,
              details_json TEXT,
              created_at TEXT NOT NULL
            );
            CREATE TABLE IF NOT EXISTS qualification_standards (
              id TEXT PRIMARY KEY,
              name TEXT NOT NULL,
              level TEXT,
              category TEXT,
              description TEXT,
              status TEXT NOT NULL DEFAULT '草稿',
              created_by TEXT,
              created_at TEXT NOT NULL,
              updated_at TEXT NOT NULL
            );
            CREATE TABLE IF NOT EXISTS job_taxonomy (
              id TEXT PRIMARY KEY,
              name TEXT NOT NULL,
              kind TEXT NOT NULL CHECK(kind IN ('family','category','subcategory')),
              parent_id TEXT REFERENCES job_taxonomy(id) ON DELETE RESTRICT,
              created_at TEXT NOT NULL,
              active INTEGER NOT NULL DEFAULT 1
            );
            CREATE UNIQUE INDEX IF NOT EXISTS taxonomy_sibling_name
              ON job_taxonomy(COALESCE(parent_id,''),name);
            CREATE TABLE IF NOT EXISTS standard_modules (
              id TEXT PRIMARY KEY,
              standard_id TEXT NOT NULL REFERENCES qualification_standards(id) ON DELETE CASCADE,
              section_type TEXT NOT NULL,
              evaluation_item TEXT,
              module_name TEXT NOT NULL,
              weight REAL,
              standard TEXT,
              certification_method TEXT,
              sort_order INTEGER NOT NULL
            );
            """
        )
        taxonomy_columns = {row["name"] for row in conn.execute("PRAGMA table_info(job_taxonomy)")}
        if "active" not in taxonomy_columns:
            conn.execute("ALTER TABLE job_taxonomy ADD COLUMN active INTEGER NOT NULL DEFAULT 1")
        submission_columns = {row["name"] for row in conn.execute("PRAGMA table_info(submissions)")}
        standard_columns = {row["name"] for row in conn.execute("PRAGMA table_info(qualification_standards)")}
        if "subcategory_id" not in standard_columns:
            conn.execute("ALTER TABLE qualification_standards ADD COLUMN subcategory_id TEXT REFERENCES job_taxonomy(id)")
        if "framework_json" not in standard_columns:
            conn.execute("ALTER TABLE qualification_standards ADD COLUMN framework_json TEXT NOT NULL DEFAULT '{}'")
        if "matrix_json" not in standard_columns:
            conn.execute("ALTER TABLE qualification_standards ADD COLUMN matrix_json TEXT")
        if "positions_json" not in standard_columns:
            conn.execute("ALTER TABLE qualification_standards ADD COLUMN positions_json TEXT NOT NULL DEFAULT '{}'")
        conn.execute("CREATE TABLE IF NOT EXISTS qualification_template(id INTEGER PRIMARY KEY CHECK(id=1),matrix_json TEXT NOT NULL,revision INTEGER NOT NULL DEFAULT 1)")
        conn.execute("INSERT OR IGNORE INTO qualification_template(id,matrix_json) VALUES(1,?)", (json.dumps(default_matrix(), ensure_ascii=False),))
        conn.execute("CREATE TABLE IF NOT EXISTS job_positions(subcategory_id TEXT NOT NULL REFERENCES job_taxonomy(id) ON DELETE CASCADE,grade INTEGER NOT NULL CHECK(grade BETWEEN 1 AND 7),name TEXT NOT NULL,PRIMARY KEY(subcategory_id,grade))")
        for column, definition in {
            "series_id": "TEXT", "version_number": "INTEGER NOT NULL DEFAULT 1",
            "parent_version_id": "TEXT", "revision": "INTEGER NOT NULL DEFAULT 0",
            "archived_at": "TEXT", "archived_by": "TEXT",
            "classification_json": "TEXT NOT NULL DEFAULT '{}'",
            "change_note": "TEXT NOT NULL DEFAULT ''",
        }.items():
            if column not in standard_columns:
                conn.execute(f"ALTER TABLE qualification_standards ADD COLUMN {column} {definition}")
        conn.execute("CREATE TABLE IF NOT EXISTS standard_history (id INTEGER PRIMARY KEY AUTOINCREMENT, standard_id TEXT NOT NULL, series_id TEXT NOT NULL, action TEXT NOT NULL, actor_id TEXT, actor_name TEXT, created_at TEXT NOT NULL, note TEXT, snapshot_json TEXT NOT NULL)")
        evidence.init_storage(conn)
        init_personnel_storage(conn)
        # Legacy publication is not an explicit finalization. Keep content and references,
        # but require an administrator to archive before accepting new applications.
        for legacy in conn.execute("SELECT * FROM qualification_standards WHERE series_id IS NULL").fetchall():
            conn.execute("UPDATE qualification_standards SET series_id=id WHERE id=?", (legacy["id"],))
            record_standard_history(conn, legacy["id"], "迁移原标准", {"name": "系统迁移"}, "保留原有内容和状态；须定稿存档后方可新增引用")
        conn.execute("CREATE UNIQUE INDEX IF NOT EXISTS standard_version_unique ON qualification_standards(series_id,version_number)")
        if "standard_snapshot_json" not in submission_columns:
            conn.execute("ALTER TABLE submissions ADD COLUMN standard_snapshot_json TEXT NOT NULL DEFAULT '{}'")
        if "owner_external_user_id" not in submission_columns:
            conn.execute("ALTER TABLE submissions ADD COLUMN owner_external_user_id TEXT")
        if "entry_mode" not in submission_columns:
            conn.execute("ALTER TABLE submissions ADD COLUMN entry_mode TEXT NOT NULL DEFAULT 'excel'")
        if "standard_id" not in submission_columns:
            conn.execute("ALTER TABLE submissions ADD COLUMN standard_id TEXT")
        if "submitted_at" not in submission_columns:
            conn.execute("ALTER TABLE submissions ADD COLUMN submitted_at TEXT")
        review_columns = {row["name"] for row in conn.execute("PRAGMA table_info(reviews)")}
        if "assignment_id" not in review_columns:
            conn.execute("ALTER TABLE reviews ADD COLUMN assignment_id TEXT")
        attachment_columns = {row["name"] for row in conn.execute("PRAGMA table_info(attachments)")}
        if "link_url" not in attachment_columns:
            conn.execute("ALTER TABLE attachments ADD COLUMN link_url TEXT")
        # Enforce exclusive ownership for online evidence, retaining legacy Excel links.
        for operation in ('INSERT', 'UPDATE'):
            conn.execute(f'''CREATE TRIGGER IF NOT EXISTS online_attachment_owner_{operation.lower()}
                BEFORE {operation} ON module_attachments
                WHEN EXISTS (SELECT 1 FROM attachments a JOIN submissions s ON s.id=a.submission_id
                             WHERE a.id=NEW.attachment_id AND s.entry_mode='online')
                BEGIN
                    SELECT CASE WHEN NOT EXISTS (
                        SELECT 1 FROM attachments a JOIN modules m ON m.submission_id=a.submission_id
                        WHERE a.id=NEW.attachment_id AND m.id=NEW.module_id
                    ) THEN RAISE(ABORT, '附件与举证项不属于同一申请') END;
                    SELECT CASE WHEN EXISTS (
                        SELECT 1 FROM module_attachments WHERE attachment_id=NEW.attachment_id
                        AND module_id<>NEW.module_id
                    ) THEN RAISE(ABORT, '同一附件只能对应一个举证项') END;
                END''')
        reviewer_columns = {row["name"] for row in conn.execute("PRAGMA table_info(reviewers)")}
        if "external_user_id" not in reviewer_columns:
            conn.execute("ALTER TABLE reviewers ADD COLUMN external_user_id TEXT")
        if "directory_provider" not in reviewer_columns:
            conn.execute("ALTER TABLE reviewers ADD COLUMN directory_provider TEXT")
        if "access_code_salt" not in reviewer_columns:
            conn.execute("ALTER TABLE reviewers ADD COLUMN access_code_salt TEXT")
        if "access_code_hash" not in reviewer_columns:
            conn.execute("ALTER TABLE reviewers ADD COLUMN access_code_hash TEXT")
        conn.execute("""CREATE TABLE IF NOT EXISTS reviewer_sessions (
                          token TEXT PRIMARY KEY,
                          reviewer_id TEXT NOT NULL REFERENCES reviewers(id) ON DELETE CASCADE,
                          expires_at TEXT NOT NULL,
                          created_at TEXT NOT NULL
                        )""")
        if dingtalk_config()["mode"] == "mock":
            mock_users = [
                ("mock-admin", "模拟管理员", "管理员", "认证运营组"),
                ("mock-reviewer", "模拟评委", "评委", "专业评审组"),
                ("mock-student", "模拟学员", "学员", "认证申请人"),
            ]
            for external_id, name, role, department in mock_users:
                conn.execute(
                    """INSERT INTO directory_users(id,provider,external_user_id,name,department,mobile_masked,active,synced_at)
                       VALUES(?,?,?,?,?,?,1,?) ON CONFLICT(external_user_id) DO UPDATE SET name=excluded.name,department=excluded.department,active=1,synced_at=excluded.synced_at""",
                    (uuid.uuid4().hex, "dingtalk-mock", external_id, name, department, "", now_text()),
                )
                conn.execute(
                    """INSERT INTO app_users(id,external_user_id,name,role,active,created_at,updated_at)
                       VALUES(?,?,?,?,1,?,?) ON CONFLICT(external_user_id) DO UPDATE SET name=excluded.name,role=excluded.role,active=1,updated_at=excluded.updated_at""",
                    (uuid.uuid4().hex, external_id, name, role, now_text(), now_text()),
                )
        conn.execute("""UPDATE modules SET module_name='知识产权', evaluation_item='知识成果贡献'
                        WHERE section_type='bonus' AND module_name LIKE '知识成果贡献%' AND standard LIKE '%逐一举证%'""")
        conn.execute("""UPDATE modules SET module_name='标准制定', evaluation_item='团队成长贡献'
                        WHERE section_type='bonus' AND module_name LIKE '团队成长贡献%' AND evidence<>''""")
        conn.execute("""UPDATE modules SET evaluation_item='知识成果贡献'
                        WHERE section_type='bonus' AND module_name IN ('知识产权','成果输出','创新成果','学术论文')""")
        conn.execute("""UPDATE modules SET evaluation_item='团队成长贡献'
                        WHERE section_type='bonus' AND (module_name='标准制定' OR module_name LIKE '担任新员工导师%')""")


def clean(value):
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d")
    return str(value).strip()


def audit_event(action: str, entity_type: str = "", entity_id: str = "", summary: str = "", details: dict | None = None, actor: dict | None = None) -> None:
    with db() as conn:
        conn.execute(
            """INSERT INTO audit_logs(id,actor_external_id,actor_name,actor_role,action,entity_type,entity_id,summary,details_json,created_at)
               VALUES(?,?,?,?,?,?,?,?,?,?)""",
            (uuid.uuid4().hex, actor.get("userId") if actor else None, actor.get("name") if actor else None,
             actor.get("role") if actor else None, action, entity_type, entity_id, summary,
             json.dumps(details or {}, ensure_ascii=False), now_text()),
        )


def number(value):
    if value is None or value == "":
        return None
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value).strip().replace("%", "")
    try:
        parsed = float(text)
        return parsed / 100 if "%" in str(value) else parsed
    except ValueError:
        return None


def sheet_by_expected_name(workbook, expected: str, index: int):
    if expected in workbook.sheetnames:
        return workbook[expected]
    for name in workbook.sheetnames:
        if expected.split("员工")[0].split("职级")[0] in name:
            return workbook[name]
    return workbook.worksheets[index] if len(workbook.worksheets) > index else None


def find_label_value(sheet, label: str, max_rows: int = 15, max_cols: int = 12) -> str:
    for row in range(1, min(sheet.max_row, max_rows) + 1):
        for col in range(1, min(sheet.max_column, max_cols) + 1):
            if label in clean(sheet.cell(row, col).value):
                return clean(sheet.cell(row, col + 1).value) if col < max_cols else ""
    return ""


def infer_name_from_filename(filename: str) -> str:
    stem = Path(filename).stem
    match = re.search(r"[-—]\s*([\u4e00-\u9fff]{2,5})(?:\d+(?:\.\d+)?)?$", stem)
    return match.group(1) if match else ""


def find_structure_rows(sheet) -> tuple[int, int, int]:
    header_row = 0
    bonus_row = 0
    summary_row = 0
    for row in range(1, min(sheet.max_row, 80) + 1):
        first = clean(sheet.cell(row, 1).value)
        second = clean(sheet.cell(row, 2).value)
        if not header_row and "评估项目" in first and "工作模块" in second:
            header_row = row
        if header_row and not bonus_row and "团队成长贡献加分项" in first:
            bonus_row = row
        if header_row and "综合得分" in first:
            summary_row = row
            break
    return header_row, bonus_row, summary_row


def extract_images(sheet, output_dir: Path) -> dict[int, str]:
    images: dict[int, str] = {}
    output_dir.mkdir(parents=True, exist_ok=True)
    for image in getattr(sheet, "_images", []):
        anchor = getattr(image, "anchor", None)
        marker = getattr(anchor, "_from", None)
        if marker is None:
            continue
        row = int(marker.row) + 1
        ext = (getattr(image, "format", None) or "png").lower()
        filename = f"row-{row}-{uuid.uuid4().hex[:8]}.{ext}"
        try:
            (output_dir / filename).write_bytes(image._data())
            images[row] = filename
        except Exception:
            continue
    return images


def parse_workbook(path: Path, batch_id: str) -> dict:
    workbook = load_workbook(path, data_only=False)
    warnings = []
    errors = []
    app = sheet_by_expected_name(workbook, SHEET_APPLICATION, 0)
    evidence_sheet = sheet_by_expected_name(workbook, SHEET_EVIDENCE, 1)
    if app is None:
        return {"errors": ["未找到申报信息 Sheet"], "warnings": [], "modules": [], "attachments": []}
    if evidence_sheet is None:
        errors.append("未找到附件材料 Sheet")

    name = find_label_value(app, "姓名")
    level = find_label_value(app, "申请认证级别")
    category = find_label_value(app, "对应职类")
    applied_at = find_label_value(app, "申请认证时间")
    if not name:
        inferred_name = infer_name_from_filename(path.name)
        if inferred_name:
            name = inferred_name
            warnings.append(f"Excel 内姓名为空，已根据文件名识别为“{name}”，请确认")
        else:
            errors.append("姓名为空")
    if not level:
        errors.append("申请认证级别为空")

    modules = []
    normal_weights = []
    header_row, bonus_start, summary_row = find_structure_rows(app)
    if not header_row:
        errors.append("未识别到包含“评估项目/工作模块”的表头")
        header_row = 4
    if not summary_row:
        errors.append("未识别到“综合得分”结束行")
        summary_row = min(app.max_row, header_row + 20)
    if not bonus_start:
        warnings.append("未识别到团队成长贡献板块")
        bonus_start = summary_row

    for row in range(header_row + 1, bonus_start):
        module_name = clean(app.cell(row, 2).value)
        if not module_name or "填写要求" in module_name:
            continue
        standard = clean(app.cell(row, 4).value)
        description = clean(app.cell(row, 5).value)
        self_score = clean(app.cell(row, 7).value)
        method = clean(app.cell(row, 8).value)
        weight = number(app.cell(row, 3).value)
        if weight is not None:
            normal_weights.append(weight)
        refs = sorted(set(re.findall(r"关联附件材料\s*([0-9]+)", description)))
        modules.append(
            {
                "sourceRow": row,
                "sectionType": "standard",
                "evaluationItem": "行为标准",
                "moduleName": module_name,
                "weight": weight,
                "maxScore": 5,
                "standard": standard,
                "evidence": description,
                "selfScore": self_score,
                "certificationMethod": method,
                "attachmentRefs": refs,
            }
        )

    current_bonus_group = ""
    for row in range(bonus_start, summary_row):
        group = clean(app.cell(row, 2).value)
        if group:
            current_bonus_group = re.sub(r"\s*（每项最多2分）\s*", "", group).strip()
        item_name = clean(app.cell(row, 3).value)
        if not item_name:
            continue
        description = clean(app.cell(row, 5).value)
        refs = sorted(set(re.findall(r"关联附件材料\s*([0-9]+)", description)))
        modules.append({
            "sourceRow": row,
            "sectionType": "bonus",
            "evaluationItem": current_bonus_group or "团队成长贡献",
            "moduleName": item_name,
            "weight": None,
            "maxScore": 2,
            "standard": clean(app.cell(row, 4).value),
            "evidence": description,
            "selfScore": clean(app.cell(row, 7).value),
            "certificationMethod": clean(app.cell(row, 8).value),
            "attachmentRefs": refs,
        })

    if normal_weights and abs(sum(normal_weights) - 1) > 0.001:
        warnings.append(f"普通工作模块权重合计为 {sum(normal_weights):.0%}，不是 100%")

    attachments = []
    image_dir = STAGING_DIR / batch_id / "images" / path.stem
    image_map = extract_images(evidence_sheet, image_dir) if evidence_sheet else {}
    if evidence_sheet:
        blank_run = 0
        current_project = ""
        for row in range(2, min(evidence_sheet.max_row, 2000) + 1):
            serial = clean(evidence_sheet.cell(row, 1).value)
            explicit_project = clean(evidence_sheet.cell(row, 2).value)
            if explicit_project:
                current_project = explicit_project
            project = explicit_project or current_project
            material = clean(evidence_sheet.cell(row, 3).value)
            storage_path = clean(evidence_sheet.cell(row, 4).value)
            link_cell = evidence_sheet.cell(row, 5)
            link_url = clean(link_cell.hyperlink.target if link_cell.hyperlink else link_cell.value)
            if not any((serial, project, material, storage_path, link_url, row in image_map)):
                blank_run += 1
                if blank_run >= 25:
                    break
                continue
            blank_run = 0
            attachments.append(
                {
                    "sourceRow": row,
                    "serial": serial,
                    "projectName": project,
                    "materialName": material,
                    "storagePath": storage_path,
                    "linkUrl": link_url,
                    "imageFile": image_map.get(row),
                }
            )

    serials = [item["serial"] for item in attachments if item["serial"]]
    duplicates = sorted({x for x in serials if serials.count(x) > 1})
    if duplicates:
        errors.append("附件序号重复：" + "、".join(duplicates))
    known = set(serials)
    referenced = {ref for module in modules for ref in module["attachmentRefs"]}
    missing = sorted(referenced - known)
    if missing:
        warnings.append("正文引用但附件表不存在：" + "、".join(missing))
    unused = sorted(known - referenced)
    if unused:
        warnings.append("未被工作模块引用的附件：" + "、".join(unused))

    result = {
        "name": name,
        "level": level,
        "category": category,
        "appliedAt": applied_at,
        "modules": modules,
        "attachments": attachments,
        "warnings": warnings,
        "errors": errors,
        "sheetNames": workbook.sheetnames,
    }
    workbook.close()
    return result


def save_preview(files: list[dict]) -> dict:
    batch_id = uuid.uuid4().hex
    batch_dir = STAGING_DIR / batch_id
    batch_dir.mkdir(parents=True)
    results = []
    for index, item in enumerate(files):
        filename = Path(item.get("name", f"workbook-{index + 1}.xlsx")).name
        if not filename.lower().endswith(".xlsx"):
            results.append({"fileName": filename, "errors": ["仅支持 .xlsx 文件"], "warnings": []})
            continue
        raw = item.get("data", "")
        if "," in raw:
            raw = raw.split(",", 1)[1]
        try:
            content = base64.b64decode(raw, validate=True)
            stored_name = f"{index:03d}-{filename}"
            workbook_path = batch_dir / stored_name
            workbook_path.write_bytes(content)
            parsed = parse_workbook(workbook_path, batch_id)
            parsed.update({"fileName": filename, "storedName": stored_name, "selected": not parsed["errors"]})
            results.append(parsed)
        except Exception as exc:
            results.append({"fileName": filename, "errors": [f"Excel 解析失败：{exc}"], "warnings": [], "selected": False})
    payload = {"batchId": batch_id, "createdAt": now_text(), "items": results}
    (batch_dir / "preview.json").write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    return payload


def create_import_batch() -> str:
    batch_id = uuid.uuid4().hex
    (STAGING_DIR / batch_id).mkdir(parents=True)
    return batch_id


def store_batch_stream(batch_id: str, filename: str, source, length: int) -> dict:
    batch_dir = STAGING_DIR / batch_id
    if not batch_dir.is_dir():
        raise ValueError("导入批次不存在")
    filename = Path(filename).name
    if not filename.lower().endswith(".xlsx"):
        raise ValueError("仅支持 .xlsx 文件")
    existing = len(list(batch_dir.glob("*.xlsx")))
    target = batch_dir / f"{existing:03d}-{filename}"
    remaining = length
    with target.open("wb") as output:
        while remaining:
            chunk = source.read(min(UPLOAD_CHUNK, remaining))
            if not chunk:
                raise ValueError("上传连接提前中断")
            output.write(chunk)
            remaining -= len(chunk)
    return {"storedName": target.name, "size": length}


def embedded_anchor_map(archive: zipfile.ZipFile) -> dict[str, dict]:
    """Map embedded package paths to their worksheet cell anchors."""
    mapping = {}
    rel_ns = {"r": "http://schemas.openxmlformats.org/package/2006/relationships"}
    office_rel = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id"
    for sheet_name in [name for name in archive.namelist() if re.fullmatch(r"xl/worksheets/sheet\d+\.xml", name)]:
        rel_name = posixpath.join(posixpath.dirname(sheet_name), "_rels", posixpath.basename(sheet_name) + ".rels")
        if rel_name not in archive.namelist():
            continue
        rel_root = ET.fromstring(archive.read(rel_name))
        relationships = {node.attrib.get("Id"): node.attrib.get("Target", "") for node in rel_root.findall("r:Relationship", rel_ns)}
        sheet_root = ET.fromstring(archive.read(sheet_name))
        seen = set()
        for ole in sheet_root.iter():
            if ole.tag.rsplit("}", 1)[-1] != "oleObject":
                continue
            rel_id = ole.attrib.get(office_rel)
            if not rel_id or rel_id in seen or rel_id not in relationships:
                continue
            anchor_row = anchor_col = None
            for node in ole.iter():
                if node.tag.rsplit("}", 1)[-1] == "from":
                    for child in node:
                        local = child.tag.rsplit("}", 1)[-1]
                        if local == "row": anchor_row = int(child.text or 0) + 1
                        elif local == "col": anchor_col = int(child.text or 0) + 1
                    break
            target = relationships[rel_id]
            normalized = posixpath.normpath(posixpath.join(posixpath.dirname(sheet_name), target))
            mapping[normalized] = {"sourceRow": anchor_row, "sourceColumn": anchor_col, "progId": ole.attrib.get("progId", ""), "relationshipId": rel_id}
            seen.add(rel_id)
    return mapping


def inspect_workbook_package(path: Path, batch_id: str, extract: bool = True) -> dict:
    result = {"fileSize": path.stat().st_size, "images": 0, "imageBytes": 0, "embeddedObjects": [], "risk": "常规"}
    try:
        with zipfile.ZipFile(path) as archive:
            entries = archive.infolist()
            anchors = embedded_anchor_map(archive)
            for entry in entries:
                if entry.filename.startswith("xl/media/") and not entry.is_dir():
                    result["images"] += 1
                    result["imageBytes"] += entry.file_size
                if entry.filename.startswith("xl/embeddings/") and not entry.is_dir():
                    raw = archive.read(entry)
                    extension, kind = ".bin", "OLE 嵌入对象"
                    if raw.startswith(b"PK\x03\x04"):
                        names = []
                        try:
                            import io
                            with zipfile.ZipFile(io.BytesIO(raw)) as nested:
                                names = nested.namelist()
                        except Exception:
                            names = []
                        if any(name.startswith("word/") for name in names): extension, kind = ".docx", "Word 文档"
                        elif any(name.startswith("ppt/") for name in names): extension, kind = ".pptx", "PowerPoint 演示文稿"
                        elif any(name.startswith("xl/") for name in names): extension, kind = ".xlsx", "Excel 工作簿"
                        else: extension, kind = ".zip", "压缩包或嵌入包"
                    index = len(result["embeddedObjects"]) + 1
                    saved_name = f"embedded-{index}{extension}"
                    relative = None
                    if extract:
                        target_dir = STAGING_DIR / batch_id / "embeddings" / path.stem
                        target_dir.mkdir(parents=True, exist_ok=True)
                        target = target_dir / saved_name
                        target.write_bytes(raw)
                        relative = str(target.relative_to(STAGING_DIR / batch_id))
                    anchor = anchors.get(posixpath.normpath(entry.filename), {})
                    result["embeddedObjects"].append({"name": saved_name, "kind": kind, "size": entry.file_size, "sourceEntry": entry.filename, "stagedPath": relative, **anchor})
            total_uncompressed = sum(entry.file_size for entry in entries)
            if result["embeddedObjects"] or path.stat().st_size > 50 * 1024 * 1024:
                result["risk"] = "复杂"
            elif total_uncompressed > 150 * 1024 * 1024 or result["images"] > 30:
                result["risk"] = "较大"
    except zipfile.BadZipFile:
        header = path.read_bytes()[:32]
        if header.startswith(b"%TSD-Header-"):
            raise ValueError("文件是加密或受保护的 TSD 封装，并非标准 xlsx；请在原应用中打开后另存为标准 .xlsx 再上传")
        raise ValueError("文件不是有效的标准 xlsx；请在 Excel/WPS 中打开并另存为 .xlsx 后重试")
    return result


def job_path(batch_id: str) -> Path:
    return STAGING_DIR / batch_id / "job.json"


def write_job(batch_id: str, payload: dict) -> None:
    with JOB_LOCK:
        target = job_path(batch_id)
        temp = target.with_suffix(".tmp")
        temp.write_text(json.dumps(payload, ensure_ascii=False), encoding="utf-8")
        temp.replace(target)


def parse_one_workbook_to_file(workbook_path: Path, batch_id: str, output_path: Path) -> None:
    filename = workbook_path.name.split("-", 1)[1] if "-" in workbook_path.name else workbook_path.name
    package = inspect_workbook_package(workbook_path, batch_id)
    parsed = parse_workbook(workbook_path, batch_id)
    attachments_by_row = {item.get("sourceRow"): item for item in parsed.get("attachments", [])}
    for embedded in package.get("embeddedObjects", []):
        matched = attachments_by_row.get(embedded.get("sourceRow"))
        if matched:
            embedded.update({
                "matched": True, "serial": matched.get("serial"),
                "projectName": matched.get("projectName"), "materialName": matched.get("materialName"),
                "storagePath": matched.get("storagePath"), "linkUrl": matched.get("linkUrl"),
            })
    parsed.update({"fileName": filename, "storedName": workbook_path.name, "selected": not parsed["errors"], "package": package})
    output_path.write_text(json.dumps(parsed, ensure_ascii=False), encoding="utf-8")


def parse_one_workbook_isolated(workbook_path: Path, batch_id: str) -> dict:
    output_path = STAGING_DIR / batch_id / f"parse-{uuid.uuid4().hex}.json"
    command = [sys.executable]
    if not getattr(sys, "frozen", False):
        command.append(str(ROOT / "server.py"))
    command.extend(["--parse-one", str(workbook_path), batch_id, str(output_path)])
    try:
        IMPORT_LOG.info("parse_start batch=%s file=%s size_bytes=%s", batch_id, workbook_path.name, workbook_path.stat().st_size)
        completed = subprocess.run(command, capture_output=True, text=True, timeout=PARSE_TIMEOUT_SECONDS)
        if completed.returncode != 0:
            message = (completed.stderr or completed.stdout or "解析子进程异常退出").strip().splitlines()[-1]
            raise RuntimeError(message)
        if not output_path.exists():
            raise RuntimeError("解析子进程没有生成结果")
        result = json.loads(output_path.read_text(encoding="utf-8"))
        IMPORT_LOG.info("parse_success batch=%s file=%s", batch_id, workbook_path.name)
        return result
    except subprocess.TimeoutExpired as exc:
        IMPORT_LOG.error("parse_timeout batch=%s file=%s timeout_seconds=%s", batch_id, workbook_path.name, PARSE_TIMEOUT_SECONDS)
        APP_LOG.exception("Excel parser timeout batch=%s file=%s", batch_id, workbook_path.name)
        raise TimeoutError(f"单文件解析超过 {PARSE_TIMEOUT_SECONDS} 秒，已自动跳过") from exc
    except Exception:
        IMPORT_LOG.exception("parse_failed batch=%s file=%s", batch_id, workbook_path.name)
        APP_LOG.exception("Excel parser failed batch=%s file=%s", batch_id, workbook_path.name)
        raise
    finally:
        output_path.unlink(missing_ok=True)


def start_parse_job(batch_id: str) -> dict:
    batch_dir = STAGING_DIR / batch_id
    if not batch_dir.is_dir():
        raise ValueError("导入批次不存在")
    files = sorted(batch_dir.glob("*.xlsx"))
    if not files:
        raise ValueError("当前批次没有可解析文件")
    existing = job_path(batch_id)
    if existing.exists():
        current = json.loads(existing.read_text(encoding="utf-8"))
        if current.get("status") in {"等待解析", "解析中"}:
            return current
    state = {"batchId": batch_id, "status": "等待解析", "total": len(files), "completed": 0, "currentFile": "", "startedAt": now_text(), "result": None}
    write_job(batch_id, state)

    def worker():
        results = []
        try:
            IMPORT_LOG.info("batch_start batch=%s files=%s", batch_id, len(files))
            state["status"] = "解析中"
            write_job(batch_id, state)
            for index, workbook_path in enumerate(files):
                filename = workbook_path.name.split("-", 1)[1] if "-" in workbook_path.name else workbook_path.name
                state.update({"completed": index, "currentFile": filename, "updatedAt": now_text()})
                write_job(batch_id, state)
                started = time.monotonic()
                try:
                    parsed = parse_one_workbook_isolated(workbook_path, batch_id)
                    parsed["parseSeconds"] = round(time.monotonic() - started, 2)
                    results.append(parsed)
                except Exception as exc:
                    results.append({"fileName": filename, "storedName": workbook_path.name, "errors": [f"Excel 解析失败：{exc}"], "warnings": [], "selected": False, "parseSeconds": round(time.monotonic() - started, 2), "package": {"fileSize": workbook_path.stat().st_size, "risk": "已跳过", "embeddedObjects": []}})
                    IMPORT_LOG.warning("file_skipped batch=%s file=%s reason=%s", batch_id, workbook_path.name, type(exc).__name__)
                state.update({"completed": index + 1, "currentFile": "", "updatedAt": now_text()})
                write_job(batch_id, state)
            payload = {"batchId": batch_id, "createdAt": now_text(), "items": results}
            (batch_dir / "preview.json").write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
            state.update({"status": "已完成", "completed": len(files), "currentFile": "", "finishedAt": now_text(), "result": parse_import_batch_result(payload)})
            IMPORT_LOG.info("batch_complete batch=%s files=%s failed=%s", batch_id, len(files), sum(1 for item in results if item.get("errors")))
        except BaseException as exc:
            state.update({"status": "任务异常", "error": f"后台任务异常：{type(exc).__name__}: {exc}", "currentFile": "", "finishedAt": now_text()})
            IMPORT_LOG.exception("batch_failed batch=%s", batch_id)
            APP_LOG.exception("Background import job failed batch=%s", batch_id)
        finally:
            write_job(batch_id, state)

    threading.Thread(target=worker, name=f"parse-{batch_id[:8]}", daemon=True).start()
    return state


def parse_import_batch_result(payload: dict) -> dict:
    return {
        "batchId": payload["batchId"], "createdAt": payload["createdAt"],
        "items": [{
            "fileName": item.get("fileName"), "name": item.get("name"), "level": item.get("level"),
            "category": item.get("category"), "appliedAt": item.get("appliedAt"),
            "warnings": item.get("warnings", []), "errors": item.get("errors", []), "selected": item.get("selected", False),
            "parseSeconds": item.get("parseSeconds"),
            "package": item.get("package", {}),
            "modules": [{"sectionType": m.get("sectionType"), "evaluationItem": m.get("evaluationItem"), "moduleName": m.get("moduleName"), "weight": m.get("weight"), "maxScore": m.get("maxScore")} for m in item.get("modules", [])],
            "attachments": [{"serial": a.get("serial"), "materialName": a.get("materialName"), "imageFile": bool(a.get("imageFile"))} for a in item.get("attachments", [])],
        } for item in payload["items"]],
    }


def parse_import_batch(batch_id: str) -> dict:
    batch_dir = STAGING_DIR / batch_id
    if not batch_dir.is_dir():
        raise ValueError("导入批次不存在")
    results = []
    for workbook_path in sorted(batch_dir.glob("*.xlsx")):
        filename = workbook_path.name.split("-", 1)[1] if "-" in workbook_path.name else workbook_path.name
        try:
            parsed = parse_workbook(workbook_path, batch_id)
            parsed.update({"fileName": filename, "storedName": workbook_path.name, "selected": not parsed["errors"]})
            results.append(parsed)
        except Exception as exc:
            results.append({"fileName": filename, "storedName": workbook_path.name, "errors": [f"Excel 解析失败：{exc}"], "warnings": [], "selected": False})
    payload = {"batchId": batch_id, "createdAt": now_text(), "items": results}
    (batch_dir / "preview.json").write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    return {
        "batchId": batch_id,
        "createdAt": payload["createdAt"],
        "items": [
            {
                "fileName": item.get("fileName"), "name": item.get("name"), "level": item.get("level"),
                "category": item.get("category"), "appliedAt": item.get("appliedAt"),
                "warnings": item.get("warnings", []), "errors": item.get("errors", []),
                "selected": item.get("selected", False),
                "modules": [{"sectionType": m.get("sectionType"), "evaluationItem": m.get("evaluationItem"), "moduleName": m.get("moduleName"), "weight": m.get("weight"), "maxScore": m.get("maxScore")} for m in item.get("modules", [])],
                "attachments": [{"serial": a.get("serial"), "materialName": a.get("materialName"), "imageFile": bool(a.get("imageFile"))} for a in item.get("attachments", [])],
            }
            for item in results
        ],
    }


def confirm_batch(batch_id: str, indexes: list[int] | None) -> list[str]:
    batch_dir = STAGING_DIR / batch_id
    preview_path = batch_dir / "preview.json"
    if not preview_path.exists():
        raise ValueError("导入批次不存在或已失效")
    payload = json.loads(preview_path.read_text(encoding="utf-8"))
    selected = set(indexes if indexes is not None else range(len(payload["items"])))
    created = []
    with db() as conn:
        for index, item in enumerate(payload["items"]):
            if index not in selected or item.get("errors"):
                continue
            submission_id = uuid.uuid4().hex
            target_dir = SUBMISSION_DIR / submission_id
            target_dir.mkdir(parents=True)
            source_path = batch_dir / item["storedName"]
            raw_target = target_dir / item["fileName"]
            shutil.copy2(source_path, raw_target)
            conn.execute(
                """INSERT INTO submissions(id,name,level,category,applied_at,source_file,status,imported_at,raw_path)
                   VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)""",
                (submission_id, item["name"], item["level"], item["category"], item["appliedAt"], item["fileName"], "材料已导入", now_text(), str(raw_target.relative_to(DATA_DIR))),
            )
            attachment_ids = {}
            attachment_ids_by_row = {}
            for order, attachment in enumerate(item["attachments"]):
                attachment_id = uuid.uuid4().hex
                image_path = None
                if attachment.get("imageFile"):
                    image_source = batch_dir / "images" / Path(item["storedName"]).stem / attachment["imageFile"]
                    if not image_source.exists():
                        image_source = batch_dir / "images" / Path(item["fileName"]).stem / attachment["imageFile"]
                    if image_source.exists():
                        image_target_dir = target_dir / "images"
                        image_target_dir.mkdir(exist_ok=True)
                        image_target = image_target_dir / attachment["imageFile"]
                        shutil.copy2(image_source, image_target)
                        image_path = str(image_target.relative_to(DATA_DIR))
                conn.execute(
                    """INSERT INTO attachments(id,submission_id,serial,project_name,material_name,storage_path,image_path,uploaded_file_path,source_row,sort_order,link_url)
                       VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)""",
                    (attachment_id, submission_id, attachment["serial"], attachment["projectName"], attachment["materialName"], attachment["storagePath"], image_path, None, attachment["sourceRow"], order, attachment.get("linkUrl")),
                )
                attachment_ids[attachment["serial"]] = attachment_id
                attachment_ids_by_row[attachment.get("sourceRow")] = attachment_id
            for embedded_index, embedded in enumerate(item.get("package", {}).get("embeddedObjects", []), start=1):
                staged_relative = embedded.get("stagedPath")
                staged_source = batch_dir / staged_relative if staged_relative else None
                uploaded_path = None
                if staged_source and staged_source.is_file() and batch_dir.resolve() in staged_source.resolve().parents:
                    embedded_dir = target_dir / "attachments"
                    embedded_dir.mkdir(exist_ok=True)
                    preferred = Path(clean(embedded.get("storagePath"))).name
                    original_suffix = Path(embedded.get("name") or staged_source.name).suffix
                    if not preferred or preferred.lower().startswith(("http://", "https://")):
                        preferred = Path(embedded.get("name") or staged_source.name).name
                    elif not Path(preferred).suffix:
                        preferred += original_suffix
                    embedded_target = embedded_dir / preferred
                    if embedded_target.exists():
                        embedded_target = embedded_dir / f"{embedded_target.stem}-{embedded_index}{embedded_target.suffix}"
                    shutil.copy2(staged_source, embedded_target)
                    uploaded_path = str(embedded_target.relative_to(DATA_DIR))
                matched_attachment_id = attachment_ids_by_row.get(embedded.get("sourceRow")) if embedded.get("matched") else None
                if matched_attachment_id:
                    conn.execute("UPDATE attachments SET uploaded_file_path=? WHERE id=?", (uploaded_path, matched_attachment_id))
                else:
                    attachment_id = uuid.uuid4().hex
                    conn.execute(
                        """INSERT INTO attachments(id,submission_id,serial,project_name,material_name,storage_path,image_path,uploaded_file_path,source_row,sort_order,link_url)
                           VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)""",
                        (attachment_id, submission_id, f"自动-{embedded_index}", embedded.get("projectName") or "Excel 嵌入对象", embedded.get("materialName") or embedded.get("name"), embedded.get("storagePath") or "从 Excel 自动提取", None, uploaded_path, embedded.get("sourceRow"), len(item["attachments"]) + embedded_index, embedded.get("linkUrl")),
                    )
            for order, module in enumerate(item["modules"]):
                module_id = uuid.uuid4().hex
                conn.execute(
                    "INSERT INTO modules VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)",
                    (module_id, submission_id, module["sectionType"], module["evaluationItem"], module["moduleName"], module["weight"], module["standard"], module["evidence"], module["selfScore"], module["certificationMethod"], order),
                )
                for ref in module["attachmentRefs"]:
                    if ref in attachment_ids:
                        conn.execute("INSERT OR IGNORE INTO module_attachments VALUES (?, ?, '自动识别')", (module_id, attachment_ids[ref]))
            created.append(submission_id)
    return created


def list_submissions() -> list[dict]:
    with db() as conn:
        rows = conn.execute(
            """
            SELECT s.*, COUNT(DISTINCT m.id) module_count, COUNT(DISTINCT a.id) attachment_count
            FROM submissions s
            LEFT JOIN modules m ON m.submission_id=s.id
            LEFT JOIN attachments a ON a.submission_id=s.id
            GROUP BY s.id ORDER BY s.imported_at DESC
            """
        ).fetchall()
    return [dict(row) for row in rows]


FRAMEWORK_SECTIONS = (
    ("basic", "基本条件"), ("knowledge", "知识标准"),
    ("skills", "技能标准"), ("competencies", "能力素质标准"),
)


def normalize_framework(value: dict) -> dict:
    if not isinstance(value, dict):
        raise ValueError("标准框架格式无效")
    result = {}
    for key, label in FRAMEWORK_SECTIONS:
        section = value.get(key, {})
        if not isinstance(section, dict):
            raise ValueError(f"{label}格式无效")
        result[key] = {"content": clean(section.get("content")), "method": clean(section.get("method"))}
    return result


def list_job_taxonomy() -> list[dict]:
    with db() as conn:
        return [dict(row) for row in conn.execute("SELECT * FROM job_taxonomy ORDER BY created_at,name,id")]


def taxonomy_path(conn: sqlite3.Connection, subcategory_id: str) -> dict:
    row = conn.execute(
        """SELECT sub.id subcategory_id,sub.name subcategory_name,c.id category_id,c.name category_name,
                  f.id family_id,f.name family_name
           FROM job_taxonomy sub JOIN job_taxonomy c ON c.id=sub.parent_id
           JOIN job_taxonomy f ON f.id=c.parent_id
           WHERE sub.id=? AND sub.kind='subcategory' AND c.kind='category' AND f.kind='family'""",
        (subcategory_id,),
    ).fetchone()
    if not row:
        raise ValueError("请选择有效的职族、职类和子类")
    return dict(row)


def save_job_taxonomy(payload: dict, node_id: str | None = None) -> dict:
    name = clean(payload.get("name"))
    if not name or len(name) > 100:
        raise ValueError("分类名称须为 1 至 100 个字符")
    with db() as conn:
        conn.execute("BEGIN IMMEDIATE")
        if node_id:
            existing = conn.execute("SELECT * FROM job_taxonomy WHERE id=?", (node_id,)).fetchone()
            if not existing:
                raise ValueError("分类不存在")
            kind, parent_id = existing["kind"], existing["parent_id"]
        else:
            kind, parent_id = clean(payload.get("kind")), clean(payload.get("parentId")) or None
            if kind not in {"family", "category", "subcategory"}:
                raise ValueError("分类层级无效")
            if kind == "family" and parent_id:
                raise ValueError("职族不能设置上级分类")
            if kind != "family":
                parent = conn.execute("SELECT kind FROM job_taxonomy WHERE id=?", (parent_id,)).fetchone()
                if not parent or parent["kind"] != {"category": "family", "subcategory": "category"}[kind]:
                    raise ValueError("职类必须属于职族，子类必须属于职类")
        duplicate = conn.execute("SELECT id FROM job_taxonomy WHERE parent_id IS ? AND name=?", (parent_id, name)).fetchone()
        if duplicate and duplicate["id"] != node_id:
            raise ValueError("同一上级下已存在该名称")
        if node_id:
            conn.execute("UPDATE job_taxonomy SET name=? WHERE id=?", (name, node_id))
        else:
            node_id = uuid.uuid4().hex
            conn.execute("INSERT INTO job_taxonomy(id,name,kind,parent_id,created_at) VALUES(?,?,?,?,?)", (node_id, name, kind, parent_id, now_text()))
        return dict(conn.execute("SELECT * FROM job_taxonomy WHERE id=?", (node_id,)).fetchone())


def delete_job_taxonomy(node_id: str) -> None:
    with db() as conn:
        conn.execute("BEGIN IMMEDIATE")
        if conn.execute("SELECT 1 FROM job_taxonomy WHERE parent_id=?", (node_id,)).fetchone():
            raise ValueError("该分类下仍有下级分类，请先处理下级分类")
        if conn.execute("SELECT 1 FROM qualification_standards WHERE subcategory_id=?", (node_id,)).fetchone():
            raise ValueError("该子类下仍有资格标准，不能直接删除；请先删除未定稿且未被引用的标准，已存档或被引用的标准请使用「停用」保留历史")
        if not conn.execute("DELETE FROM job_taxonomy WHERE id=?", (node_id,)).rowcount:
            raise ValueError("分类不存在")


def set_job_taxonomy_active(node_id: str, active: bool) -> dict:
    """停用分类以保留历史；停用后不再出现在新建标准的选择中，但已有引用不受影响。"""
    if not isinstance(active, bool):
        raise ValueError("启用状态无效")
    with db() as conn:
        conn.execute("BEGIN IMMEDIATE")
        if not conn.execute("UPDATE job_taxonomy SET active=? WHERE id=?", (1 if active else 0, node_id)).rowcount:
            raise ValueError("分类不存在")
    return {"id": node_id, "active": active}


def enrich_standard(conn: sqlite3.Connection, row) -> dict:
    result = dict(row)
    result["matrix"] = json.loads(result.pop("matrix_json", None) or "null")
    result["positions"] = json.loads(result.pop("positions_json", "{}"))
    result["framework"] = normalize_framework(json.loads(result.pop("framework_json", "{}")))
    classification = json.loads(result.pop("classification_json", "{}"))
    if result.get("status") == "已存档" and classification:
        result.update(classification)
        result["category"] = result["subcategory_name"]
    elif result.get("subcategory_id"):
        result.update(taxonomy_path(conn, result["subcategory_id"]))
        result["category"] = result["subcategory_name"]
    return result


def list_qualification_standards(published_only: bool = False) -> list[dict]:
    where = "WHERE q.status='已存档'" if published_only else ""
    with db() as conn:
        rows = conn.execute(
            f"""SELECT q.*, COUNT(sm.id) module_count,
                       SUM(CASE WHEN sm.section_type='standard' THEN COALESCE(sm.weight,0) ELSE 0 END) weight_total,
                       (SELECT COUNT(*) FROM submissions s WHERE s.standard_id=q.id) application_count
                FROM qualification_standards q
                LEFT JOIN standard_modules sm ON sm.standard_id=q.id
                {where}
                GROUP BY q.id ORDER BY q.updated_at DESC"""
        ).fetchall()
        return [enrich_standard(conn, row) for row in rows]


def qualification_standard_detail(standard_id: str, published_only: bool = False) -> dict | None:
    with db() as conn:
        sql = "SELECT * FROM qualification_standards WHERE id=?" + (" AND status='已存档'" if published_only else "")
        standard = conn.execute(sql, (standard_id,)).fetchone()
        if not standard:
            return None
        modules = [dict(row) for row in conn.execute(
            "SELECT * FROM standard_modules WHERE standard_id=? ORDER BY sort_order", (standard_id,)
        )]
        return {"standard": enrich_standard(conn, standard), "modules": modules}


_CHINESE_DIGITS = "零一二三四五六七八九"


def _chinese_numeral(number: int) -> str:
    if number < 1:
        return str(number)
    if number <= 10:
        return _CHINESE_DIGITS[number]
    if number < 20:
        return "十" + (_CHINESE_DIGITS[number - 10] if number > 10 else "")
    tens, ones = divmod(number, 10)
    return _CHINESE_DIGITS[tens] + "十" + (_CHINESE_DIGITS[ones] if ones else "")


def export_subclass_standard_matrix(subcategory_id: str) -> tuple[bytes, str]:
    """把某个子类最新存档标准按维度分块导出为单个 Excel Sheet。"""
    with db() as conn:
        path = taxonomy_path(conn, subcategory_id)
        row = conn.execute(
            """SELECT * FROM qualification_standards
               WHERE subcategory_id=? AND status='已存档'
               ORDER BY COALESCE(archived_at,'') DESC, version_number DESC LIMIT 1""",
            (subcategory_id,),
        ).fetchone()
        if not row:
            raise ValueError("该子类下暂无已存档的任职资格标准")
        standard = enrich_standard(conn, row)

    matrix = standard.get("matrix") or {}
    dimensions = matrix.get("dimensions") or []
    tier_columns = [tier["name"] for tier in TIERS[:4]]
    grade_columns = list(GRADE_NAMES)
    grade_groups = [
        ("初做者", grade_columns[0:2]),
        ("专业级", grade_columns[2:4]),
        ("骨干级", grade_columns[4:6]),
        ("专家级", grade_columns[6:7]),
    ]

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "任职资格标准"

    sheet.cell(1, 1, " / ".join(filter(None, [
        standard.get("family_name"), standard.get("category_name"), standard.get("subcategory_name"),
        f"{standard.get('name') or ''} V{standard.get('version_number') or ''}",
    ])))
    sheet.cell(2, 1, f"存档：{standard.get('archived_by') or ''} · {standard.get('archived_at') or ''}")
    sheet.cell(3, 1, f"导出时间：{now_text()}；空白单元格表示该等级尚未编写。")

    title_font = Font(bold=True, size=14, color="18323B")
    subtitle_font = Font(color="64767C")
    section_font = Font(bold=True, size=12, color="18323B")
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="446A68")
    section_fill = PatternFill("solid", fgColor="DDE9E7")
    thin_border = Border(left=Side(style="thin", color="D6DEE5"), right=Side(style="thin", color="D6DEE5"),
                         top=Side(style="thin", color="D6DEE5"), bottom=Side(style="thin", color="D6DEE5"))

    sheet["A1"].font = title_font
    sheet["A2"].font = subtitle_font
    sheet["A3"].font = subtitle_font

    def style(cell, *, font=None, fill=None, alignment=None, border=None):
        if font is not None:
            cell.font = font
        if fill is not None:
            cell.fill = fill
        if alignment is not None:
            cell.alignment = alignment
        if border is not None:
            cell.border = border

    row = 5
    for index, dim in enumerate(dimensions, 1):
        is_tier = dim.get("basis") == "tier"
        columns = tier_columns if is_tier else grade_columns
        level_count = len(columns)
        total_cols = 1 + level_count

        sheet.cell(row, 1, f"{_chinese_numeral(index)}、{dim.get('name') or ''}")
        sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=total_cols)
        style(sheet.cell(row, 1), font=section_font, fill=section_fill, border=thin_border)
        row += 1

        if is_tier:
            sheet.cell(row, 1, "子维度")
            for offset, name in enumerate(columns):
                sheet.cell(row, 2 + offset, name)
            for col in range(1, total_cols + 1):
                style(sheet.cell(row, col), font=header_font, fill=header_fill,
                      alignment=Alignment(horizontal="center", vertical="center"), border=thin_border)
            row += 1
        else:
            group_row = row
            sheet.cell(group_row, 1, "子维度")
            for group_name, group_cols in grade_groups:
                start_col = 2 + grade_columns.index(group_cols[0])
                end_col = 2 + grade_columns.index(group_cols[-1])
                sheet.cell(group_row, start_col, group_name)
                if len(group_cols) > 1:
                    sheet.merge_cells(start_row=group_row, start_column=start_col, end_row=group_row, end_column=end_col)
            sheet.merge_cells(start_row=group_row, start_column=1, end_row=group_row + 1, end_column=1)
            style(sheet.cell(group_row, 1), font=header_font, fill=header_fill,
                  alignment=Alignment(horizontal="center", vertical="center"), border=thin_border)
            for group_name, group_cols in grade_groups:
                start_col = 2 + grade_columns.index(group_cols[0])
                style(sheet.cell(group_row, start_col), font=header_font, fill=header_fill,
                      alignment=Alignment(horizontal="center", vertical="center"), border=thin_border)
            row += 1
            for offset, name in enumerate(columns):
                style(sheet.cell(row, 2 + offset, name), font=header_font, fill=header_fill,
                      alignment=Alignment(horizontal="center", vertical="center"), border=thin_border)
            row += 1

        for sub in dim.get("rows", []) or [{"name": "", "cells": {}}]:
            sheet.cell(row, 1, sub.get("name", ""))
            cells = sub.get("cells") or {}
            for offset in range(level_count):
                sheet.cell(row, 2 + offset, cells.get(str(offset + 1), ""))
            for col in range(1, total_cols + 1):
                style(sheet.cell(row, col), alignment=Alignment(vertical="top", wrap_text=True), border=thin_border)
            row += 1

        row += 1

    sheet.column_dimensions["A"].width = 30
    for column in range(2, 9):
        sheet.column_dimensions[get_column_letter(column)].width = 34
    sheet.freeze_panes = "A5"

    output = io.BytesIO()
    workbook.save(output)
    safe_name = re.sub(r'[\\/:*?"<>|]+', "_", path["subcategory_name"]).strip() or "子类"
    return output.getvalue(), f"{safe_name}-任职资格标准.xlsx"


def get_qualification_template():
    with db() as conn:
        row = conn.execute("SELECT * FROM qualification_template WHERE id=1").fetchone()
        return {"matrix": json.loads(row["matrix_json"]), "revision": row["revision"], "tiers": TIERS, "gradeNames": GRADE_NAMES}


def save_qualification_template(payload):
    matrix = normalize_matrix(payload.get("matrix"), template=True)
    with db() as conn:
        conn.execute("BEGIN IMMEDIATE")
        row = conn.execute("SELECT revision FROM qualification_template WHERE id=1").fetchone()
        if payload.get("revision") != row["revision"]:
            raise ValueError("维度模板已被更新，请重新打开后修改")
        conn.execute("UPDATE qualification_template SET matrix_json=?,revision=revision+1 WHERE id=1", (json.dumps(matrix, ensure_ascii=False),))
    return get_qualification_template()


def get_job_positions(subcategory_id):
    with db() as conn:
        taxonomy_path(conn, subcategory_id)
        return {str(r["grade"]): r["name"] for r in conn.execute("SELECT grade,name FROM job_positions WHERE subcategory_id=? ORDER BY grade", (subcategory_id,))}


def save_job_positions(subcategory_id, payload):
    positions = payload.get("positions")
    if not isinstance(positions, dict) or any(k not in {str(i) for i in range(1,8)} for k in positions):
        raise ValueError("岗位名称必须按一级至七级配置")
    if any(not isinstance(v, str) or len(v) > 200 for v in positions.values()):
        raise ValueError("岗位名称必须为不超过 200 字的文本")
    normalized = {str(i): positions.get(str(i), "").strip() for i in range(1,8)}
    with db() as conn:
        conn.execute("BEGIN IMMEDIATE")
        taxonomy_path(conn, subcategory_id)
        current = {str(r["grade"]): r["name"] for r in conn.execute("SELECT grade,name FROM job_positions WHERE subcategory_id=?", (subcategory_id,))}
        if payload.get("original") != current:
            raise ValueError("岗位名称已被更新，请重新打开后修改")
        for grade, name in normalized.items():
            conn.execute("INSERT INTO job_positions VALUES(?,?,?) ON CONFLICT(subcategory_id,grade) DO UPDATE SET name=excluded.name", (subcategory_id,int(grade),name))
    return {"positions": normalized}


def save_matrix_standard(payload, actor, standard_id=None):
    name = clean(payload.get("name"))
    if not name:
        raise ValueError("请填写标准名称")
    archived = payload.get("status") == "已存档"
    if payload.get("status", "草稿") not in ("草稿", "迭代中", "已存档"):
        raise ValueError("标准状态无效")
    matrix = normalize_matrix(payload.get("matrix"), archive=archived)
    subcategory_id = clean(payload.get("subcategoryId")) or None
    if archived and not subcategory_id:
        raise ValueError("存档前请选择所属子类")
    with db() as conn:
        conn.execute("BEGIN IMMEDIATE")
        existing = conn.execute("SELECT * FROM qualification_standards WHERE id=?", (standard_id,)).fetchone() if standard_id else None
        if standard_id and not existing:
            raise ValueError("资格标准不存在")
        if existing and existing["status"] == "已存档":
            raise ValueError("已存档版本不可修改，请发起迭代")
        if existing and payload.get("revision") != existing["revision"]:
            raise ValueError("标准已被更新，请重新打开后编辑")
        path = taxonomy_path(conn, subcategory_id) if subcategory_id else {}
        positions = {str(r["grade"]):r["name"] for r in conn.execute("SELECT grade,name FROM job_positions WHERE subcategory_id=?", (subcategory_id,))}
        sid, timestamp = standard_id or uuid.uuid4().hex, now_text()
        status = "已存档" if archived else "迭代中" if existing and existing["parent_version_id"] else "草稿"
        if not existing:
            conn.execute("INSERT INTO qualification_standards(id,name,created_by,created_at,updated_at,series_id) VALUES(?,?,?,?,?,?)", (sid,name,actor.get("userId"),timestamp,timestamp,sid))
        conn.execute("""UPDATE qualification_standards SET name=?,subcategory_id=?,category=?,status=?,updated_at=?,matrix_json=?,positions_json=?,revision=revision+1,change_note=? WHERE id=?""",
                     (name,subcategory_id,path.get("subcategory_name", existing["category"] if existing else ""),status,timestamp,json.dumps(matrix,ensure_ascii=False),json.dumps(positions,ensure_ascii=False),clean(payload.get("changeNote")),sid))
        if archived:
            conn.execute("UPDATE qualification_standards SET archived_at=?,archived_by=?,classification_json=? WHERE id=?", (timestamp,actor.get("name") or actor.get("userId"),json.dumps(path,ensure_ascii=False),sid))
        record_standard_history(conn,sid,"定稿存档" if archived else "保存分级矩阵",actor,clean(payload.get("changeNote")))
    return qualification_standard_detail(sid)


def record_standard_history(conn, standard_id: str, action: str, actor: dict, note: str = ""):
    row = conn.execute("SELECT * FROM qualification_standards WHERE id=?", (standard_id,)).fetchone()
    snapshot = {"standard": enrich_standard(conn, row), "modules": [dict(m) for m in conn.execute(
        "SELECT * FROM standard_modules WHERE standard_id=? ORDER BY sort_order", (standard_id,))]}
    conn.execute("INSERT INTO standard_history(standard_id,series_id,action,actor_id,actor_name,created_at,note,snapshot_json) VALUES(?,?,?,?,?,?,?,?)",
                 (standard_id, row["series_id"], action, actor.get("userId"), actor.get("name"), now_text(), note, json.dumps(snapshot, ensure_ascii=False)))


def qualification_standard_history(standard_id: str) -> list[dict]:
    with db() as conn:
        row = conn.execute("SELECT series_id FROM qualification_standards WHERE id=?", (standard_id,)).fetchone()
        if not row:
            raise ValueError("资格标准不存在")
        items = []
        for record in conn.execute("SELECT * FROM standard_history WHERE series_id=? ORDER BY id DESC", (row["series_id"],)):
            item = dict(record)
            item["snapshot"] = json.loads(item.pop("snapshot_json"))
            items.append(item)
        return items


def iterate_qualification_standard(standard_id: str, actor: dict, note: str) -> dict:
    note = clean(note)
    if not note:
        raise ValueError("请填写迭代说明")
    with db() as conn:
        conn.execute("BEGIN IMMEDIATE")
        source = conn.execute("SELECT * FROM qualification_standards WHERE id=?", (standard_id,)).fetchone()
        if not source or source["status"] != "已存档":
            raise ValueError("只有已存档标准可以发起迭代")
        if conn.execute("SELECT 1 FROM qualification_standards WHERE series_id=? AND status!='已存档'", (source["series_id"],)).fetchone():
            raise ValueError("该标准已有开发中的版本，请在标准开发中继续编辑")
        number = conn.execute("SELECT MAX(version_number)+1 FROM qualification_standards WHERE series_id=?", (source["series_id"],)).fetchone()[0]
        new_id, timestamp = uuid.uuid4().hex, now_text()
        conn.execute("""INSERT INTO qualification_standards(id,name,level,category,description,status,created_by,created_at,updated_at,subcategory_id,framework_json,series_id,version_number,parent_version_id,change_note,revision)
            VALUES(?,?,?,?,?,'迭代中',?,?,?,?,?,?,?,?,?,1)""",
            (new_id, source["name"], source["level"], source["category"], source["description"], actor.get("userId"), timestamp, timestamp, source["subcategory_id"], source["framework_json"], source["series_id"], number, standard_id, note))
        for m in conn.execute("SELECT * FROM standard_modules WHERE standard_id=? ORDER BY sort_order", (standard_id,)).fetchall():
            conn.execute("INSERT INTO standard_modules VALUES(?,?,?,?,?,?,?,?,?)", (uuid.uuid4().hex,new_id,m["section_type"],m["evaluation_item"],m["module_name"],m["weight"],m["standard"],m["certification_method"],m["sort_order"]))
        conn.execute("UPDATE qualification_standards SET matrix_json=?,positions_json=? WHERE id=?", (source["matrix_json"],source["positions_json"],new_id))
        record_standard_history(conn, new_id, "发起迭代", actor, note)
    return qualification_standard_detail(new_id)


def save_qualification_standard(payload: dict, actor: dict, standard_id: str | None = None) -> dict:
    if "matrix" in payload:
        return save_matrix_standard(payload, actor, standard_id)
    if standard_id:
        current = qualification_standard_detail(standard_id)
        if current and current["standard"].get("matrix"):
            raise ValueError("分级标准必须通过矩阵编辑保存")
    name = clean(payload.get("name"))
    level = clean(payload.get("level"))
    category = clean(payload.get("category"))
    description = clean(payload.get("description"))
    status = clean(payload.get("status")) or "草稿"
    if not name:
        raise ValueError("请填写标准名称")
    if status not in {"草稿", "迭代中", "已存档"}:
        raise ValueError("标准状态无效")
    modules = payload.get("modules") or []
    normalized = []
    for index, item in enumerate(modules):
        section_type = clean(item.get("sectionType")) or "standard"
        if section_type not in {"standard", "bonus"}:
            raise ValueError(f"第 {index + 1} 项的类型无效")
        module_name = clean(item.get("moduleName"))
        if not module_name:
            raise ValueError(f"第 {index + 1} 项请填写工作模块")
        weight = None
        if section_type == "standard":
            raw_weight = item.get("weight")
            if raw_weight not in (None, "", 0, 0.0):
                try:
                    weight = float(raw_weight)
                except (TypeError, ValueError):
                    raise ValueError(f"第 {index + 1} 项请填写有效权重")
                if not 0 < weight <= 1:
                    raise ValueError(f"第 {index + 1} 项权重必须大于 0 且不超过 100%")
        normalized.append({
            "sectionType": section_type,
            "evaluationItem": clean(item.get("evaluationItem")),
            "moduleName": module_name,
            "weight": weight,
            "standard": clean(item.get("standard")),
            "certificationMethod": clean(item.get("certificationMethod")),
        })
    if status == "已存档":
        normal = [item for item in normalized if item["sectionType"] == "standard"]
        if not normal:
            raise ValueError("存档前至少需要一个行为标准项")
        incomplete = [item["moduleName"] for item in normal if item["weight"] is None or not item["standard"]]
        if incomplete:
            raise ValueError("存档前请补全权重和职级定义标准：" + "、".join(incomplete[:5]))
        if abs(sum(item["weight"] for item in normal) - 1) > .001:
            raise ValueError("存档前，行为标准权重合计必须为 100%")
    updating = bool(standard_id)
    standard_id = standard_id or uuid.uuid4().hex
    timestamp = now_text()
    with db() as conn:
        conn.execute("BEGIN IMMEDIATE")
        existing = conn.execute("SELECT * FROM qualification_standards WHERE id=?", (standard_id,)).fetchone()
        if updating and not existing:
            raise ValueError("资格标准不存在")
        if existing and existing["status"] == "已存档":
            raise ValueError("已存档版本不可修改，请发起迭代创建新版本")
        if existing and "revision" in payload and payload["revision"] != existing["revision"]:
            raise ValueError("标准已被更新，请重新打开后编辑")
        if status != "已存档":
            status = "迭代中" if existing and existing["parent_version_id"] else "草稿"
        if existing and "category" not in payload:
            category = existing["category"]
        subcategory_id = clean(payload.get("subcategoryId", existing["subcategory_id"] if existing else "")) or None
        framework = normalize_framework(payload.get("framework", json.loads(existing["framework_json"]) if existing else {}))
        if subcategory_id:
            path = taxonomy_path(conn, subcategory_id)
            category = path["subcategory_name"]
        if status == "已存档" and not subcategory_id:
            raise ValueError("存档前请选择所属子类（职族 → 职类 → 子类）")
        if existing:
            conn.execute(
                "UPDATE qualification_standards SET name=?,level=?,category=?,description=?,status=?,updated_at=?,subcategory_id=?,framework_json=? WHERE id=?",
                (name, level, category, description, status, timestamp, subcategory_id, json.dumps(framework, ensure_ascii=False), standard_id),
            )
            conn.execute("DELETE FROM standard_modules WHERE standard_id=?", (standard_id,))
        else:
            conn.execute(
                "INSERT INTO qualification_standards(id,name,level,category,description,status,created_by,created_at,updated_at,subcategory_id,framework_json) VALUES(?,?,?,?,?,?,?,?,?,?,?)",
                (standard_id, name, level, category, description, status, actor.get("userId"), timestamp, timestamp, subcategory_id, json.dumps(framework, ensure_ascii=False)),
            )
        for index, item in enumerate(normalized):
            conn.execute(
                "INSERT INTO standard_modules VALUES(?,?,?,?,?,?,?,?,?)",
                (uuid.uuid4().hex, standard_id, item["sectionType"], item["evaluationItem"], item["moduleName"], item["weight"], item["standard"], item["certificationMethod"], index),
            )
        conn.execute("UPDATE qualification_standards SET series_id=COALESCE(series_id,id),revision=revision+1,change_note=? WHERE id=?",
                     (clean(payload.get("changeNote")) or (existing["change_note"] if existing else ""), standard_id))
        if status == "已存档":
            conn.execute("UPDATE qualification_standards SET archived_at=?,archived_by=?,classification_json=? WHERE id=?",
                         (timestamp, actor.get("name") or actor.get("userId"), json.dumps(path, ensure_ascii=False), standard_id))
        record_standard_history(conn, standard_id, "定稿存档" if status == "已存档" else "保存迭代" if existing and existing["parent_version_id"] else "保存草稿", actor, clean(payload.get("changeNote")))
    return qualification_standard_detail(standard_id)


def copy_submission_to_standard(submission_id: str, actor: dict) -> dict:
    detail = submission_detail(submission_id)
    if not detail:
        raise ValueError("认证材料不存在")
    submission = detail["submission"]
    snapshot = json.loads(submission.get("standard_snapshot_json") or "{}")
    if snapshot.get("matrix"):
        source = qualification_standard_detail(submission.get("standard_id"))
        return save_matrix_standard({"name":(snapshot.get("name") or "分级标准")+"（副本）", "status":"草稿", "matrix":snapshot["matrix"],
            "subcategoryId":source["standard"]["subcategory_id"] if source else "", "changeNote":"从举证材料引用的矩阵快照复制"},actor)
    return save_qualification_standard({
        "name": f"{submission.get('category') or '通用'} {submission.get('level') or ''}任职资格标准".strip(),
        "level": submission.get("level"),
        "category": submission.get("category"),
        "framework": json.loads(submission.get("standard_snapshot_json") or "{}").get("framework", {}),
        "description": f"从 {submission.get('name')} 的材料结构复制，请确认后定稿存档。",
        "status": "草稿",
        "modules": [{
            "sectionType": item["section_type"], "evaluationItem": item["evaluation_item"],
            "moduleName": item["module_name"], "weight": item["weight"],
            "standard": item["standard"], "certificationMethod": item["certification_method"],
        } for item in detail["modules"]],
    }, actor)


def delete_qualification_standard(standard_id: str) -> None:
    """只允许删除未定稿且未被引用的标准；已存档版本永久保留。"""
    with db() as conn:
        conn.execute("BEGIN IMMEDIATE")
        row = conn.execute("SELECT status FROM qualification_standards WHERE id=?", (standard_id,)).fetchone()
        if not row:
            raise ValueError("资格标准不存在")
        if row["status"] == "已存档":
            raise ValueError("已存档标准为保留版本记录不可删除；如需调整请发起迭代")
        if conn.execute("SELECT 1 FROM submissions WHERE standard_id=?", (standard_id,)).fetchone():
            raise ValueError("该标准已被材料引用，不能删除")
        conn.execute("DELETE FROM standard_history WHERE standard_id=?", (standard_id,))
        conn.execute("DELETE FROM qualification_standards WHERE id=?", (standard_id,))


def list_my_applications(actor: dict) -> list[dict]:
    with db() as conn:
        rows = conn.execute(
            """SELECT s.*, COUNT(DISTINCT m.id) module_count, COUNT(DISTINCT a.id) attachment_count
               FROM submissions s LEFT JOIN modules m ON m.submission_id=s.id
               LEFT JOIN attachments a ON a.submission_id=s.id
               WHERE s.owner_external_user_id=? GROUP BY s.id ORDER BY s.imported_at DESC""",
            (actor["userId"],),
        ).fetchall()
    return [dict(row) for row in rows]


def create_online_application(standard_id: str, actor: dict, grade=None, task_config=None) -> dict:
    if actor.get('evidenceId'):
        raise PermissionError('举证人只能填写管理员配置的任务')
    detail = qualification_standard_detail(standard_id, published_only=True)
    if not detail:
        raise ValueError("资格标准不存在或尚未存档，只有已存档标准可以引用")
    application_id = uuid.uuid4().hex
    standard = detail["standard"]
    matrix = standard.get("matrix")
    if matrix:
        grade = grade_number(grade)
        standard["level"] = GRADE_NAMES[grade-1]
        resolved = resolve_requirements(matrix, grade)
        detail["modules"] = [{"section_type":dim["reviewType"], "evaluation_item":dim["name"], "module_name":row["name"], "weight":None,
                              "standard":row["text"] or ("（该等级标准待补充）" if row["pending"] else "无新增要求"), "certification_method":""}
                             for dim in resolved for row in dim["rows"]]
    timestamp = now_text()
    with db() as conn:
        conn.execute('BEGIN IMMEDIATE')
        duplicate = conn.execute(
            "SELECT id FROM submissions WHERE owner_external_user_id=? AND standard_id=? AND status='填报中' AND COALESCE(level,'')=?",
            (actor["userId"], standard_id, standard.get("level") or ""),
        ).fetchone()
        if duplicate:
            return submission_detail(duplicate["id"])
        conn.execute(
            """INSERT INTO submissions(id,name,level,category,applied_at,source_file,status,imported_at,raw_path,
                       owner_external_user_id,entry_mode,standard_id,submitted_at)
               VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?)""",
            (application_id, actor["name"], standard["level"], standard["category"], timestamp[:10], "在线填报", "填报中", timestamp, "", actor["userId"], "online", standard_id, None),
        )
        snapshot = {key: standard.get(key) for key in ("id", "series_id", "version_number", "archived_at", "name", "level", "family_name", "category_name", "subcategory_name", "framework")}
        if matrix:
            snapshot.update({"matrix":matrix,"resolved":resolved,"grade":grade,"positions":standard["positions"],"positionName":standard["positions"].get(str(grade),"")})
        conn.execute("UPDATE submissions SET standard_snapshot_json=? WHERE id=?", (json.dumps(snapshot, ensure_ascii=False), application_id))
        for index, item in enumerate(detail["modules"]):
            conn.execute(
                "INSERT INTO modules VALUES(?,?,?,?,?,?,?,?,?,?,?)",
                (uuid.uuid4().hex, application_id, item["section_type"], item["evaluation_item"], item["module_name"], item["weight"], item["standard"], "", "", item["certification_method"], index),
            )
        if task_config:
            provision_evidence_personnel(conn,application_id,actor['name'],task_config)
            evidence.insert_task(conn,application_id,task_config)
            evidence.event(conn,application_id,'配置举证任务',task_config['instructions'],{'name':task_config.get('createdByName') or '管理员'})
    return submission_detail(application_id)


def require_application_access(application_id: str, actor: dict, editable: bool = False) -> dict:
    with db() as conn:
        row = conn.execute("SELECT * FROM submissions WHERE id=? AND entry_mode='online'", (application_id,)).fetchone()
    if not row:
        raise ValueError("在线申请不存在")
    if actor.get("role") != "管理员" and row["owner_external_user_id"] != actor.get("userId"):
        raise PermissionError("你无权访问该申请")
    if editable and row["status"] not in evidence.EDITABLE:
        raise ValueError("该申请已提交，不能再修改")
    return dict(row)


def save_online_application(application_id: str, payload: dict, actor: dict) -> dict:
    require_application_access(application_id, actor, editable=True)
    name = clean(payload.get("name")) or actor.get("name")
    if not name:
        raise ValueError("请填写姓名")
    submit = bool(payload.get("submit"))
    with db() as conn:
        conn.execute('BEGIN IMMEDIATE')
        task=evidence.task(conn,application_id)
        current=conn.execute('SELECT status,name FROM submissions WHERE id=?',(application_id,)).fetchone()
        if current['status'] not in evidence.EDITABLE:raise ValueError('材料状态已改变，请重新打开')
        if task:
            if payload.get('revision')!=task['revision']:raise ValueError('材料或反馈已更新，请重新打开后填写')
            name=current['name']
        valid_modules = {row["id"] for row in conn.execute("SELECT id FROM modules WHERE submission_id=?", (application_id,))}
        for item in payload.get("modules") or []:
            module_id = clean(item.get("moduleId"))
            if task and (len(clean(item.get('evidence')))>30000 or len(clean(item.get('selfScore')))>1000):raise ValueError('单项举证最多30000字，自评最多1000字')
            if module_id in valid_modules:
                conn.execute("UPDATE modules SET evidence=?,self_score=? WHERE id=?", (clean(item.get("evidence")), clean(item.get("selfScore")), module_id))
        if submit:
            validate_application_attachment_links(conn, application_id)
            missing = conn.execute(
                "SELECT module_name FROM modules WHERE submission_id=? AND section_type IN ('standard','reference') AND evaluation_item<>? AND TRIM(COALESCE(evidence,''))='' ORDER BY sort_order",
                (application_id, BASIC_CONDITION_NAME),
            ).fetchall()
            if missing:
                raise ValueError("请先完成所有必填维度的举证：" + "、".join(row["module_name"] for row in missing[:5]))
        timestamp = now_text()
        conn.execute(
            "UPDATE submissions SET name=?,applied_at=?,status=?,submitted_at=? WHERE id=?",
            (name, clean(payload.get("appliedAt")) or timestamp[:10], ("待管理员确认" if task else "待评审") if submit else current['status'], timestamp if submit else None, application_id),
        )
        if task:evidence.event(conn,application_id,'提交举证' if submit else '保存草稿','',actor)
    return submission_detail(application_id)


def application_attachment_module(conn, application_id, module_id):
    if not module_id:
        raise ValueError('请为每份附件选择一个对应的举证项')
    module = conn.execute('SELECT id,module_name FROM modules WHERE id=? AND submission_id=?', (module_id, application_id)).fetchone()
    if not module:
        raise ValueError('关联的举证项不存在')
    return module


def validate_application_attachment_links(conn, application_id):
    invalid = conn.execute('''SELECT a.material_name FROM attachments a
        LEFT JOIN module_attachments ma ON ma.attachment_id=a.id
        LEFT JOIN modules m ON m.id=ma.module_id AND m.submission_id=a.submission_id
        WHERE a.submission_id=? GROUP BY a.id
        HAVING COUNT(ma.module_id)<>1 OR COUNT(m.id)<>1''', (application_id,)).fetchall()
    if invalid:
        raise ValueError('每份附件必须且只能对应一个举证项，请调整附件归属：' + '、'.join(r['material_name'] or '未命名附件' for r in invalid[:5]))


def reassign_application_attachment(application_id, attachment_id, module_id, actor):
    require_application_access(application_id, actor, editable=True)
    with db() as conn:
        conn.execute('BEGIN IMMEDIATE')
        if conn.execute('SELECT status FROM submissions WHERE id=?', (application_id,)).fetchone()['status'] not in evidence.EDITABLE:
            raise ValueError('材料已提交，不能修改附件归属')
        module = application_attachment_module(conn, application_id, module_id)
        if not conn.execute('SELECT id FROM attachments WHERE id=? AND submission_id=?', (attachment_id, application_id)).fetchone():
            raise ValueError('附件不存在')
        conn.execute('DELETE FROM module_attachments WHERE attachment_id=?', (attachment_id,))
        conn.execute('INSERT INTO module_attachments VALUES(?,?,?)', (module_id, attachment_id, '在线填报'))
        conn.execute('UPDATE attachments SET project_name=? WHERE id=?', (module['module_name'], attachment_id))
        if evidence.task(conn, application_id):
            evidence.event(conn, application_id, '调整附件归属', '', actor)


def store_application_attachment(application_id: str, filename: str, source, length: int, actor: dict, module_id: str = "", material_name: str = "", link_url: str = "") -> dict:
    if link_url and link_url.strip().lower().split(":", 1)[0] in {"javascript", "data", "vbscript"}:
        raise ValueError('访问地址不能使用脚本或数据协议')
    require_application_access(application_id, actor, editable=True)
    filename = Path(filename).name
    if not filename:
        raise ValueError("附件文件名不能为空")
    with db() as conn:
        task=evidence.task(conn,application_id)
        if task:evidence.validate_attachment(filename,length)
        module = application_attachment_module(conn, application_id, module_id)
        count = conn.execute("SELECT COUNT(*) FROM attachments WHERE submission_id=?", (application_id,)).fetchone()[0]
    target_dir = SUBMISSION_DIR / application_id / "attachments"
    target_dir.mkdir(parents=True, exist_ok=True)
    target = target_dir / f"{uuid.uuid4().hex[:8]}-{filename}"
    remaining = length
    with target.open("wb") as output:
        while remaining:
            chunk = source.read(min(UPLOAD_CHUNK, remaining))
            if not chunk:
                target.unlink(missing_ok=True)
                raise ValueError("附件上传连接提前中断")
            output.write(chunk)
            remaining -= len(chunk)
    attachment_id = uuid.uuid4().hex
    project_name = module["module_name"] if module else "其他举证材料"
    with db() as conn:
        conn.execute('BEGIN IMMEDIATE')
        if conn.execute('SELECT status FROM submissions WHERE id=?',(application_id,)).fetchone()['status'] not in evidence.EDITABLE:
            target.unlink(missing_ok=True)
            raise ValueError('材料已提交，附件未保存')
        conn.execute(
            """INSERT INTO attachments(id,submission_id,serial,project_name,material_name,storage_path,image_path,uploaded_file_path,source_row,sort_order,link_url)
               VALUES(?,?,?,?,?,?,?,?,?,?,?)""",
            (attachment_id, application_id, str(count + 1), project_name, material_name or filename, filename, None, str(target.relative_to(DATA_DIR)), None, count, clean(link_url)),
        )
        if module:
            conn.execute("INSERT INTO module_attachments VALUES(?,?,?)", (module["id"], attachment_id, "在线填报"))
        if task:evidence.event(conn,application_id,'上传附件',filename,actor)
    return {"id": attachment_id, "name": filename, "size": length}


def store_application_link(application_id: str, payload: dict, actor: dict) -> dict:
    require_application_access(application_id, actor, editable=True)
    module_id = clean(payload.get("moduleId"))
    material_name = clean(payload.get("materialName"))
    link_url = clean(payload.get("linkUrl"))
    storage_path = clean(payload.get("storagePath")) or link_url
    if not material_name:
        raise ValueError("请填写材料名称")
    if not link_url and not storage_path:
        raise ValueError("请填写材料存放说明或访问地址")
    if link_url and link_url.strip().lower().split(":", 1)[0] in {"javascript", "data", "vbscript"}:
        raise ValueError("访问地址不能使用脚本或数据协议")
    with db() as conn:
        conn.execute('BEGIN IMMEDIATE')
        if conn.execute('SELECT status FROM submissions WHERE id=?',(application_id,)).fetchone()['status'] not in evidence.EDITABLE:raise ValueError('材料已提交，不能修改附件')
        module = application_attachment_module(conn, application_id, module_id)
        count = conn.execute("SELECT COUNT(*) FROM attachments WHERE submission_id=?", (application_id,)).fetchone()[0]
        attachment_id = uuid.uuid4().hex
        conn.execute(
            """INSERT INTO attachments(id,submission_id,serial,project_name,material_name,storage_path,image_path,uploaded_file_path,source_row,sort_order,link_url)
               VALUES(?,?,?,?,?,?,?,?,?,?,?)""",
            (attachment_id, application_id, str(count + 1), module["module_name"] if module else "其他举证材料", material_name, storage_path, None, None, None, count, link_url),
        )
        if module:
            conn.execute("INSERT INTO module_attachments VALUES(?,?,?)", (module["id"], attachment_id, "在线填报"))
        if evidence.task(conn,application_id):evidence.event(conn,application_id,'添加附件链接',material_name,actor)
    return {"id": attachment_id, "name": material_name, "size": 0}


def delete_application_attachment(application_id: str, attachment_id: str, actor: dict) -> None:
    require_application_access(application_id, actor, editable=True)
    with db() as conn:
        conn.execute('BEGIN IMMEDIATE')
        if conn.execute('SELECT status FROM submissions WHERE id=?',(application_id,)).fetchone()['status'] not in evidence.EDITABLE:raise ValueError('材料已提交，不能删除附件')
        assigned=evidence.task(conn,application_id)
        row = conn.execute("SELECT uploaded_file_path FROM attachments WHERE id=? AND submission_id=?", (attachment_id, application_id)).fetchone()
        if not row:
            raise ValueError("附件不存在")
        conn.execute("DELETE FROM attachments WHERE id=?", (attachment_id,))
        if assigned:evidence.event(conn,application_id,'移除附件','已从当前材料中移除，历史文件保留',actor)
    if row["uploaded_file_path"] and not assigned:
        target = (DATA_DIR / row["uploaded_file_path"]).resolve()
        if SUBMISSION_DIR.resolve() in target.parents and target.is_file():
            target.unlink()


def calculate_review_summary(conn: sqlite3.Connection, review_id: str) -> dict:
    rows = conn.execute(
        """SELECT m.id, m.module_name, m.section_type, m.weight, ms.score
           FROM modules m
           JOIN reviews r ON r.submission_id=m.submission_id
           LEFT JOIN module_scores ms ON ms.module_id=m.id AND ms.review_id=r.id
           WHERE r.id=? ORDER BY m.sort_order""",
        (review_id,),
    ).fetchall()
    behavior_total = 0.0
    bonus_total = 0.0
    items = []
    for row in rows:
        raw_score = float(row["score"] or 0)
        if row["section_type"] == "bonus":
            converted = raw_score
            bonus_total += converted
            formula = f"{raw_score:g} 分直接计入"
        else:
            weight = float(row["weight"] or 0)
            converted = raw_score / 5 * 100 * weight
            behavior_total += converted
            formula = f"({raw_score:g} ÷ 5) × 100 × {weight:.0%}"
        items.append({"moduleId": row["id"], "name": row["module_name"], "sectionType": row["section_type"], "rawScore": raw_score, "convertedScore": round(converted, 2), "formula": formula})
    return {"behaviorTotal": round(behavior_total, 2), "bonusTotal": round(bonus_total, 2), "total": round(behavior_total + bonus_total, 2), "items": items}


def dingtalk_status() -> dict:
    config = dingtalk_config()
    required = ["corpId", "clientId", "clientSecret", "agentId"]
    missing = [key for key in required if not config[key]] if config["mode"] == "live" else []
    return {
        "mode": config["mode"],
        "ready": config["mode"] == "mock" or (config["mode"] == "live" and not missing),
        "missing": missing,
        "baseUrl": config["baseUrl"],
        "capabilities": {"directory": True, "notifications": True, "login": config["mode"] == "live" and not missing},
    }


def get_app_user(external_user_id: str) -> dict | None:
    with db() as conn:
        row = conn.execute("SELECT external_user_id userId,name,role FROM app_users WHERE external_user_id=? AND active=1", (external_user_id,)).fetchone()
    return dict(row) if row else None


def mock_identities() -> list[dict]:
    with db() as conn:
        rows = conn.execute("SELECT external_user_id userId,name,role FROM app_users WHERE active=1 ORDER BY CASE role WHEN '管理员' THEN 1 WHEN '评委' THEN 2 ELSE 3 END,name").fetchall()
    return [dict(row) for row in rows]


def my_review_tasks(external_user_id: str) -> list[dict]:
    with db() as conn:
        rows = conn.execute(
            """SELECT a.id,a.token,a.status,a.pushed_at,a.completed_at,s.id submission_id,s.name submission_name,
                      s.level,s.category
               FROM review_assignments a JOIN reviewers r ON r.id=a.reviewer_id
               JOIN submissions s ON s.id=a.submission_id
               WHERE r.external_user_id=?
                 AND NOT EXISTS (SELECT 1 FROM evidence_tasks e
                                 WHERE e.submission_id=s.id AND e.finalized_at IS NULL)
               ORDER BY CASE WHEN a.status IN ('已评价','已完成') THEN 1 ELSE 0 END,a.created_at DESC""",
            (external_user_id,),
        ).fetchall()
    return [dict(row) for row in rows]


def list_directory_users(query: str = "") -> list[dict]:
    config = dingtalk_config()
    with db() as conn:
        if config["mode"] == "mock":
            existing_reviewers = conn.execute("SELECT id,name,external_user_id FROM reviewers ORDER BY name").fetchall()
            for reviewer in existing_reviewers:
                external_id = reviewer["external_user_id"] or f"mock-{reviewer['id']}"
                conn.execute("UPDATE reviewers SET external_user_id=?,directory_provider='dingtalk-mock' WHERE id=?", (external_id, reviewer["id"]))
                conn.execute(
                    """INSERT INTO directory_users(id,provider,external_user_id,name,department,mobile_masked,active,synced_at)
                       VALUES(?,?,?,?,?,?,1,?) ON CONFLICT(external_user_id) DO UPDATE SET name=excluded.name,active=1,synced_at=excluded.synced_at""",
                    (uuid.uuid4().hex, "dingtalk-mock", external_id, reviewer["name"], "模拟通讯录", "", now_text()),
                )
        sql = "SELECT external_user_id userId,name,department,provider FROM directory_users WHERE active=1"
        params = []
        if query.strip():
            sql += " AND (name LIKE ? OR department LIKE ?)"
            term = f"%{query.strip()}%"
            params.extend([term, term])
        sql += " ORDER BY name LIMIT 100"
        return [dict(row) for row in conn.execute(sql, params)]


def send_assignment_notification(assignment_id: str) -> dict:
    config = dingtalk_config()
    with db() as conn:
        row = conn.execute(
            """SELECT a.token,s.name submission_name,r.name reviewer_name,r.external_user_id
               FROM review_assignments a JOIN submissions s ON s.id=a.submission_id
               JOIN reviewers r ON r.id=a.reviewer_id WHERE a.id=?""", (assignment_id,)
        ).fetchone()
        if not row:
            raise ValueError("评审分配不存在")
        target_url = f"{config['baseUrl']}/#review={row['token']}"
        status = "模拟成功" if config["mode"] == "mock" else "待接入"
        error = None
        if config["mode"] == "off":
            status, error = "未发送", "钉钉集成已关闭"
        elif config["mode"] == "live":
            missing = dingtalk_status()["missing"]
            if missing:
                status, error = "发送失败", f"缺少钉钉配置：{', '.join(missing)}"
            else:
                status, error = "待接入", "真实钉钉接口将在取得企业应用权限后启用"
        log_id = uuid.uuid4().hex
        conn.execute(
            """INSERT INTO notification_logs(id,assignment_id,channel,recipient_external_id,recipient_name,title,content,target_url,status,provider_message_id,error_message,created_at)
               VALUES(?,?,?,?,?,?,?,?,?,?,?,?)""",
            (log_id, assignment_id, "dingtalk", row["external_user_id"], row["reviewer_name"], "任职资格评审提醒", f"请评审 {row['submission_name']} 的认证材料", target_url, status, None, error, now_text()),
        )
    return {"id": log_id, "status": status, "targetUrl": target_url, "error": error}


def retry_notification(log_id: str) -> dict:
    with db() as conn:
        row = conn.execute("SELECT assignment_id FROM notification_logs WHERE id=?", (log_id,)).fetchone()
    if not row or not row["assignment_id"]:
        raise ValueError("通知记录不存在或无法重发")
    return send_assignment_notification(row["assignment_id"])


def system_health() -> dict:
    DATA_DIR.mkdir(exist_ok=True)
    with db() as conn:
        integrity = conn.execute("PRAGMA integrity_check").fetchone()[0]
        counts = {
            "submissions": conn.execute("SELECT COUNT(*) FROM submissions").fetchone()[0],
            "assignments": conn.execute("SELECT COUNT(*) FROM review_assignments").fetchone()[0],
            "notifications": conn.execute("SELECT COUNT(*) FROM notification_logs").fetchone()[0],
        }
    usage = shutil.disk_usage(DATA_DIR)
    backup_dir = ROOT / "backups"
    backups = sorted(backup_dir.glob("certification-backup-*.zip"), reverse=True) if backup_dir.exists() else []
    config = dingtalk_config()
    log_files = sorted(LOG_DIR.glob("*.log*"), key=lambda item: item.stat().st_mtime, reverse=True)
    return {
        "status": "healthy" if integrity == "ok" else "degraded",
        "database": {"integrity": integrity, "path": str(DB_PATH), "sizeBytes": DB_PATH.stat().st_size if DB_PATH.exists() else 0},
        "storage": {"freeBytes": usage.free, "totalBytes": usage.total, "writable": os.access(DATA_DIR, os.W_OK)},
        "counts": counts,
        "backup": {"latest": backups[0].name if backups else None, "count": len(backups)},
        "logs": {
            "directory": str(LOG_DIR), "count": len(log_files),
            "latest": log_files[0].name if log_files else None,
            "sizeBytes": sum(item.stat().st_size for item in log_files),
            "retentionDays": 30,
        },
        "configuration": {
            "host": os.environ.get("CERT_HOST", "127.0.0.1"), "port": int(os.environ.get("CERT_PORT", "8787")),
            "baseUrl": config["baseUrl"], "dingtalk": dingtalk_status(),
            "productionWarnings": [warning for warning in [
                "服务当前只监听本机" if os.environ.get("CERT_HOST", "127.0.0.1") in {"127.0.0.1", "localhost"} else None,
                "仍在使用钉钉模拟模式" if config["mode"] == "mock" else None,
                "尚未生成备份" if not backups else None,
            ] if warning],
        },
        "checkedAt": now_text(),
    }


def create_diagnostic_bundle() -> tuple[bytes, str]:
    """Create a support bundle containing logs and non-sensitive runtime metadata only."""
    health = system_health()
    summary = {
        "generatedAt": now_text(),
        "version": (ROOT / "VERSION").read_text(encoding="utf-8").strip() if (ROOT / "VERSION").exists() else "development",
        "status": health["status"],
        "database": {"integrity": health["database"]["integrity"], "sizeBytes": health["database"]["sizeBytes"]},
        "storage": {"freeBytes": health["storage"]["freeBytes"], "writable": health["storage"]["writable"]},
        "counts": health["counts"],
        "backup": health["backup"],
        "logs": health["logs"],
        "runtime": {"python": sys.version, "platform": sys.platform, "frozen": bool(getattr(sys, "frozen", False))},
        "privacy": "No database, attachments, configuration, access codes, cookies, or request bodies are included.",
    }
    stamp = datetime.now().strftime("%Y%m%d-%H%M%S")
    filename = f"CertificationOps-Diagnostics-{stamp}.zip"
    output = io.BytesIO()
    with zipfile.ZipFile(output, "w", zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("diagnostic-summary.json", json.dumps(summary, ensure_ascii=False, indent=2))
        archive.writestr("README.txt", "诊断包仅包含运行日志和非敏感系统摘要，不包含数据库、附件、配置文件、访问码或请求正文。\n")
        for path in sorted(LOG_DIR.glob("*.log*")):
            if not path.is_file():
                continue
            data = path.read_bytes()
            if len(data) > 5 * 1024 * 1024:
                data = b"[truncated to last 5 MB]\n" + data[-5 * 1024 * 1024:]
            archive.writestr(f"logs/{path.name}", data)
    APP_LOG.info("Diagnostic bundle generated files=%s", len(list(LOG_DIR.glob("*.log*"))))
    return output.getvalue(), filename


def create_runtime_backup(retention: int = 14) -> Path:
    # Keep backups beside the active data directory. This also supports
    # deployments that place persistent data outside the application tree.
    backup_dir = DATA_DIR.parent / "backups"
    backup_dir.mkdir(parents=True, exist_ok=True)
    stamp = datetime.now().strftime("%Y%m%d-%H%M%S")
    target = backup_dir / f"certification-backup-{stamp}.zip"
    suffix = 1
    while target.exists():
        target = backup_dir / f"certification-backup-{stamp}-{suffix}.zip"
        suffix += 1
    with tempfile.TemporaryDirectory(prefix="cert-backup-") as temp_name:
        snapshot = Path(temp_name) / "certification.db"
        source = sqlite3.connect(DB_PATH)
        destination = sqlite3.connect(snapshot)
        try:
            source.backup(destination)
        finally:
            destination.close()
            source.close()
        with zipfile.ZipFile(target, "w", compression=zipfile.ZIP_DEFLATED, compresslevel=6) as archive:
            archive.write(snapshot, "data/certification.db")
            archive.writestr("metadata.json", json.dumps({"createdAt": now_text()}, ensure_ascii=False, indent=2))
            if SUBMISSION_DIR.exists():
                for path in SUBMISSION_DIR.rglob("*"):
                    if path.is_file():
                        archive.write(path, f"data/{path.relative_to(DATA_DIR).as_posix()}")
    for old in sorted(backup_dir.glob("certification-backup-*.zip"), reverse=True)[max(1, retention):]:
        old.unlink()
    return target


def restore_runtime_backup(backup: Path) -> None:
    backup = backup.resolve()
    if not backup.is_file() or backup.suffix.lower() != ".zip":
        raise ValueError("备份文件无效")
    if DB_PATH.exists():
        create_runtime_backup(retention=30)
    with tempfile.TemporaryDirectory(prefix="cert-restore-") as temp_name:
        temp = Path(temp_name)
        with zipfile.ZipFile(backup) as archive:
            root = temp.resolve()
            for member in archive.infolist():
                destination = (temp / member.filename).resolve()
                if root != destination and root not in destination.parents:
                    raise ValueError("备份文件包含不安全路径")
            archive.extractall(temp)
        restored_db = temp / "data" / "certification.db"
        if not restored_db.exists():
            raise ValueError("备份中没有数据库")
        conn = sqlite3.connect(restored_db)
        try:
            if conn.execute("PRAGMA integrity_check").fetchone()[0] != "ok":
                raise ValueError("备份数据库校验失败")
        finally:
            conn.close()
        DATA_DIR.mkdir(exist_ok=True)
        shutil.copy2(restored_db, DB_PATH)
        restored_files = temp / "data" / "submissions"
        if restored_files.exists():
            if SUBMISSION_DIR.exists():
                shutil.rmtree(SUBMISSION_DIR)
            shutil.copytree(restored_files, SUBMISSION_DIR)


def list_assignments() -> dict:
    with db() as conn:
        assignments = [dict(row) for row in conn.execute(
            """SELECT a.id,a.submission_id,a.weight,a.token,a.status,a.pushed_at,a.completed_at,
                      s.name submission_name,s.level,s.category,r.id reviewer_id,r.name reviewer_name,r.external_user_id reviewer_user_id,
                      CASE WHEN r.access_code_hash IS NULL THEN 0 ELSE 1 END access_code_set
               FROM review_assignments a JOIN submissions s ON s.id=a.submission_id
               JOIN reviewers r ON r.id=a.reviewer_id ORDER BY a.created_at DESC"""
        )]
        reviewers = [dict(row) for row in conn.execute("SELECT id,name,CASE WHEN access_code_hash IS NULL THEN 0 ELSE 1 END access_code_set FROM reviewers ORDER BY name")]
    with db() as conn:pending_evidence={r[0] for r in conn.execute('SELECT submission_id FROM evidence_tasks WHERE finalized_at IS NULL')}
    ready_submissions = [item for item in list_submissions() if item["status"] != "填报中" and item['id'] not in pending_evidence]
    with db() as conn:
        for item in ready_submissions:
            item['reviewWeightsReady'] = matrix_review_weights_ready(conn, item['id'])
    return {"assignments": assignments, "reviewers": reviewers, "contacts": list_directory_users(), "integration": dingtalk_status(), "submissions": ready_submissions}


def public_reviewer_directory() -> list[dict]:
    with db() as conn:
        rows = conn.execute(
            """SELECT r.id,r.name,CASE WHEN r.access_code_hash IS NULL THEN 0 ELSE 1 END access_code_set
               FROM reviewers r
               WHERE EXISTS (SELECT 1 FROM review_assignments a WHERE a.reviewer_id=r.id)
               ORDER BY r.name"""
        ).fetchall()
    return [dict(row) for row in rows]


def set_reviewer_access_code(reviewer_id: str, code: str) -> None:
    code = clean(code)
    if not re.fullmatch(r"\d{4,8}", code):
        raise ValueError("访问码必须是 4 至 8 位数字")
    salt = secrets.token_hex(16)
    digest = hashlib.pbkdf2_hmac("sha256", code.encode(), bytes.fromhex(salt), 120_000).hex()
    with db() as conn:
        if not conn.execute("SELECT 1 FROM reviewers WHERE id=?", (reviewer_id,)).fetchone():
            raise ValueError("评价人不存在")
        conn.execute("UPDATE reviewers SET access_code_salt=?,access_code_hash=? WHERE id=?", (salt, digest, reviewer_id))
        conn.execute("DELETE FROM reviewer_sessions WHERE reviewer_id=?", (reviewer_id,))


def create_reviewer(name: str) -> dict:
    name = clean(name)
    if not name:
        raise ValueError("请填写评价人姓名")
    with db() as conn:
        existing = conn.execute("SELECT id,name FROM reviewers WHERE name=?", (name,)).fetchone()
        if existing:
            raise ValueError("该评价人已存在")
        reviewer_id = uuid.uuid4().hex
        conn.execute(
            "INSERT INTO reviewers(id,name,created_at,directory_provider) VALUES(?,?,?,?)",
            (reviewer_id, name, now_text(), "manual"),
        )
    return {"id": reviewer_id, "name": name}


def delete_reviewer(reviewer_id: str) -> None:
    with db() as conn:
        reviewer = conn.execute("SELECT id FROM reviewers WHERE id=?", (reviewer_id,)).fetchone()
        if not reviewer:
            raise ValueError("评价人不存在")
        conn.execute(
            "DELETE FROM reviews WHERE assignment_id IN (SELECT id FROM review_assignments WHERE reviewer_id=?)",
            (reviewer_id,),
        )
        conn.execute("DELETE FROM reviewers WHERE id=?", (reviewer_id,))


def verify_reviewer_access(reviewer_id: str, code: str) -> dict:
    with db() as conn:
        reviewer = conn.execute("SELECT id,name,access_code_salt,access_code_hash FROM reviewers WHERE id=?", (reviewer_id,)).fetchone()
        if not reviewer or not reviewer["access_code_hash"]:
            raise ValueError("该评价人尚未设置访问码，请联系管理员")
        digest = hashlib.pbkdf2_hmac("sha256", clean(code).encode(), bytes.fromhex(reviewer["access_code_salt"]), 120_000).hex()
        if not hmac.compare_digest(digest, reviewer["access_code_hash"]):
            raise ValueError("访问码错误")
        token = secrets.token_hex(32)
        expires_at = datetime.fromtimestamp(time.time() + 12 * 3600).astimezone().isoformat(timespec="seconds")
        conn.execute("DELETE FROM reviewer_sessions WHERE expires_at<?", (now_text(),))
        conn.execute("INSERT INTO reviewer_sessions(token,reviewer_id,expires_at,created_at) VALUES(?,?,?,?)", (token, reviewer_id, expires_at, now_text()))
    return {"token": token, "reviewerName": reviewer["name"], "expiresAt": expires_at}


def authorized_reviewer_tasks(reviewer_id: str, session_token: str) -> list[dict]:
    with db() as conn:
        valid = conn.execute("SELECT 1 FROM reviewer_sessions WHERE token=? AND reviewer_id=? AND expires_at>?", (session_token, reviewer_id, now_text())).fetchone()
    if not valid:
        raise ValueError("身份验证已失效，请重新输入访问码")
    return public_reviewer_tasks(reviewer_id)


def assignment_import_template() -> tuple[bytes, str]:
    template_path = STATIC_DIR / "评价关系批量导入模板.xlsx"
    if template_path.is_file():
        return template_path.read_bytes(), template_path.name
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "评价关系"
    sheet.append(["被评价人", "评价人", "权重（%）", "说明"])
    sheet.append(["张三", "评委甲", 40, "同一被评价人有多位评价人时，请重复填写多行"])
    sheet.append(["张三", "评委乙", 60, "权重填写 1-100 的数字"])
    for cell in sheet[1]:
        cell.fill = PatternFill("solid", fgColor="2F75B5")
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center")
    sheet.freeze_panes = "A2"
    sheet.column_dimensions["A"].width = 20
    sheet.column_dimensions["B"].width = 20
    sheet.column_dimensions["C"].width = 14
    sheet.column_dimensions["D"].width = 46
    notes = workbook.create_sheet("填写说明")
    notes.append(["填写规则"])
    notes.append(["1. 一行代表一条评价关系；同一被评价人对应多人时重复填写多行。"])
    notes.append(["2. 被评价人姓名必须与系统已导入姓名完全一致。"])
    notes.append(["3. 权重填写 1 至 100，例如 40 表示 40%。"])
    notes.append(["4. 已存在关系会更新权重，不会覆盖该评委已经提交的评分。"])
    output = io.BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue(), "评价关系批量导入模板.xlsx"


def import_assignment_workbook(data: bytes) -> dict:
    try:
        workbook = load_workbook(io.BytesIO(data), data_only=True, read_only=True)
    except Exception as exc:
        raise ValueError(f"无法读取 Excel：{exc}") from exc
    sheet = workbook["评价关系"] if "评价关系" in workbook.sheetnames else workbook.active
    rows, errors = [], []
    for row_number, values in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
        candidate, reviewer, weight_value = clean(values[0] if len(values)>0 else ""), clean(values[1] if len(values)>1 else ""), values[2] if len(values)>2 else None
        if not candidate and not reviewer and weight_value in {None, ""}:
            continue
        try:
            weight = float(str(weight_value).replace("%", ""))
            if 0 < weight <= 1:
                weight *= 100
        except (TypeError, ValueError):
            errors.append(f"第 {row_number} 行：权重不是有效数字")
            continue
        if not candidate or not reviewer or not 1 <= weight <= 100:
            errors.append(f"第 {row_number} 行：请填写被评价人、评价人及 1-100 的权重")
        else:
            rows.append((row_number, candidate, reviewer, weight / 100))
    workbook.close()
    with db() as conn:
        submission_by_name = {row["name"]: row["id"] for row in conn.execute("SELECT id,name FROM submissions")}
    for row_number, candidate, reviewer, weight in rows:
        if candidate not in submission_by_name:
            errors.append(f"第 {row_number} 行：系统中找不到被评价人“{candidate}”")
        else:
            try:
                with db() as conn:ensure_assignment_ready(conn, submission_by_name[candidate])
            except ValueError as exc:
                errors.append(f"第 {row_number} 行：{candidate}：{exc}")
    if errors:
        return {"created": 0, "updated": 0, "errors": errors}
    created = updated = 0
    for _, candidate, reviewer, weight in rows:
        with db() as conn:
            existing = conn.execute("""SELECT a.id FROM review_assignments a JOIN reviewers r ON r.id=a.reviewer_id
                                       WHERE a.submission_id=? AND r.name=?""", (submission_by_name[candidate], reviewer)).fetchone()
        upsert_assignment({"submissionId": submission_by_name[candidate], "reviewerName": reviewer, "weight": weight})
        updated += bool(existing)
        created += not bool(existing)
    return {"created": created, "updated": updated, "errors": []}


def public_reviewer_tasks(reviewer_id: str) -> list[dict]:
    with db() as conn:
        rows = conn.execute(
            """SELECT a.token,a.status,s.name submission_name,s.level,s.category
               FROM review_assignments a JOIN submissions s ON s.id=a.submission_id
               WHERE a.reviewer_id=?
                 AND NOT EXISTS (SELECT 1 FROM evidence_tasks e
                                 WHERE e.submission_id=s.id AND e.finalized_at IS NULL)
               ORDER BY s.name""", (reviewer_id,)
        ).fetchall()
    items = []
    for row in rows:
        item = dict(row)
        item["status"] = "已完成" if item["status"] in {"已评价", "已完成"} else "评价中" if item["status"] in {"评价中", "待评价"} else "未开始"
        items.append(item)
    return items


def admin_assignment_review_detail(assignment_id: str) -> dict:
    with db() as conn:
        assignment = conn.execute(
            """SELECT a.id,a.weight,a.status,s.name submission_name,r.name reviewer_name,rv.id review_id,
                      rv.strengths,rv.improvements,rv.opinion,rv.opinion_note,rv.status review_status,rv.updated_at
               FROM review_assignments a JOIN submissions s ON s.id=a.submission_id
               JOIN reviewers r ON r.id=a.reviewer_id
               LEFT JOIN reviews rv ON rv.assignment_id=a.id WHERE a.id=?""", (assignment_id,)
        ).fetchone()
        if not assignment:
            raise ValueError("评价关系不存在")
        scores = []
        if assignment["review_id"]:
            rows = conn.execute(
                """SELECT m.section_type,m.evaluation_item,m.module_name,m.weight,ms.score,ms.comment
                   FROM modules m LEFT JOIN module_scores ms ON ms.module_id=m.id AND ms.review_id=?
                   WHERE m.submission_id=(SELECT submission_id FROM review_assignments WHERE id=?)
                   ORDER BY m.sort_order""", (assignment["review_id"], assignment_id)
            ).fetchall()
            for row in rows:
                item = dict(row)
                raw = item["score"]
                item["convertedScore"] = None if raw is None else round(raw if item["section_type"] == "bonus" else raw / 5 * 100 * float(item["weight"] or 0), 2)
                scores.append(item)
    result = dict(assignment)
    result["scores"] = scores
    return result


def update_submission_basic(submission_id: str, payload: dict) -> None:
    with db() as conn:
        if evidence.task(conn,submission_id):raise ValueError('已配置的在线举证任务身份和标准已锁定，请重新配置任务')
    name = clean(payload.get("name"))
    level, category = clean(payload.get("level")), clean(payload.get("category"))
    if not name:
        raise ValueError("被评价人姓名不能为空")
    with db() as conn:
        existing = conn.execute("SELECT level,category,standard_snapshot_json FROM submissions WHERE id=?", (submission_id,)).fetchone()
        if not existing:
            raise ValueError("被评价人不存在")
        if json.loads(existing["standard_snapshot_json"] or "{}").get("matrix") and (level != existing["level"] or category != existing["category"]):
            raise ValueError("分级标准申请的职级和职类已与标准快照绑定，请重新发起对应职级的申请")
        conn.execute("UPDATE submissions SET name=?,level=?,category=? WHERE id=?", (name, level, category, submission_id))


def delete_submission(submission_id: str) -> None:
    with db() as conn:
        row = conn.execute("SELECT raw_path FROM submissions WHERE id=?", (submission_id,)).fetchone()
        if not row:
            raise ValueError("被评价人不存在")
        conn.execute("DELETE FROM submissions WHERE id=?", (submission_id,))
    folder = SUBMISSION_DIR / submission_id
    if folder.exists() and folder.resolve().parent == SUBMISSION_DIR.resolve():
        shutil.rmtree(folder)


def matrix_review_weights_ready(conn, submission_id):
    submission = conn.execute('SELECT standard_snapshot_json FROM submissions WHERE id=?', (submission_id,)).fetchone()
    if not submission or not json.loads(submission['standard_snapshot_json'] or '{}').get('matrix'):
        return True
    weights = [r['weight'] for r in conn.execute("SELECT weight FROM modules WHERE submission_id=? AND section_type IN ('standard','reference')", (submission_id,))]
    return not weights or (all(w is not None and 0 < w <= 1 for w in weights) and abs(sum(weights)-1) <= .001)


def ensure_assignment_ready(conn, submission_id):
    evidence.ensure_finalized(conn, submission_id)
    if not matrix_review_weights_ready(conn, submission_id):
        raise ValueError('请先完整配置子维度权重（合计100%），再进行评委配置')


def save_matrix_review_weights(submission_id, payload):
    with db() as conn:
        conn.execute("BEGIN IMMEDIATE")
        submission = conn.execute("SELECT standard_snapshot_json FROM submissions WHERE id=?", (submission_id,)).fetchone()
        if not submission or not json.loads(submission["standard_snapshot_json"] or "{}").get("matrix"):
            raise ValueError("仅分级标准申请使用此权重配置")
        if conn.execute("SELECT 1 FROM reviews WHERE submission_id=?", (submission_id,)).fetchone():
            raise ValueError("已有评审记录，不能修改评分权重")
        modules = [r["id"] for r in conn.execute("SELECT id FROM modules WHERE submission_id=? AND section_type IN ('standard','reference')", (submission_id,))]
        supplied = payload.get("weights", {})
        if not isinstance(supplied, dict) or set(supplied) != set(modules):
            raise ValueError("请完整配置本标准所有普通评分子项的权重（成果贡献加分项除外）")
        try:
            weights = {key:float(value) for key,value in supplied.items()}
        except (TypeError, ValueError):
            raise ValueError("权重必须是数字")
        if any(not 0 < value <= 1 for value in weights.values()) or (weights and abs(sum(weights.values())-1) > .001):
            raise ValueError("各项权重必须大于0，合计为100%")
        for mid,weight in weights.items():
            conn.execute("UPDATE modules SET weight=? WHERE id=?", (weight,mid))
    return {"ok":True}


def store_manual_attachment(submission_id: str, filename: str, source, length: int, project_name: str = "补充附件") -> dict:
    with db() as conn:
        if evidence.task(conn,submission_id):raise ValueError('在线举证附件由举证人维护；请先反馈或退回修改')
    filename = Path(filename).name
    if not filename:
        raise ValueError("附件文件名不能为空")
    with db() as conn:
        if not conn.execute("SELECT 1 FROM submissions WHERE id=?", (submission_id,)).fetchone():
            raise ValueError("被评价人不存在")
        count = conn.execute("SELECT COUNT(*) FROM attachments WHERE submission_id=?", (submission_id,)).fetchone()[0]
    target_dir = SUBMISSION_DIR / submission_id / "attachments"
    target_dir.mkdir(parents=True, exist_ok=True)
    target = target_dir / f"{uuid.uuid4().hex[:8]}-{filename}"
    remaining = length
    with target.open("wb") as output:
        while remaining:
            chunk = source.read(min(UPLOAD_CHUNK, remaining))
            if not chunk:
                target.unlink(missing_ok=True)
                raise ValueError("附件上传连接提前中断")
            output.write(chunk)
            remaining -= len(chunk)
    attachment_id = uuid.uuid4().hex
    with db() as conn:
        conn.execute(
            """INSERT INTO attachments(id,submission_id,serial,project_name,material_name,storage_path,image_path,uploaded_file_path,source_row,sort_order,link_url)
               VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)""",
            (attachment_id, submission_id, f"补充-{count + 1}", project_name, filename, "管理员手动上传", None, str(target.relative_to(DATA_DIR)), None, count, None),
        )
    return {"id": attachment_id, "name": filename, "size": length}


def update_assignment(assignment_id: str, payload: dict) -> None:
    reviewer_name = clean(payload.get("reviewerName"))
    try:
        weight = float(payload.get("weight"))
    except (TypeError, ValueError):
        raise ValueError("评价人权重必须是数字")
    if not reviewer_name or not 0 < weight <= 1:
        raise ValueError("请填写评价人姓名，并设置 1% 至 100% 的权重")
    with db() as conn:
        conn.execute('BEGIN IMMEDIATE')
        row = conn.execute("SELECT reviewer_id,submission_id FROM review_assignments WHERE id=?", (assignment_id,)).fetchone()
        if not row:
            raise ValueError("评价关系不存在")
        ensure_assignment_ready(conn, row['submission_id'])
        conflict = conn.execute("SELECT id FROM reviewers WHERE name=? AND id<>?", (reviewer_name, row["reviewer_id"])).fetchone()
        if conflict:
            raise ValueError("该评价人姓名已存在，请删除错误关系后重新配置")
        conn.execute("UPDATE reviewers SET name=? WHERE id=?", (reviewer_name, row["reviewer_id"]))
        conn.execute("UPDATE reviews SET reviewer_name=? WHERE assignment_id IN (SELECT id FROM review_assignments WHERE reviewer_id=?)", (reviewer_name, row["reviewer_id"]))
        conn.execute("UPDATE review_assignments SET weight=? WHERE id=?", (weight, assignment_id))


def delete_assignment(assignment_id: str) -> None:
    with db() as conn:
        row = conn.execute("SELECT reviewer_id FROM review_assignments WHERE id=?", (assignment_id,)).fetchone()
        if not row:
            raise ValueError("评价关系不存在")
        conn.execute("DELETE FROM reviews WHERE assignment_id=?", (assignment_id,))
        conn.execute("DELETE FROM review_assignments WHERE id=?", (assignment_id,))
        remaining = conn.execute("SELECT COUNT(*) FROM review_assignments WHERE reviewer_id=?", (row["reviewer_id"],)).fetchone()[0]
        if not remaining:
            conn.execute("DELETE FROM reviewers WHERE id=?", (row["reviewer_id"],))


def upsert_assignment(payload: dict) -> dict:
    submission_id = clean(payload.get("submissionId"))
    reviewer_name = clean(payload.get("reviewerName"))
    external_user_id = clean(payload.get("reviewerUserId"))
    try:
        weight = float(payload.get("weight"))
    except (TypeError, ValueError):
        raise ValueError("评委权重必须是数字")
    if not submission_id or not reviewer_name:
        raise ValueError("请选择被评价人并填写评委姓名")
    if not 0 < weight <= 1:
        raise ValueError("评委权重必须大于 0% 且不超过 100%")
    with db() as conn:
        conn.execute('BEGIN IMMEDIATE')
        if not conn.execute("SELECT 1 FROM submissions WHERE id=?", (submission_id,)).fetchone():
            raise ValueError("被评价人不存在")
        ensure_assignment_ready(conn, submission_id)
        reviewer = conn.execute("SELECT id FROM reviewers WHERE external_user_id=?", (external_user_id,)).fetchone() if external_user_id else conn.execute("SELECT id FROM reviewers WHERE name=?", (reviewer_name,)).fetchone()
        reviewer_id = reviewer["id"] if reviewer else uuid.uuid4().hex
        if not reviewer:
            conn.execute("INSERT INTO reviewers(id,name,created_at,external_user_id,directory_provider) VALUES(?,?,?,?,?)", (reviewer_id, reviewer_name, now_text(), external_user_id or None, "dingtalk" if external_user_id else "manual"))
        elif external_user_id:
            conn.execute("UPDATE reviewers SET name=?,external_user_id=?,directory_provider='dingtalk' WHERE id=?", (reviewer_name, external_user_id, reviewer_id))
        if external_user_id:
            conn.execute(
                """INSERT INTO app_users(id,external_user_id,name,role,active,created_at,updated_at)
                   VALUES(?,?,?,?,1,?,?) ON CONFLICT(external_user_id) DO UPDATE SET name=excluded.name,active=1,updated_at=excluded.updated_at""",
                (uuid.uuid4().hex, external_user_id, reviewer_name, "评委", now_text(), now_text()),
            )
        existing = conn.execute("SELECT id,token,status FROM review_assignments WHERE submission_id=? AND reviewer_id=?", (submission_id, reviewer_id)).fetchone()
        if existing:
            conn.execute("UPDATE review_assignments SET weight=? WHERE id=?", (weight, existing["id"]))
            assignment_id, token, status = existing["id"], existing["token"], existing["status"]
        else:
            assignment_id, token, status = uuid.uuid4().hex, uuid.uuid4().hex, "未开始"
            conn.execute("INSERT INTO review_assignments VALUES(?,?,?,?,?,?,?,?,?)", (assignment_id, submission_id, reviewer_id, weight, token, status, None, None, now_text()))
    return {"id": assignment_id, "token": token, "status": status}


def push_assignment(assignment_id: str) -> dict:
    with db() as conn:
        row = conn.execute("SELECT token,status FROM review_assignments WHERE id=?", (assignment_id,)).fetchone()
        if not row:
            raise ValueError("评审分配不存在")
        if row["status"] not in {"已评价", "已完成"}:
            conn.execute("UPDATE review_assignments SET status='未开始',pushed_at=? WHERE id=?", (now_text(), assignment_id))
    notification = send_assignment_notification(assignment_id)
    return {"token": row["token"], "reviewPath": f"/#review={row['token']}", "notification": notification}


def assigned_review_detail(token: str) -> dict | None:
    with db() as conn:
        assignment = conn.execute(
            """SELECT a.id,a.submission_id,a.status,r.name reviewer_name
               FROM review_assignments a JOIN reviewers r ON r.id=a.reviewer_id
               WHERE a.token=?""", (token,)
        ).fetchone()
    if not assignment:
        return None
    detail = submission_detail(assignment["submission_id"], assignment["id"], include_workflow=False)
    if not detail:
        return None
    detail["assignment"] = {"id": assignment["id"], "reviewerName": assignment["reviewer_name"], "status": assignment["status"]}
    return detail


def dashboard_data() -> dict:
    with db() as conn:
        submissions = [dict(row) for row in conn.execute("SELECT * FROM submissions ORDER BY imported_at DESC")]
        results = []
        for submission in submissions:
            assignments = conn.execute(
                """SELECT a.id,a.weight,a.status,r.name reviewer_name,rv.id review_id
                   FROM review_assignments a JOIN reviewers r ON r.id=a.reviewer_id
                   LEFT JOIN reviews rv ON rv.assignment_id=a.id
                   WHERE a.submission_id=? ORDER BY r.name""", (submission["id"],)
            ).fetchall()
            reviewer_results = []
            weighted_total = 0.0
            completed = 0
            weight_total = 0.0
            for assignment in assignments:
                weight_total += float(assignment["weight"])
                total = None
                contribution = None
                if assignment["status"] in {"已评价", "已完成"} and assignment["review_id"]:
                    total = calculate_review_summary(conn, assignment["review_id"])["total"]
                    contribution = total * float(assignment["weight"])
                    weighted_total += contribution
                    completed += 1
                display_status = "已完成" if assignment["status"] in {"已评价", "已完成"} else "评价中" if assignment["status"] in {"评价中", "待评价"} else "未开始"
                reviewer_results.append({"assignmentId": assignment["id"], "reviewerName": assignment["reviewer_name"], "status": display_status, "weight": assignment["weight"], "score": total, "weightedContribution": None if contribution is None else round(contribution, 2)})
            all_done = bool(assignments) and completed == len(assignments)
            any_started = any(item["status"] in {"评价中", "已完成"} for item in reviewer_results)
            results.append({
                "submissionId": submission["id"], "name": submission["name"], "level": submission["level"], "category": submission["category"],
                "status": "已完成" if all_done else "评价中" if any_started else "未开始", "assignmentCount": len(assignments), "completedCount": completed,
                "weightTotal": round(weight_total, 4), "weightedTotal": round(weighted_total, 2) if completed else None,
                "reviewers": reviewer_results,
            })
    return {"items": results}


def _export_status(assignment_status: str, review_status: str | None) -> str:
    if assignment_status in {"已评价", "已完成"} or review_status == "已提交":
        return "已完成"
    if review_status == "草稿" or assignment_status in {"评价中", "待评价"}:
        return "评价中"
    return "未开始"


def export_submission_reviews(submission_id: str) -> tuple[bytes, str]:
    """Export one candidate's reviewer scores without mutating business data."""
    with db() as conn:
        submission = conn.execute("SELECT * FROM submissions WHERE id=?", (submission_id,)).fetchone()
        if not submission:
            raise ValueError("被评价人不存在")
        modules = [dict(row) for row in conn.execute(
            "SELECT * FROM modules WHERE submission_id=? ORDER BY sort_order", (submission_id,)
        )]
        assignments = [dict(row) for row in conn.execute(
            """SELECT a.id,a.weight,a.status,a.completed_at,r.name reviewer_name
               FROM review_assignments a JOIN reviewers r ON r.id=a.reviewer_id
               WHERE a.submission_id=? ORDER BY r.name""", (submission_id,)
        )]
        export_rows = []
        for assignment in assignments:
            review = conn.execute(
                """SELECT * FROM reviews
                   WHERE submission_id=? AND (assignment_id=? OR (assignment_id IS NULL AND reviewer_name=?))
                   ORDER BY CASE WHEN assignment_id=? THEN 0 ELSE 1 END, updated_at DESC LIMIT 1""",
                (submission_id, assignment["id"], assignment["reviewer_name"], assignment["id"]),
            ).fetchone()
            scores = {}
            summary = None
            if review:
                scores = {row["module_id"]: dict(row) for row in conn.execute(
                    "SELECT * FROM module_scores WHERE review_id=?", (review["id"],)
                )}
                summary = calculate_review_summary(conn, review["id"])
            export_rows.append({
                "assignment": assignment,
                "review": dict(review) if review else None,
                "scores": scores,
                "summary": summary,
                "status": _export_status(assignment["status"], review["status"] if review else None),
            })

    standard_modules = [module for module in modules if module["section_type"] in {"standard", "reference"}]
    bonus_modules = [module for module in modules if module["section_type"] == "bonus"]
    scoring_label = "各维度评分" if json.loads(submission['standard_snapshot_json'] or '{}').get('matrix') else "行为标准"
    workbook = Workbook()
    summary_sheet = workbook.active
    summary_sheet.title = "评委评分汇总"
    detail_sheet = workbook.create_sheet("逐项评分明细")
    notes_sheet = workbook.create_sheet("字段与计算说明")

    navy, blue = "17365D", "2F75B5"
    pale_blue, pale_green, pale_gold, pale_gray = "DDEBF7", "E2F0D9", "FFF2CC", "F2F2F2"
    white, text, line = "FFFFFF", "243447", "C9D2DF"
    thin = Side(style="thin", color=line)
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    base_headers = ["评价人", "评审状态", "评价人权重"]
    standard_headers = [module["module_name"] or module["evaluation_item"] or "未命名维度" for module in standard_modules]
    bonus_headers = [module["module_name"] or module["evaluation_item"] or "未命名加分项" for module in bonus_modules]
    tail_headers = [scoring_label + "折算总分", "成长贡献加分合计", "评价人总分", "加权后总分", "提交时间", "优势", "不足或待改进项", "评估结论", "评估意见补充"]
    headers = base_headers + standard_headers + tail_headers[:1] + bonus_headers + tail_headers[1:]
    max_col = len(headers)

    summary_sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=max_col)
    summary_sheet.cell(1, 1, "任职资格认证｜评委评分导出表")
    summary_sheet.merge_cells(start_row=2, start_column=1, end_row=2, end_column=max_col)
    summary_sheet.cell(2, 1, "工作模块列根据该被评价人的实际评分维度动态生成；本文件为导出时点的数据快照。")
    summary_sheet.merge_cells(start_row=4, start_column=1, end_row=4, end_column=max_col)
    summary_sheet.cell(4, 1, "一、被评价人基本信息")
    basic = [
        "姓名", submission["name"], "申报职级", submission["level"] or "",
        "职类", submission["category"] or "", "申请时间", submission["applied_at"] or "",
    ]
    for index, value in enumerate(basic, start=1):
        summary_sheet.cell(5, index, value)
    progress = "未开始"
    if export_rows and all(row["status"] == "已完成" for row in export_rows):
        progress = "已完成"
    elif any(row["status"] != "未开始" for row in export_rows):
        progress = "评价中"
    secondary = ["导出时间", now_text(), "评价人数", len(export_rows), "评价进度", progress, "材料状态", submission["status"] or ""]
    for index, value in enumerate(secondary, start=1):
        summary_sheet.cell(6, index, value)

    groups = []
    col = 1
    groups.append((col, col + 2, "评价信息")); col += 3
    if standard_modules:
        groups.append((col, col + len(standard_modules), scoring_label + "（每项原始分满分 5 分）")); col += len(standard_modules) + 1
    else:
        groups.append((col, col, scoring_label + "折算")); col += 1
    if bonus_modules:
        groups.append((col, col + len(bonus_modules), "团队成长贡献加分（每项最多 2 分）")); col += len(bonus_modules) + 1
    else:
        groups.append((col, col, "成长贡献加分")); col += 1
    groups.append((col, col + 2, "分数汇总")); col += 3
    groups.append((col, max_col, "定性评价与结论"))
    for start, end, label in groups:
        summary_sheet.merge_cells(start_row=8, start_column=start, end_row=8, end_column=end)
        summary_sheet.cell(8, start, label)
    for index, header in enumerate(headers, start=1):
        summary_sheet.cell(9, index, header)
    annotations = ["", "", ""]
    annotations += [f"模块权重 {float(module['weight'] or 0):.0%}" for module in standard_modules]
    annotations += ["自动折算"]
    annotations += ["原始加分" for _ in bonus_modules]
    annotations += ["直接相加", scoring_label + "折算+加分", "仅已完成评价计入", "", "", "", "通过/不通过/其他", ""]
    for index, value in enumerate(annotations, start=1):
        summary_sheet.cell(10, index, value)

    data_start = 11
    for row_index, item in enumerate(export_rows, start=data_start):
        assignment, review, scores, score_summary = item["assignment"], item["review"], item["scores"], item["summary"]
        values = [assignment["reviewer_name"], item["status"], float(assignment["weight"])]
        values += [scores.get(module["id"], {}).get("score") for module in standard_modules]
        values += [score_summary["behaviorTotal"] if score_summary else None]
        values += [scores.get(module["id"], {}).get("score") for module in bonus_modules]
        bonus_total = score_summary["bonusTotal"] if score_summary else None
        total = score_summary["total"] if score_summary else None
        weighted = round(total * float(assignment["weight"]), 2) if total is not None and item["status"] == "已完成" else None
        values += [
            bonus_total, total, weighted,
            (assignment["completed_at"] or review["updated_at"]) if review else assignment["completed_at"],
            review["strengths"] if review else "", review["improvements"] if review else "",
            review["opinion"] if review else "", review["opinion_note"] if review else "",
        ]
        for column_index, value in enumerate(values, start=1):
            summary_sheet.cell(row_index, column_index, value)

    if not export_rows:
        summary_sheet.cell(data_start, 1, "尚未配置评价人")
    data_end = max(data_start, data_start + len(export_rows) - 1)

    for cell in summary_sheet[1]:
        cell.fill = PatternFill("solid", fgColor=navy); cell.font = Font(color=white, bold=True, size=16)
    summary_sheet[1][0].alignment = Alignment(vertical="center")
    for cell in summary_sheet[2]:
        cell.fill = PatternFill("solid", fgColor="EAF1F8"); cell.font = Font(color="52657A", italic=True, size=10)
    for cell in summary_sheet[4]:
        cell.fill = PatternFill("solid", fgColor=blue); cell.font = Font(color=white, bold=True)
    for row_index in (5, 6):
        for column_index in range(1, 9):
            cell = summary_sheet.cell(row_index, column_index); cell.border = border
            if column_index % 2 == 1:
                cell.fill = PatternFill("solid", fgColor=pale_gray); cell.font = Font(bold=True, color=text)
    for cell in summary_sheet[8]:
        cell.fill = PatternFill("solid", fgColor=navy); cell.font = Font(color=white, bold=True); cell.alignment = Alignment(horizontal="center", vertical="center"); cell.border = Border(left=Side(style="thin", color=white), right=Side(style="thin", color=white), top=Side(style="thin", color=white), bottom=Side(style="thin", color=white))
    for cell in summary_sheet[9]:
        cell.fill = PatternFill("solid", fgColor=pale_blue); cell.font = Font(bold=True, color=text, size=10); cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True); cell.border = border
    for cell in summary_sheet[10]:
        cell.fill = PatternFill("solid", fgColor="FAFBFC"); cell.font = Font(italic=True, color="6B7785", size=9); cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True); cell.border = border
    for row in summary_sheet.iter_rows(min_row=data_start, max_row=data_end, min_col=1, max_col=max_col):
        for cell in row:
            cell.border = border; cell.alignment = Alignment(vertical="top", wrap_text=True); cell.font = Font(color=text, size=10)
    summary_sheet.column_dimensions["A"].width = 14
    summary_sheet.column_dimensions["B"].width = 12
    summary_sheet.column_dimensions["C"].width = 12
    for column_index in range(4, max_col + 1):
        header = summary_sheet.cell(9, column_index).value or ""
        summary_sheet.column_dimensions[get_column_letter(column_index)].width = 28 if header in {"优势", "不足或待改进项", "评估意见补充"} else 23 if header == "提交时间" else 16
    summary_sheet.freeze_panes = "A11"
    summary_sheet.sheet_view.showGridLines = False
    summary_sheet.row_dimensions[1].height = 32
    summary_sheet.row_dimensions[2].height = 24
    summary_sheet.row_dimensions[9].height = 42
    summary_sheet.row_dimensions[10].height = 28
    summary_sheet.auto_filter.ref = f"A9:{get_column_letter(max_col)}{data_end}"
    for row_index in range(data_start, data_end + 1):
        summary_sheet.cell(row_index, 3).number_format = "0%"
        for column_index in range(4, max_col + 1):
            if isinstance(summary_sheet.cell(row_index, column_index).value, (int, float)):
                summary_sheet.cell(row_index, column_index).number_format = "0.00"

    detail_headers = ["被评价人", "申报职级", "职类", "评价人", "评价人权重", "板块", "类别", "评分维度", "满分", "维度分数权重", "原始评分", "折算/计入分", "单项评价", "评审状态", "最后保存时间"]
    detail_sheet.append(["逐项评分明细（补充表）"])
    detail_sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(detail_headers))
    detail_sheet.append(["用于保留每一位评价人在每一个评分维度上的原始分、折算分和单项评价。"])
    detail_sheet.merge_cells(start_row=2, start_column=1, end_row=2, end_column=len(detail_headers))
    detail_sheet.append([])
    detail_sheet.append(detail_headers)
    for item in export_rows:
        assignment, review, scores = item["assignment"], item["review"], item["scores"]
        for module in modules:
            score = scores.get(module["id"], {})
            raw = score.get("score")
            converted = None if raw is None else (float(raw) if module["section_type"] == "bonus" else float(raw) / 5 * 100 * float(module["weight"] or 0))
            detail_sheet.append([
                submission["name"], submission["level"] or "", submission["category"] or "",
                assignment["reviewer_name"], float(assignment["weight"]), "团队成长贡献加分" if module["section_type"] == "bonus" else scoring_label,
                module["evaluation_item"] or "", module["module_name"] or "", 2 if module["section_type"] == "bonus" else 5,
                module["weight"] if module["section_type"] != "bonus" else None, raw, round(converted, 2) if converted is not None else None,
                score.get("comment") or "", item["status"], review["updated_at"] if review else "",
            ])
    for cell in detail_sheet[1]:
        cell.fill = PatternFill("solid", fgColor=navy); cell.font = Font(color=white, bold=True, size=16)
    for cell in detail_sheet[2]:
        cell.fill = PatternFill("solid", fgColor="EAF1F8"); cell.font = Font(color="52657A", italic=True)
    for cell in detail_sheet[4]:
        cell.fill = PatternFill("solid", fgColor=blue); cell.font = Font(color=white, bold=True); cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True); cell.border = border
    for row in detail_sheet.iter_rows(min_row=5, max_row=max(5, detail_sheet.max_row), min_col=1, max_col=len(detail_headers)):
        for cell in row:
            cell.border = border; cell.alignment = Alignment(vertical="top", wrap_text=True)
    for column_index in range(1, len(detail_headers) + 1):
        detail_sheet.column_dimensions[get_column_letter(column_index)].width = 30 if column_index == 13 else 23 if column_index == 15 else 16
    detail_sheet.freeze_panes = "A5"
    detail_sheet.sheet_view.showGridLines = False
    for row_index in range(5, detail_sheet.max_row + 1):
        detail_sheet.cell(row_index, 5).number_format = "0%"
        detail_sheet.cell(row_index, 10).number_format = "0%"
        for column_index in (9, 11, 12):
            detail_sheet.cell(row_index, column_index).number_format = "0.00"

    notes = [
        ["区域", "字段", "来源/计算逻辑", "说明"],
        ["基本信息", "姓名、申报职级、职类、申请时间", "来自被评价人申报信息", "与原始申请表的基本信息字段保持一致。"],
        ["评价权重", "评价人权重", "来自管理端评审配置", "评价人端不可见。"],
        [scoring_label, "各子项原始分", "评价人录入，单项满分 5 分", "子项列动态生成，包含所有普通评分维度。"],
        [scoring_label, scoring_label + "折算总分", "Σ[(原始分÷5)×100×子项权重]", "普通评分子项权重合计100%。"],
        ["成长贡献", "各加分项", "评价人录入，单项最多 2 分", "原始分直接计入，不进行百分制折算。"],
        ["分数汇总", "评价人总分", scoring_label + "折算总分＋成长贡献加分合计", "对应单个评价人的当前总分。"],
        ["分数汇总", "加权后总分", "评价人总分×评价人权重", "只有已完成的评价显示加权后总分。"],
        ["定性评价", "优势、不足、评估结论、补充意见", "来自评价人提交内容", "评估结论保留通过、不通过或其他。"],
    ]
    notes_sheet.append(["字段与计算说明"])
    notes_sheet.merge_cells("A1:D1")
    notes_sheet.append(["导出数据来源与当前系统计分规则。"])
    notes_sheet.merge_cells("A2:D2")
    notes_sheet.append([])
    for row in notes:
        notes_sheet.append(row)
    for cell in notes_sheet[1]:
        cell.fill = PatternFill("solid", fgColor=navy); cell.font = Font(color=white, bold=True, size=16)
    for cell in notes_sheet[2]:
        cell.fill = PatternFill("solid", fgColor="EAF1F8"); cell.font = Font(color="52657A", italic=True)
    for cell in notes_sheet[4]:
        cell.fill = PatternFill("solid", fgColor=blue); cell.font = Font(color=white, bold=True); cell.border = border
    for row in notes_sheet.iter_rows(min_row=5, max_row=notes_sheet.max_row, min_col=1, max_col=4):
        for cell in row:
            cell.border = border; cell.alignment = Alignment(vertical="top", wrap_text=True)
    for column, width in zip("ABCD", (15, 28, 42, 48)):
        notes_sheet.column_dimensions[column].width = width
    notes_sheet.freeze_panes = "A5"
    notes_sheet.sheet_view.showGridLines = False

    workbook.calculation.fullCalcOnLoad = True
    workbook.calculation.forceFullCalc = True
    output = io.BytesIO()
    workbook.save(output)
    safe_name = re.sub(r"[\\/:*?\"<>|]+", "_", submission["name"]).strip() or "被评价人"
    return output.getvalue(), f"{safe_name}-评委评分导出.xlsx"


def export_all_submission_reviews() -> tuple[bytes, str]:
    """Export every candidate shown on the dashboard as individual workbooks in one ZIP."""
    with db() as conn:
        submission_ids = [
            row["id"]
            for row in conn.execute("SELECT id FROM submissions ORDER BY imported_at DESC, name")
        ]
    if not submission_ids:
        raise ValueError("看板中暂无可导出的被评价人")

    output = io.BytesIO()
    used_names: set[str] = set()
    with zipfile.ZipFile(output, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        for submission_id in submission_ids:
            content, filename = export_submission_reviews(submission_id)
            candidate_name = filename
            suffix = 2
            while candidate_name.casefold() in used_names:
                stem, extension = os.path.splitext(filename)
                candidate_name = f"{stem}-{suffix}{extension}"
                suffix += 1
            used_names.add(candidate_name.casefold())
            archive.writestr(candidate_name, content)

    timestamp = datetime.now().astimezone().strftime("%Y%m%d-%H%M%S")
    return output.getvalue(), f"全部被评价人-评委评分导出-{timestamp}.zip"


def submission_detail(submission_id: str, assignment_id: str | None = None, *, include_workflow: bool = True) -> dict | None:
    with db() as conn:
        submission = conn.execute("SELECT * FROM submissions WHERE id=?", (submission_id,)).fetchone()
        if not submission:
            return None
        modules = [dict(row) for row in conn.execute("SELECT * FROM modules WHERE submission_id=? ORDER BY sort_order", (submission_id,))]
        attachments = [dict(row) for row in conn.execute("SELECT * FROM attachments WHERE submission_id=? ORDER BY sort_order", (submission_id,))]
        links = conn.execute("SELECT module_id, attachment_id, source FROM module_attachments WHERE module_id IN (SELECT id FROM modules WHERE submission_id=?)", (submission_id,)).fetchall()
        review = conn.execute("SELECT * FROM reviews WHERE submission_id=? AND assignment_id=? LIMIT 1", (submission_id, assignment_id)).fetchone() if assignment_id else conn.execute("SELECT * FROM reviews WHERE submission_id=? ORDER BY updated_at DESC LIMIT 1", (submission_id,)).fetchone()
        scores = [] if not review else [dict(row) for row in conn.execute("SELECT * FROM module_scores WHERE review_id=?", (review["id"],))]
        score_summary = None if not review else calculate_review_summary(conn, review["id"])
        if include_workflow:
            evidence_info=evidence.info(conn,submission_id)
        else:
            # Reviewers see finalized material, not the applicant/admin conversation or change history.
            task=conn.execute('SELECT finalized_at FROM evidence_tasks WHERE submission_id=?',(submission_id,)).fetchone()
            evidence_info={'finalizedAt':task['finalized_at']} if task else None
    link_map = {}
    for row in links:
        link_map.setdefault(row["module_id"], []).append(row["attachment_id"])
    for module in modules:
        module["attachment_ids"] = link_map.get(module["id"], [])
    return {"submission": dict(submission), "modules": modules, "attachments": attachments, "review": dict(review) if review else None, "scores": scores, "scoreSummary": score_summary,"evidenceTask":evidence_info}


def save_review(submission_id: str, payload: dict, assignment_id: str | None = None) -> dict:
    with db() as guard_conn:evidence.ensure_finalized(guard_conn,submission_id)
    reviewer_name = clean(payload.get("reviewerName")) or "当前评委"
    status = "已提交" if payload.get("submit") else "草稿"
    with db() as conn:
        conn.execute('BEGIN IMMEDIATE')
        submission = conn.execute("SELECT standard_snapshot_json FROM submissions WHERE id=?", (submission_id,)).fetchone()
        if submission and json.loads(submission["standard_snapshot_json"] or "{}").get("matrix"):
            weights = [r["weight"] for r in conn.execute("SELECT weight FROM modules WHERE submission_id=? AND section_type IN ('standard','reference')", (submission_id,))]
            if weights and (any(w is None or not 0 < w <= 1 for w in weights) or abs(sum(w or 0 for w in weights)-1) > .001):
                raise ValueError("此分级标准尚未完整配置各维度的评分权重，请联系管理员在材料详情中配置后再保存或提交评审")
        existing = conn.execute("SELECT id FROM reviews WHERE submission_id=? AND reviewer_name=?", (submission_id, reviewer_name)).fetchone()
        review_id = existing["id"] if existing else uuid.uuid4().hex
        conn.execute("""INSERT INTO reviews(id,submission_id,reviewer_name,strengths,improvements,opinion,opinion_note,status,updated_at)
          VALUES(?,?,?,?,?,?,?,?,?) ON CONFLICT(submission_id,reviewer_name) DO UPDATE SET strengths=excluded.strengths,
          improvements=excluded.improvements,opinion=excluded.opinion,opinion_note=excluded.opinion_note,status=excluded.status,updated_at=excluded.updated_at""",
          (review_id, submission_id, reviewer_name, clean(payload.get("strengths")), clean(payload.get("improvements")), clean(payload.get("opinion")), clean(payload.get("opinionNote")), status, now_text()))
        if assignment_id:
            conn.execute("UPDATE reviews SET assignment_id=? WHERE id=?", (assignment_id, review_id))
        valid_modules = {row["id"]: row["section_type"] for row in conn.execute("SELECT id,section_type FROM modules WHERE submission_id=?", (submission_id,))}
        for item in payload.get("moduleScores", []):
            if item.get("moduleId") not in valid_modules:
                continue
            score = None if item.get("score") in (None, "") else float(item["score"])
            max_score = 2 if valid_modules[item["moduleId"]] == "bonus" else 5
            if score is not None and not 0 <= score <= max_score:
                raise ValueError(f"该项评分必须在 0 至 {max_score} 分之间")
            conn.execute("""INSERT INTO module_scores(review_id,module_id,score,comment) VALUES(?,?,?,?)
              ON CONFLICT(review_id,module_id) DO UPDATE SET score=excluded.score,comment=excluded.comment""",
              (review_id, item["moduleId"], score, clean(item.get("comment"))))
        if payload.get("submit") and submission and json.loads(submission["standard_snapshot_json"] or "{}").get("matrix"):
            missing = conn.execute("""SELECT m.module_name FROM modules m
                LEFT JOIN module_scores ms ON ms.module_id=m.id AND ms.review_id=?
                WHERE m.submission_id=? AND ms.score IS NULL
                  AND (TRIM(COALESCE(m.evidence,''))<>'' OR EXISTS
                       (SELECT 1 FROM module_attachments ma WHERE ma.module_id=m.id))
                ORDER BY m.sort_order""", (review_id, submission_id)).fetchall()
            if missing:
                raise ValueError("请为已举证的子项逐项评分：" + "、".join(row['module_name'] for row in missing))
        summary = calculate_review_summary(conn, review_id)
        if assignment_id:
            conn.execute("UPDATE review_assignments SET status=?,completed_at=? WHERE id=?", ("已完成" if payload.get("submit") else "评价中", now_text() if payload.get("submit") else None, assignment_id))
    return {"reviewId": review_id, "status": status, "scoreSummary": summary}


def save_assigned_review(token: str, payload: dict) -> dict:
    with db() as conn:
        assignment = conn.execute(
            """SELECT a.id,a.submission_id,r.name reviewer_name FROM review_assignments a
               JOIN reviewers r ON r.id=a.reviewer_id WHERE a.token=?""", (token,)
        ).fetchone()
    if not assignment:
        raise ValueError("评审入口不存在或尚未推送")
    payload = dict(payload)
    payload["reviewerName"] = assignment["reviewer_name"]
    return save_review(assignment["submission_id"], payload, assignment["id"])


def save_public_review(submission_id: str, payload: dict) -> dict:
    """Save an intranet self-service review and create its assignment on demand."""
    with db() as conn:
        if evidence.task(conn,submission_id):raise PermissionError('在线举证必须通过管理员分配的评审入口评分')
    reviewer_name = clean(payload.get("reviewerName"))
    if not reviewer_name:
        raise ValueError("请填写评价官姓名")
    with db() as conn:
        if not conn.execute("SELECT 1 FROM submissions WHERE id=?", (submission_id,)).fetchone():
            raise ValueError("被评价人不存在")
    assignment = upsert_assignment({
        "submissionId": submission_id,
        "reviewerName": reviewer_name,
        "weight": 1,
    })
    result = save_review(submission_id, payload, assignment["id"])
    audit_event(
        "临时内网评审提交" if payload.get("submit") else "临时内网评审草稿",
        "review",
        result["reviewId"],
        f"{reviewer_name}：{result['status']}",
        {"submissionId": submission_id, "total": result["scoreSummary"]["total"]},
        {"name": reviewer_name, "role": "评价官", "userId": "intranet-self-service"},
    )
    return result


class Handler(BaseHTTPRequestHandler):
    server_version = "CertificationOps/0.8.0"

    def log_message(self, fmt, *args):
        message = fmt % args
        message = re.sub(r"[a-f0-9]{24,}", "<redacted-id>", message, flags=re.I)
        message = re.sub(r"\?[^ ]+", "?<redacted>", message)
        APP_LOG.info("http client=%s %s", self.client_address[0], message)

    def current_actor(self) -> dict | None:
        if self.headers.get('X-Evidence-Session'):
            return evidence.session_actor(self.headers['X-Evidence-Session'])
        cookie = self.headers.get("Cookie", "")
        values = {}
        for part in cookie.split(";"):
            if "=" in part:
                key, value = part.strip().split("=", 1)
                values[key] = unquote(value)
        external_id = values.get("cert_mock_user", "")
        if self.client_address[0] not in ('127.0.0.1','::1'):
            return None  # Local mock identities are never remote login credentials.
        if not external_id and dingtalk_config()["mode"] == "mock":
            external_id = "mock-admin"
        return get_app_user(external_id) if external_id else None

    def require_admin(self) -> dict | None:
        actor = self.current_actor()
        if not actor or actor.get("role") != "管理员":
            self.json_response({"error": "仅管理员可以执行此操作"}, 403)
            return None
        return actor

    def require_user(self) -> dict | None:
        actor = self.current_actor()
        if not actor:
            self.json_response({"error": "请先登录"}, 401)
            return None
        return actor

    def json_response(self, payload, status=200, headers=None):
        data = json.dumps(payload, ensure_ascii=False).encode("utf-8")
        self.send_response(status)
        self.send_header("Content-Type", "application/json; charset=utf-8")
        self.send_header("Content-Length", str(len(data)))
        self.send_header('Cache-Control','no-store')
        self.send_header('X-Content-Type-Options','nosniff')
        for key, value in (headers or {}).items():
            self.send_header(key, value)
        self.end_headers()
        self.wfile.write(data)

    def read_json(self):
        length = int(self.headers.get("Content-Length", "0"))
        if length > MAX_BODY:
            raise ValueError(f"请求内容超过 {MAX_FILE_MB}MB")
        payload = json.loads(self.rfile.read(length).decode("utf-8"))
        safe_path = re.sub(r"[a-f0-9]{24,}", "<redacted-id>", urlparse(self.path).path, flags=re.I)
        APP_LOG.info("request_json method=%s path=%s content_length=%s data=%s", self.command, safe_path, length, safe_log_json(payload))
        return payload

    def do_GET(self):
        parsed_url = urlparse(self.path)
        path = unquote(parsed_url.path)
        if self.headers.get('X-Evidence-Session') and path.startswith('/api/'):
            if not self.current_actor():return self.json_response({'error':'登录已失效，请重新输入举证码'},401)
            if not (path.startswith('/api/evidence/') or path.startswith('/api/my/applications')):return self.json_response({'error':'无权访问该接口'},403)
        if evidence.http_get(self,path):return
        if personnel_http_get(self,path,parsed_url.query):return
        APP_LOG.info("request_start method=GET path=%s query=%s client=%s", path, safe_log_value(parsed_url.query, "query"), self.client_address[0])
        if path == "/api/auth/me":
            actor = self.current_actor()
            return self.json_response({"user": actor, "mode": dingtalk_config()["mode"], "identities": mock_identities() if dingtalk_config()["mode"] == "mock" else []})
        if path == "/api/my/reviews":
            actor = self.current_actor()
            if not actor:
                return self.json_response({"error": "请先登录"}, 401)
            return self.json_response({"items": my_review_tasks(actor["userId"]), "user": actor})
        if path == "/api/admin/standards":
            if not self.require_admin():
                return
            return self.json_response({"items": list_qualification_standards()})
        if path == "/api/admin/job-taxonomy":
            if not self.require_admin():
                return
            return self.json_response({"items": list_job_taxonomy()})
        if path == "/api/admin/qualification-template":
            if not self.require_admin():
                return
            return self.json_response(get_qualification_template())
        positions_match = re.fullmatch(r"/api/admin/job-taxonomy/([a-f0-9]+)/positions", path)
        if positions_match:
            if not self.require_admin():
                return
            try:
                return self.json_response({"positions":get_job_positions(positions_match.group(1))})
            except ValueError as exc:
                return self.json_response({"error":str(exc)},400)
        standard_export_match = re.fullmatch(r"/api/admin/job-taxonomy/([a-f0-9]+)/standard-export", path)
        if standard_export_match:
            if not self.require_admin():
                return
            try:
                content, filename = export_subclass_standard_matrix(standard_export_match.group(1))
            except ValueError as exc:
                return self.json_response({"error":str(exc)},404)
            return self.send_download(
                content,
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                filename,
            )
        standard_history_match = re.fullmatch(r"/api/admin/standards/([a-f0-9]+)/history", path)
        if standard_history_match:
            if not self.require_admin():
                return
            if not qualification_standard_detail(standard_history_match.group(1)):
                return self.json_response({"error": "资格标准不存在"}, 404)
            return self.json_response({"items": qualification_standard_history(standard_history_match.group(1))})
        admin_standard_match = re.fullmatch(r"/api/admin/standards/([a-f0-9]+)", path)
        if admin_standard_match:
            if not self.require_admin():
                return
            detail = qualification_standard_detail(admin_standard_match.group(1))
            return self.json_response(detail or {"error": "资格标准不存在"}, 200 if detail else 404)
        if path == "/api/standards":
            if not self.require_user():
                return
            return self.json_response({"items": list_qualification_standards(published_only=True)})
        if path == "/api/my/applications":
            actor = self.require_user()
            if not actor:
                return
            return self.json_response({"items": list_my_applications(actor), "user": actor})
        my_application_match = re.fullmatch(r"/api/my/applications/([a-f0-9]+)", path)
        if my_application_match:
            actor = self.require_user()
            if not actor:
                return
            try:
                require_application_access(my_application_match.group(1), actor)
                result=submission_detail(my_application_match.group(1))
                if actor.get('evidenceId'):result.update(review=None,scores=[],scoreSummary=None)
                return self.json_response(result)
            except PermissionError as exc:
                return self.json_response({"error": str(exc)}, 403)
            except ValueError as exc:
                return self.json_response({"error": str(exc)}, 404)
        if path == "/api/system/health":
            if not self.require_admin():
                return
            return self.json_response(system_health())
        if path == "/api/admin/diagnostics/download":
            if not self.require_admin():
                return
            content, filename = create_diagnostic_bundle()
            return self.send_download(content, "application/zip", filename)
        if path == "/api/admin/audit-logs":
            if not self.require_admin():
                return
            with db() as conn:
                rows = [dict(row) for row in conn.execute("SELECT * FROM audit_logs ORDER BY created_at DESC LIMIT 200")]
            return self.json_response({"items": rows})
        if path == "/api/integrations/dingtalk/status":
            return self.json_response(dingtalk_status())
        if path == "/api/integrations/dingtalk/contacts":
            query = ""
            if parsed_url.query:
                from urllib.parse import parse_qs
                query = parse_qs(parsed_url.query).get("q", [""])[0]
            return self.json_response({"items": list_directory_users(query), "integration": dingtalk_status()})
        if path == "/api/admin/notifications":
            if not self.require_admin():
                return
            with db() as conn:
                rows = [dict(row) for row in conn.execute("SELECT * FROM notification_logs ORDER BY created_at DESC LIMIT 100")]
            return self.json_response({"items": rows})
        if path == "/api/admin/assignments":
            if not self.require_admin():
                return
            return self.json_response(list_assignments())
        if path == "/api/admin/dashboard":
            if not self.require_admin():
                return
            return self.json_response(dashboard_data())
        if path == "/api/import/jobs/latest":
            if not self.require_admin():
                return
            jobs = sorted(STAGING_DIR.glob("*/job.json"), key=lambda item: item.stat().st_mtime, reverse=True)
            if not jobs:
                return self.json_response({"error": "暂无可恢复的解析任务"}, 404)
            return self.json_response(json.loads(jobs[0].read_text(encoding="utf-8")))
        import_job_match = re.fullmatch(r"/api/import/jobs/([a-f0-9]+)", path)
        if import_job_match:
            if not self.require_admin():
                return
            target = job_path(import_job_match.group(1))
            if not target.exists():
                return self.json_response({"error": "解析任务不存在"}, 404)
            return self.json_response(json.loads(target.read_text(encoding="utf-8")))
        admin_review_match = re.fullmatch(r"/api/admin/assignments/([a-f0-9]+)/review-detail", path)
        if admin_review_match:
            if not self.require_admin():
                return
            return self.json_response(admin_assignment_review_detail(admin_review_match.group(1)))
        export_match = re.fullmatch(r"/api/admin/submissions/([a-f0-9]+)/reviews-export", path)
        if export_match:
            if not self.require_admin():
                return
            try:
                content, filename = export_submission_reviews(export_match.group(1))
            except ValueError as exc:
                return self.json_response({"error": str(exc)}, 404)
            return self.send_download(
                content,
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                filename,
            )
        if path == "/api/admin/submissions/reviews-export-all":
            if not self.require_admin():
                return
            try:
                content, filename = export_all_submission_reviews()
            except ValueError as exc:
                return self.json_response({"error": str(exc)}, 404)
            return self.send_download(content, "application/zip", filename)
        if path == "/api/admin/assignments/import-template":
            if not self.require_admin():
                return
            content, filename = assignment_import_template()
            return self.send_download(content, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", filename)
        assigned_match = re.fullmatch(r"/api/review/([a-f0-9]+)", path)
        if assigned_match:
            detail = assigned_review_detail(assigned_match.group(1))
            return self.json_response(detail or {"error": "评审入口不存在或尚未推送"}, 200 if detail else 404)
        if path == "/api/submissions":
            if not self.require_admin():return
            return self.json_response({"items": list_submissions()})
        if path == "/api/public/submissions":
            items = list_submissions()
            with db() as conn:private_ids={r[0] for r in conn.execute('SELECT submission_id FROM evidence_tasks')}
            items=[item for item in items if item['id'] not in private_ids]
            return self.json_response({"items": [{
                "id": item["id"], "name": item["name"], "level": item["level"],
                "category": item["category"], "module_count": item["module_count"],
            } for item in items]})
        if path == "/api/public/reviewers":
            return self.json_response({"items": public_reviewer_directory()})
        public_match = re.fullmatch(r"/api/public/submissions/([a-f0-9]+)", path)
        if public_match:
            with db() as conn:
                if evidence.task(conn,public_match.group(1)):return self.json_response({'error':'请使用管理员分配的评审入口'},403)
            detail = submission_detail(public_match.group(1))
            return self.json_response(detail or {"error": "被评价人不存在"}, 200 if detail else 404)
        match = re.fullmatch(r"/api/submissions/([a-f0-9]+)", path)
        if match:
            if not self.require_admin():return
            detail = submission_detail(match.group(1))
            return self.json_response(detail or {"error": "记录不存在"}, 200 if detail else 404)
        if path.startswith("/media/"):
            from urllib.parse import parse_qs
            media_query = parse_qs(parsed_url.query)
            target = (DATA_DIR / path.removeprefix("/media/")).resolve()
            if DATA_DIR.resolve() not in target.parents or not target.is_file():
                return self.send_error(404)
            parts=target.relative_to(DATA_DIR.resolve()).parts
            with db() as conn:private_task=evidence.task(conn,parts[1]) if len(parts)>1 and parts[0]=='submissions' else None
            if private_task:
                review_token=media_query.get('review',[''])[0]
                if not evidence.media_access(private_task['submission_id'],self.current_actor(),review_token):return self.json_response({'error':'无权访问附件'},403)
                if target.suffix.lower() not in {'.png','.jpg','.jpeg','.gif','.webp','.bmp'} and not (target.suffix.lower()=='.pdf' and media_query.get('preview')==['1']):
                    return self.send_download(target.read_bytes(),'application/octet-stream',target.name)
            if media_query.get('download')==['1']:
                return self.send_download(target.read_bytes(),'application/octet-stream',target.name)
            return self.send_file(target)
        if path == "/" or not Path(path).suffix:
            return self.send_file(STATIC_DIR / "index.html")
        target = (STATIC_DIR / path.lstrip("/")).resolve()
        if STATIC_DIR.resolve() not in target.parents or not target.is_file():
            return self.send_error(404)
        return self.send_file(target)

    def do_POST(self):
        path = urlparse(self.path).path
        if self.headers.get('X-Evidence-Session') and path not in ('/api/evidence/login','/api/evidence/logout'):
            if not self.current_actor():return self.json_response({'error':'登录已失效，请重新输入举证码'},401)
            if not path.startswith('/api/my/applications/'):return self.json_response({'error':'无权访问该接口'},403)
        APP_LOG.info(
            "request_start method=POST path=%s client=%s content_length=%s content_type=%s metadata=%s",
            re.sub(r"[a-f0-9]{24,}", "<redacted-id>", path, flags=re.I), self.client_address[0],
            self.headers.get("Content-Length", "0"), self.headers.get("Content-Type", ""),
            safe_log_json({"filename": unquote(self.headers.get("X-Filename", "")), "project": unquote(self.headers.get("X-Project-Name", "")), "batch": self.headers.get("X-Batch-Id", "")}),
        )
        try:
            if path == "/api/import/batch":
                if not self.require_admin():
                    return
                return self.json_response({"batchId": create_import_batch()})
            application_upload_match = re.fullmatch(r"/api/my/applications/([a-f0-9]+)/attachments/upload", path)
            if application_upload_match:
                actor = self.require_user()
                if not actor:
                    return
                length = int(self.headers.get("Content-Length", "0"))
                if length > MAX_BODY:
                    raise ValueError(f"单个附件不能超过 {MAX_FILE_MB}MB")
                result = store_application_attachment(
                    application_upload_match.group(1), unquote(self.headers.get("X-Filename", "")), self.rfile, length, actor,
                    unquote(self.headers.get("X-Module-Id", "")), unquote(self.headers.get("X-Material-Name", "")), unquote(self.headers.get("X-Link-Url", "")),
                )
                audit_event("上传在线举证附件", "submission", application_upload_match.group(1), result["name"], {"size": result["size"]}, actor)
                return self.json_response(result)
            if path == "/api/import/upload":
                if not self.require_admin():
                    return
                length = int(self.headers.get("Content-Length", "0"))
                if length > MAX_BODY:
                    raise ValueError(f"单个文件不能超过 {MAX_FILE_MB}MB")
                result = store_batch_stream(self.headers.get("X-Batch-Id", ""), unquote(self.headers.get("X-Filename", "")), self.rfile, length)
                return self.json_response({"ok": True, **result})
            if path == "/api/admin/assignments/import":
                actor = self.require_admin()
                if not actor:
                    return
                length = int(self.headers.get("Content-Length", "0"))
                if length > 10 * 1024 * 1024:
                    raise ValueError("评价关系 Excel 不能超过 10MB")
                result = import_assignment_workbook(self.rfile.read(length))
                audit_event("批量导入评价关系", "review_assignment", "", f"新增 {result['created']} 条，更新 {result['updated']} 条", result, actor)
                return self.json_response(result, 200 if not result["errors"] else 400)
            manual_attachment_match = re.fullmatch(r"/api/admin/submissions/([a-f0-9]+)/attachments/upload", path)
            if manual_attachment_match:
                actor = self.require_admin()
                if not actor:
                    return
                length = int(self.headers.get("Content-Length", "0"))
                if length > MAX_BODY:
                    raise ValueError(f"单个附件不能超过 {MAX_FILE_MB}MB")
                result = store_manual_attachment(
                    manual_attachment_match.group(1), unquote(self.headers.get("X-Filename", "")),
                    self.rfile, length, unquote(self.headers.get("X-Project-Name", "补充附件")),
                )
                audit_event("上传补充附件", "submission", manual_attachment_match.group(1), result["name"], {"size": result["size"]}, actor)
                return self.json_response(result)
            body = self.read_json()
            if evidence.http_post(self,path,body):return
            if personnel_http_post(self,path,body):return
            matrix_weights_match = re.fullmatch(r"/api/admin/submissions/([a-f0-9]+)/matrix-weights", path)
            if matrix_weights_match:
                actor = self.require_admin()
                if not actor:
                    return
                result = save_matrix_review_weights(matrix_weights_match.group(1),body)
                audit_event("配置分级申请评分权重", "submission", matrix_weights_match.group(1), "", {},actor)
                return self.json_response(result)
            if path == "/api/admin/qualification-template":
                actor = self.require_admin()
                if not actor:
                    return
                result = save_qualification_template(body)
                audit_event("保存认证维度模板", "qualification_template", "1", "", {"revision":result["revision"]},actor)
                return self.json_response(result)
            positions_match = re.fullmatch(r"/api/admin/job-taxonomy/([a-f0-9]+)/positions", path)
            if positions_match:
                actor = self.require_admin()
                if not actor:
                    return
                result = save_job_positions(positions_match.group(1),body)
                audit_event("配置职级岗位名称", "job_taxonomy", positions_match.group(1), "", {},actor)
                return self.json_response(result)
            taxonomy_active_match = re.fullmatch(r"/api/admin/job-taxonomy/([a-f0-9]+)/toggle-active", path)
            if taxonomy_active_match:
                actor = self.require_admin()
                if not actor:
                    return
                node_id = taxonomy_active_match.group(1)
                result = set_job_taxonomy_active(node_id, bool(body.get("active")))
                audit_event("启用职族分类" if result["active"] else "停用职族分类", "job_taxonomy", node_id, "", {}, actor)
                return self.json_response(result)
            taxonomy_match = re.fullmatch(r"/api/admin/job-taxonomy(?:/([a-f0-9]+)/(save|delete))?", path)
            if taxonomy_match:
                actor = self.require_admin()
                if not actor:
                    return
                node_id, action = taxonomy_match.groups()
                if action == "delete":
                    delete_job_taxonomy(node_id)
                    result = {"ok": True}
                else:
                    result = save_job_taxonomy(body, node_id)
                audit_event("删除职族分类" if action == "delete" else "保存职族分类", "job_taxonomy", node_id or result["id"], "", {}, actor)
                return self.json_response(result)
            if path == "/api/admin/standards":
                actor = self.require_admin()
                if not actor:
                    return
                result = save_qualification_standard(body, actor)
                audit_event("新建资格标准", "qualification_standard", result["standard"]["id"], result["standard"]["name"], {"status": result["standard"]["status"]}, actor)
                return self.json_response(result)
            standard_iterate_match = re.fullmatch(r"/api/admin/standards/([a-f0-9]+)/iterate", path)
            if standard_iterate_match:
                actor = self.require_admin()
                if not actor:
                    return
                result = iterate_qualification_standard(standard_iterate_match.group(1), actor, body.get("changeNote"))
                audit_event("发起标准迭代", "qualification_standard", result["standard"]["id"], result["standard"]["name"], {"sourceId": standard_iterate_match.group(1)}, actor)
                return self.json_response(result)
            standard_save_match = re.fullmatch(r"/api/admin/standards/([a-f0-9]+)/save", path)
            if standard_save_match:
                actor = self.require_admin()
                if not actor:
                    return
                result = save_qualification_standard(body, actor, standard_save_match.group(1))
                audit_event("修改资格标准", "qualification_standard", standard_save_match.group(1), result["standard"]["name"], {"status": result["standard"]["status"]}, actor)
                return self.json_response(result)
            standard_copy_match = re.fullmatch(r"/api/admin/standards/copy-submission/([a-f0-9]+)", path)
            if standard_copy_match:
                actor = self.require_admin()
                if not actor:
                    return
                result = copy_submission_to_standard(standard_copy_match.group(1), actor)
                audit_event("从材料复制资格标准", "qualification_standard", result["standard"]["id"], result["standard"]["name"], {"submissionId": standard_copy_match.group(1)}, actor)
                return self.json_response(result)
            standard_delete_match = re.fullmatch(r"/api/admin/standards/([a-f0-9]+)/delete", path)
            if standard_delete_match:
                actor = self.require_admin()
                if not actor:
                    return
                delete_qualification_standard(standard_delete_match.group(1))
                audit_event("删除资格标准", "qualification_standard", standard_delete_match.group(1), "", {}, actor)
                return self.json_response({"ok": True})
            if path == "/api/my/applications":
                actor = self.require_user()
                if not actor:
                    return
                result = create_online_application(clean(body.get("standardId")), actor, body.get("grade"))
                audit_event("创建在线认证申请", "submission", result["submission"]["id"], result["submission"]["name"], {"standardId": body.get("standardId")}, actor)
                return self.json_response(result)
            application_save_match = re.fullmatch(r"/api/my/applications/([a-f0-9]+)/save", path)
            if application_save_match:
                actor = self.require_user()
                if not actor:
                    return
                result = save_online_application(application_save_match.group(1), body, actor)
                audit_event("提交在线认证申请" if body.get("submit") else "保存在线认证草稿", "submission", application_save_match.group(1), result["submission"]["name"], {}, actor)
                return self.json_response(result)
            application_attachment_assign_match = re.fullmatch(r"/api/my/applications/([a-f0-9]+)/attachments/([a-f0-9]+)/assign", path)
            if application_attachment_assign_match:
                actor = self.require_user()
                if not actor:
                    return
                reassign_application_attachment(application_attachment_assign_match.group(1), application_attachment_assign_match.group(2), clean(body.get('moduleId')), actor)
                return self.json_response({'ok': True})
            application_attachment_delete_match = re.fullmatch(r"/api/my/applications/([a-f0-9]+)/attachments/([a-f0-9]+)/delete", path)
            if application_attachment_delete_match:
                actor = self.require_user()
                if not actor:
                    return
                delete_application_attachment(application_attachment_delete_match.group(1), application_attachment_delete_match.group(2), actor)
                audit_event("删除在线举证附件", "submission", application_attachment_delete_match.group(1), "", {}, actor)
                return self.json_response({"ok": True})
            application_link_match = re.fullmatch(r"/api/my/applications/([a-f0-9]+)/attachments/link", path)
            if application_link_match:
                actor = self.require_user()
                if not actor:
                    return
                result = store_application_link(application_link_match.group(1), body, actor)
                audit_event("添加在线举证链接", "submission", application_link_match.group(1), result["name"], {}, actor)
                return self.json_response(result)
            reviewer_login_match = re.fullmatch(r"/api/public/reviewers/([a-f0-9]+)/login", path)
            if reviewer_login_match:
                return self.json_response(verify_reviewer_access(reviewer_login_match.group(1), body.get("code")))
            reviewer_tasks_match = re.fullmatch(r"/api/public/reviewers/([a-f0-9]+)/tasks", path)
            if reviewer_tasks_match:
                return self.json_response({"items": authorized_reviewer_tasks(reviewer_tasks_match.group(1), clean(body.get("sessionToken")))})
            if path == "/api/auth/mock-login":
                if self.client_address[0] not in ('127.0.0.1','::1'):raise PermissionError('模拟身份仅限本机开发使用')
                if dingtalk_config()["mode"] != "mock":
                    raise ValueError("仅模拟模式允许切换身份")
                user_id = clean(body.get("userId"))
                actor = get_app_user(user_id)
                if not actor:
                    raise ValueError("模拟用户不存在")
                audit_event("切换模拟身份", "user", user_id, f"切换为 {actor['name']}（{actor['role']}）", {}, actor)
                return self.json_response({"user": actor}, headers={"Set-Cookie": f"cert_mock_user={user_id}; Path=/; SameSite=Lax"})
            if path == "/api/import/parse":
                if not self.require_admin():
                    return
                return self.json_response(start_parse_job(body.get("batchId", "")), 202)
            if path == "/api/import/preview":
                if not self.require_admin():
                    return
                return self.json_response(save_preview(body.get("files", [])))
            if path == "/api/import/confirm":
                actor = self.require_admin()
                if not actor:
                    return
                created = confirm_batch(body.get("batchId", ""), body.get("indexes"))
                audit_event("批量导入", "submission", "", f"导入 {len(created)} 份认证材料", {"created": created}, actor)
                return self.json_response({"created": created})
            if path == "/api/admin/assignments":
                actor = self.require_admin()
                if not actor:
                    return
                result = upsert_assignment(body)
                audit_event("配置评委", "review_assignment", result["id"], f"配置评委 {clean(body.get('reviewerName'))}", {"weight": body.get("weight")}, actor)
                return self.json_response(result)
            submission_update_match = re.fullmatch(r"/api/admin/submissions/([a-f0-9]+)/update", path)
            if submission_update_match:
                actor = self.require_admin()
                if not actor:
                    return
                update_submission_basic(submission_update_match.group(1), body)
                audit_event("修改被评价人", "submission", submission_update_match.group(1), clean(body.get("name")), {}, actor)
                return self.json_response({"ok": True})
            submission_delete_match = re.fullmatch(r"/api/admin/submissions/([a-f0-9]+)/delete", path)
            if submission_delete_match:
                actor = self.require_admin()
                if not actor:
                    return
                delete_submission(submission_delete_match.group(1))
                audit_event("删除被评价人", "submission", submission_delete_match.group(1), "删除材料、评价关系与评分", {}, actor)
                return self.json_response({"ok": True})
            assignment_update_match = re.fullmatch(r"/api/admin/assignments/([a-f0-9]+)/update", path)
            if assignment_update_match:
                actor = self.require_admin()
                if not actor:
                    return
                update_assignment(assignment_update_match.group(1), body)
                audit_event("修改评价配置", "review_assignment", assignment_update_match.group(1), clean(body.get("reviewerName")), {"weight": body.get("weight")}, actor)
                return self.json_response({"ok": True})
            reviewer_code_match = re.fullmatch(r"/api/admin/reviewers/([a-f0-9]+)/access-code", path)
            if reviewer_code_match:
                actor = self.require_admin()
                if not actor:
                    return
                set_reviewer_access_code(reviewer_code_match.group(1), body.get("code"))
                audit_event("设置评价人访问码", "reviewer", reviewer_code_match.group(1), "访问码已设置或重置", {}, actor)
                return self.json_response({"ok": True})
            if path == "/api/admin/reviewers":
                actor = self.require_admin()
                if not actor:
                    return
                reviewer = create_reviewer(body.get("name"))
                audit_event("添加评价人", "reviewer", reviewer["id"], reviewer["name"], {}, actor)
                return self.json_response(reviewer)
            reviewer_delete_match = re.fullmatch(r"/api/admin/reviewers/([a-f0-9]+)/delete", path)
            if reviewer_delete_match:
                actor = self.require_admin()
                if not actor:
                    return
                delete_reviewer(reviewer_delete_match.group(1))
                audit_event("删除评价人", "reviewer", reviewer_delete_match.group(1), "删除评价人及其评价关系", {}, actor)
                return self.json_response({"ok": True})
            assignment_delete_match = re.fullmatch(r"/api/admin/assignments/([a-f0-9]+)/delete", path)
            if assignment_delete_match:
                actor = self.require_admin()
                if not actor:
                    return
                delete_assignment(assignment_delete_match.group(1))
                audit_event("删除评价配置", "review_assignment", assignment_delete_match.group(1), "删除评价关系及其评分", {}, actor)
                return self.json_response({"ok": True})
            push_match = re.fullmatch(r"/api/admin/assignments/([a-f0-9]+)/push", path)
            if push_match:
                actor = self.require_admin()
                if not actor:
                    return
                result = push_assignment(push_match.group(1))
                audit_event("推送评审", "review_assignment", push_match.group(1), "生成评审入口并发送提醒", {"notificationStatus": result["notification"]["status"]}, actor)
                return self.json_response(result)
            retry_match = re.fullmatch(r"/api/admin/notifications/([a-f0-9]+)/retry", path)
            if retry_match:
                actor = self.require_admin()
                if not actor:
                    return
                result = retry_notification(retry_match.group(1))
                audit_event("重发提醒", "notification", retry_match.group(1), "重新发送评审提醒", {"status": result["status"]}, actor)
                return self.json_response(result)
            assigned_save_match = re.fullmatch(r"/api/review/([a-f0-9]+)/save", path)
            if assigned_save_match:
                result = save_assigned_review(assigned_save_match.group(1), body)
                actor = self.current_actor()
                audit_event("提交评审" if body.get("submit") else "保存评审草稿", "review", result["reviewId"], f"评审状态：{result['status']}", {"total": result["scoreSummary"]["total"]}, actor)
                return self.json_response(result)
            public_save_match = re.fullmatch(r"/api/public/submissions/([a-f0-9]+)/review", path)
            if public_save_match:
                return self.json_response(save_public_review(public_save_match.group(1), body))
            review_match = re.fullmatch(r"/api/submissions/([a-f0-9]+)/review", path)
            if review_match:
                actor = self.require_admin()
                if not actor:
                    return
                result = save_review(review_match.group(1), body)
                audit_event("管理员记录评审", "review", result["reviewId"], f"评审状态：{result['status']}", {"total": result["scoreSummary"]["total"]}, actor)
                return self.json_response(result)
            return self.json_response({"error": "接口不存在"}, 404)
        except ValueError as exc:
            return self.json_response({"error": str(exc)}, 400)
        except PermissionError as exc:
            return self.json_response({"error": str(exc)}, 403)
        except Exception as exc:
            safe_path = re.sub(r"[a-f0-9]{24,}", "<redacted-id>", path, flags=re.I)
            APP_LOG.exception("Unhandled request error method=POST path=%s error_type=%s", safe_path, type(exc).__name__)
            return self.json_response({"error": "服务器处理失败", "detail": str(exc)}, 500)

    def send_file(self, target: Path):
        data = target.read_bytes()
        mime = mimetypes.guess_type(target.name)[0] or "application/octet-stream"
        self.send_response(200)
        self.send_header("Content-Type", mime)
        self.send_header("Content-Length", str(len(data)))
        self.send_header("Cache-Control", "no-cache")
        self.send_header('X-Content-Type-Options','nosniff')
        self.send_header('Referrer-Policy','no-referrer')
        self.end_headers()
        self.wfile.write(data)

    def send_download(self, data: bytes, mime: str, filename: str):
        self.send_response(200)
        self.send_header("Content-Type", mime)
        self.send_header("Content-Length", str(len(data)))
        self.send_header("Content-Disposition", f"attachment; filename*=UTF-8''{quote(filename)}")
        self.send_header("Cache-Control", "no-store")
        self.end_headers()
        self.wfile.write(data)


def init_personnel_storage(conn):
    conn.executescript('''
        CREATE TABLE IF NOT EXISTS personnel (
          id TEXT PRIMARY KEY, employee_no TEXT UNIQUE, name TEXT NOT NULL, department TEXT NOT NULL,
          initial_level TEXT NOT NULL, initial_date TEXT NOT NULL, note TEXT NOT NULL,
          revision INTEGER NOT NULL DEFAULT 0, created_at TEXT NOT NULL, updated_at TEXT NOT NULL);
        CREATE TABLE IF NOT EXISTS personnel_certifications (
          id TEXT PRIMARY KEY, person_id TEXT NOT NULL REFERENCES personnel(id),
          submission_id TEXT UNIQUE REFERENCES submissions(id) ON DELETE SET NULL,
          source_json TEXT NOT NULL, outcome TEXT NOT NULL DEFAULT '待确认',
          decided_on TEXT NOT NULL DEFAULT '', effective_on TEXT NOT NULL DEFAULT '',
          result_level TEXT NOT NULL DEFAULT '', promoted INTEGER NOT NULL DEFAULT 0,
          score_json TEXT NOT NULL DEFAULT '{}', note TEXT NOT NULL DEFAULT '',
          revision INTEGER NOT NULL DEFAULT 0, created_at TEXT NOT NULL, updated_at TEXT NOT NULL);
        CREATE INDEX IF NOT EXISTS personnel_certifications_person ON personnel_certifications(person_id);
        CREATE TABLE IF NOT EXISTS personnel_history (
          id INTEGER PRIMARY KEY AUTOINCREMENT, person_id TEXT NOT NULL,
          record_id TEXT, action TEXT NOT NULL, actor_id TEXT, actor_name TEXT,
          created_at TEXT NOT NULL, note TEXT NOT NULL, snapshot_json TEXT NOT NULL);
    ''')


def personnel_text(payload, key, limit=200):
    value=clean(payload.get(key))
    if len(value)>limit:raise ValueError(f'{key}内容过长')
    return value


def personnel_date(value, required=False):
    value=clean(value)
    if not value and not required:return ''
    try:
        if datetime.strptime(value,'%Y-%m-%d').strftime('%Y-%m-%d')!=value:raise ValueError()
    except ValueError:raise ValueError('请填写有效日期，格式为 YYYY-MM-DD')
    return value


def personnel_grade_rank(value):
    if value in GRADE_NAMES:return ('级',GRADE_NAMES.index(value)+1)
    match=re.fullmatch(r'([A-Za-z]*)(\d+)(级)?',value)
    if match:return (match[1].upper() or '级',int(match[2]))
    return None


def personnel_event(conn, person_id, record_id, action, actor, note=''):
    person=dict(conn.execute('SELECT * FROM personnel WHERE id=?',(person_id,)).fetchone())
    records=[dict(r) for r in conn.execute('SELECT * FROM personnel_certifications WHERE person_id=? ORDER BY created_at,id',(person_id,))]
    conn.execute('INSERT INTO personnel_history(person_id,record_id,action,actor_id,actor_name,created_at,note,snapshot_json) VALUES(?,?,?,?,?,?,?,?)',
                 (person_id,record_id,action,actor.get('userId'),actor.get('name'),now_text(),note,json.dumps({'person':person,'records':records},ensure_ascii=False)))


def save_personnel(payload, actor, person_id=None):
    name=personnel_text(payload,'name');number=personnel_text(payload,'employeeNo',80).upper()
    if not name:raise ValueError('请填写人员姓名')
    values=(number or None,name,personnel_text(payload,'department'),personnel_text(payload,'initialLevel',80),
            personnel_date(payload.get('initialDate')),personnel_text(payload,'note',3000))
    if values[4] and values[4]>now_text()[:10]:raise ValueError('初始职级日期不能晚于今天')
    with db() as conn:
        conn.execute('BEGIN IMMEDIATE')
        if number and conn.execute('SELECT id FROM personnel WHERE UPPER(employee_no)=? AND id<>?',(number,person_id or '')).fetchone():
            raise ValueError('该工号已存在，请关联已有人员档案')
        if person_id:
            old=conn.execute('SELECT * FROM personnel WHERE id=?',(person_id,)).fetchone()
            if not old:raise ValueError('人员不存在')
            if payload.get('revision')!=old['revision']:raise ValueError('档案已更新，请重新打开')
            if conn.execute("SELECT 1 FROM personnel_certifications WHERE person_id=? AND outcome<>'待确认'",(person_id,)).fetchone() and (old['initial_level']!=values[3] or old['initial_date']!=values[4]):
                raise ValueError('已有确认的认证记录，初始职级及日期已锁定')
            conn.execute('UPDATE personnel SET employee_no=?,name=?,department=?,initial_level=?,initial_date=?,note=?,revision=revision+1,updated_at=? WHERE id=?',(*values,now_text(),person_id))
        else:
            person_id=uuid.uuid4().hex
            conn.execute('INSERT INTO personnel VALUES(?,?,?,?,?,?,?,0,?,?)',(person_id,*values,now_text(),now_text()))
        personnel_event(conn,person_id,None,'维护人员档案',actor)
    return {'id':person_id}


def provision_evidence_personnel(conn, submission_id, name, config):
    number=clean(config.get('employeeNo')).upper();current=clean(config.get('currentLevel'))
    if not number or not current:raise ValueError('配置举证任务时必须填写工号和当前职级')
    matches=conn.execute('SELECT * FROM personnel WHERE UPPER(employee_no)=?',(number,)).fetchall()
    if len(matches)>1:raise ValueError('该工号存在多份历史档案，请先核对工号后再配置')
    actor={'userId':config.get('createdBy'),'name':config.get('createdByName')}
    if matches:
        person=matches[0];person_id=person['id']
        if person['name']!=name:raise ValueError('该工号已关联其他姓名，请核对工号或先维护人员档案中的姓名')
        level=person['initial_level']
        for row in conn.execute("SELECT result_level FROM personnel_certifications WHERE person_id=? AND outcome='通过' AND effective_on<=? ORDER BY effective_on,decided_on,created_at,id",(person_id,now_text()[:10])):
            level=row['result_level'] or level
        if level and level!=current:raise ValueError(f'该工号档案中的当前职级为{level}，与填写内容不一致，请核对或先补录已生效的职级变迁')
        if not level:
            conn.execute('UPDATE personnel SET initial_level=?,revision=revision+1,updated_at=? WHERE id=?',(current,now_text(),person_id))
            personnel_event(conn,person_id,None,'补充初始职级',actor)
    else:
        person_id=uuid.uuid4().hex
        conn.execute('INSERT INTO personnel VALUES(?,?,?,?,?,?,?,0,?,?)',(person_id,number,name,'',current,'','由首次举证配置同步建档',now_text(),now_text()))
        personnel_event(conn,person_id,None,'举证配置同步建档',actor)
    source=dict(conn.execute('SELECT id,name,level,category,applied_at,standard_snapshot_json FROM submissions WHERE id=?',(submission_id,)).fetchone())
    source.update({'origin':'系统认证','employeeNo':number,'currentLevel':current})
    record_id=uuid.uuid4().hex
    conn.execute('INSERT INTO personnel_certifications(id,person_id,submission_id,source_json,created_at,updated_at) VALUES(?,?,?,?,?,?)',
                 (record_id,person_id,submission_id,json.dumps(source,ensure_ascii=False),now_text(),now_text()))
    personnel_event(conn,person_id,record_id,'举证配置自动关联认证',actor)


def link_personnel_submission(person_id, submission_id, actor):
    with db() as conn:
        conn.execute('BEGIN IMMEDIATE')
        if not conn.execute('SELECT 1 FROM personnel WHERE id=?',(person_id,)).fetchone():raise ValueError('人员不存在')
        source=conn.execute('SELECT id,name,level,category,applied_at,standard_snapshot_json FROM submissions WHERE id=?',(submission_id,)).fetchone()
        if not source:raise ValueError('认证材料不存在')
        if conn.execute('SELECT 1 FROM personnel_certifications WHERE submission_id=?',(submission_id,)).fetchone():raise ValueError('该材料已关联人员，请勿重复关联')
        source=dict(source);source['origin']='系统认证'
        record_id=uuid.uuid4().hex
        conn.execute('INSERT INTO personnel_certifications(id,person_id,submission_id,source_json,created_at,updated_at) VALUES(?,?,?,?,?,?)',
                     (record_id,person_id,submission_id,json.dumps(source,ensure_ascii=False),now_text(),now_text()))
        personnel_event(conn,person_id,record_id,'关联认证材料',actor)
    return {'id':record_id}


def personnel_data(person_id=None):
    results={r['submissionId']:r for r in dashboard_data()['items']}
    with db() as conn:
        people=[dict(r) for r in conn.execute('SELECT * FROM personnel ORDER BY name,id')]
        records=[dict(r) for r in conn.execute('SELECT * FROM personnel_certifications ORDER BY decided_on,created_at,id')]
        unlinked=[dict(r) for r in conn.execute('''SELECT id,name,level,category,applied_at,status FROM submissions s
            WHERE NOT EXISTS (SELECT 1 FROM personnel_certifications c WHERE c.submission_id=s.id) ORDER BY imported_at DESC''')]
        history=[dict(r) for r in conn.execute('SELECT id,record_id,action,actor_name,created_at,note FROM personnel_history WHERE person_id=? ORDER BY id DESC',(person_id,))] if person_id else []
    today=now_text()[:10]
    for row in records:
        row['source']=json.loads(row.pop('source_json'));stored=json.loads(row.pop('score_json'))
        standard=json.loads(row['source'].pop('standard_snapshot_json','{}') or '{}')
        row['source'].update({'standardName':standard.get('name',''),'standardVersion':standard.get('version_number','')})
        row['scoreData']=stored if row['outcome']!='待确认' else results.get(row['submission_id'],{})
        row['reviewReady']=bool(row['scoreData'].get('status')=='已完成' and abs(row['scoreData'].get('weightTotal',0)-1)<=.001)
    for person in people:
        own=[r for r in records if r['person_id']==person['id']]
        own.sort(key=lambda r:(r['effective_on'] or r['decided_on'] or '9999',r['decided_on'],r['created_at'],r['id']))
        level=person['initial_level'];projected=level;promotions=0
        for row in own:
            row['beforeLevel']=projected
            effective=row['outcome']=='通过' and row['effective_on'] and row['effective_on']<=today
            if row['outcome']=='通过':projected=row['result_level'] or projected
            if effective:
                level=row['result_level'] or level
                promotions+=bool(row['promoted'])
            row['afterLevel']=level
            row['effectStatus']='已生效' if effective else '待生效' if row['outcome']=='通过' else '未通过' if row['outcome']=='未通过' else '待确认'
        person.update({'currentLevel':level,'certificationCount':len(own),'passedCount':sum(r['outcome']=='通过' for r in own),
                       'failedCount':sum(r['outcome']=='未通过' for r in own),'pendingCount':sum(r['outcome']=='待确认' for r in own),
                       'promotionCount':promotions,'records':own})
    if person_id:
        person=next((p for p in people if p['id']==person_id),None)
        if not person:raise ValueError('人员不存在')
        return {'person':person,'history':history}
    return {'items':people,'unlinked':unlinked,'asOf':today}


def save_personnel_certification(person_id, payload, actor, record_id=None):
    outcome=clean(payload.get('outcome'))
    if outcome not in {'通过','未通过'}:raise ValueError('请选择通过或未通过的正式认证结论')
    decided=personnel_date(payload.get('decidedOn'),True)
    if decided>now_text()[:10]:raise ValueError('认证结论日期不能晚于今天')
    effective=personnel_date(payload.get('effectiveOn'),outcome=='通过') if outcome=='通过' else ''
    target=personnel_text(payload,'resultLevel',80) if outcome=='通过' else ''
    promoted=payload.get('promoted') is True
    note=personnel_text(payload,'note',3000)
    if outcome=='通过' and not target:raise ValueError('请填写认证通过后的职级；不晋级请填写保持的职级')
    if effective and effective<decided:raise ValueError('职级生效日期不能早于认证结论日期')
    if outcome!='通过' and promoted:raise ValueError('未通过认证不能记为晋级')
    with db() as conn:
        conn.execute('BEGIN IMMEDIATE')
        person=conn.execute('SELECT * FROM personnel WHERE id=?',(person_id,)).fetchone()
        if not person:raise ValueError('人员不存在')
        if not person['initial_level']:raise ValueError('请先维护人员的初始职级，再确认认证结论')
        if person['initial_date'] and decided<person['initial_date']:raise ValueError('认证记录不能早于初始职级日期')
        old=conn.execute('SELECT * FROM personnel_certifications WHERE id=? AND person_id=?',(record_id,person_id)).fetchone() if record_id else None
        if record_id and not old:raise ValueError('认证记录不存在')
        if old and payload.get('revision')!=old['revision']:raise ValueError('认证记录已更新，请重新打开')
        if old and old['outcome']!='待确认' and not note:raise ValueError('修订已确认记录时请填写修订说明')
        source=json.loads(old['source_json']) if old else {'origin':'历史补录','level':personnel_text(payload,'appliedLevel',80),'category':personnel_text(payload,'category')}
        live={r['submissionId']:r for r in dashboard_data()['items']} if old and old['submission_id'] and old['outcome']=='待确认' else {}
        score=live.get(old['submission_id']) if old and old['submission_id'] else None
        if old and old['outcome']=='待确认' and old['submission_id'] and (not score or score['status']!='已完成' or abs(score['weightTotal']-1)>.001):
            raise ValueError('请先完成全部评委评审，且评委汇总权重合计100%，再确认认证结论')
        if old and old['outcome']=='待确认' and old['submission_id']:ensure_assignment_ready(conn,old['submission_id'])
        if old and old['outcome']!='待确认':score=json.loads(old['score_json'])
        if not old:
            if not note:raise ValueError('历史补录请填写记录来源或说明')
            manual=payload.get('manualScore')
            try:manual=float(manual) if manual not in (None,'') else None
            except (ValueError,TypeError):raise ValueError('历史认证分数必须是数字')
            if manual is not None and not 0<=manual<=1000:raise ValueError('历史认证分数须在0至1000之间')
            score={'status':'历史补录','weightedTotal':manual,'reviewers':[]}
        if not score:raise ValueError('认证数据已缺失，无法确认；请核对原记录')
        record_id=record_id or uuid.uuid4().hex
        if old:
            conn.execute('UPDATE personnel_certifications SET outcome=?,decided_on=?,effective_on=?,result_level=?,promoted=?,score_json=?,note=?,revision=revision+1,updated_at=? WHERE id=?',
                         (outcome,decided,effective,target,int(promoted),json.dumps(score,ensure_ascii=False),note,now_text(),record_id))
        else:
            conn.execute('''INSERT INTO personnel_certifications(id,person_id,source_json,outcome,decided_on,effective_on,result_level,promoted,score_json,note,created_at,updated_at)
                VALUES(?,?,?,?,?,?,?,?,?,?,?,?)''',(record_id,person_id,json.dumps(source,ensure_ascii=False),outcome,decided,effective,target,int(promoted),json.dumps(score,ensure_ascii=False),note,now_text(),now_text()))
        # Replay the chronological grade chain, rather than treating application grades as promotions.
        level=person['initial_level']
        rows=conn.execute("SELECT * FROM personnel_certifications WHERE person_id=? AND outcome='通过' ORDER BY effective_on,decided_on,created_at,id",(person_id,)).fetchall()
        for row in rows:
            if row['promoted'] and (not level or level==row['result_level']):raise ValueError('晋级记录需要已知的原职级，且晋级后职级应与原职级不同')
            previous_rank,next_rank=personnel_grade_rank(level),personnel_grade_rank(row['result_level'])
            if row['promoted'] and previous_rank and next_rank and previous_rank[0]==next_rank[0] and next_rank[1]<=previous_rank[1]:
                raise ValueError('晋级后的职级必须高于此前职级')
            if not row['promoted'] and level and level!=row['result_level']:raise ValueError('保持职级的记录应与此前职级一致；如职级发生变化，请勾选晋级')
            level=row['result_level']
        personnel_event(conn,person_id,record_id,'确认认证结论' if not old or old['outcome']=='待确认' else '修订认证结论',actor,note)
    return {'id':record_id}


def export_personnel(person_id=None, query='', level=''):
    data=personnel_data(person_id)
    people=[data['person']] if person_id else data['items']
    people=[p for p in people if (not query or query.lower() in ' '.join([p['name'],p['employee_no'] or '',p['department']]).lower()) and (not level or p['currentLevel']==level)]
    workbook=Workbook();overview=workbook.active;overview.title='人员职级清单'
    overview.append(['工号','姓名','部门','初始职级','初始职级日期','当前职级','认证次数','通过次数','未通过次数','待确认次数','已生效晋级次数','统计日期'])
    history=workbook.create_sheet('认证与职级变迁')
    history.append(['工号','姓名','来源','认证材料编号','职类','申报职级','认证结论','结论日期','生效日期','原职级','认证后职级','是否晋级','生效状态','加权总分','评委完成数','评委总数','评委权重合计','说明','标准名称','标准版本'])
    reviews=workbook.create_sheet('评委评分明细');reviews.append(['工号','姓名','认证材料编号','评委','进度','评委权重','评委总分','加权贡献'])
    for person in people:
        overview.append([person['employee_no'],person['name'],person['department'],person['initial_level'],person['initial_date'],person['currentLevel'],person['certificationCount'],person['passedCount'],person['failedCount'],person['pendingCount'],person['promotionCount'],now_text()[:10]])
        for record in person['records']:
            source=record['source'];score=record['scoreData']
            history.append([person['employee_no'],person['name'],source.get('origin'),source.get('id'),source.get('category'),source.get('level'),record['outcome'],record['decided_on'],record['effective_on'],record['beforeLevel'],record['result_level'] if record['outcome']=='通过' else record['beforeLevel'], '是' if record['promoted'] else '否',record['effectStatus'],score.get('weightedTotal'),score.get('completedCount'),score.get('assignmentCount'),score.get('weightTotal'),record['note'],source.get('standardName'),source.get('standardVersion')])
            for reviewer in score.get('reviewers',[]):
                reviews.append([person['employee_no'],person['name'],source.get('id'),reviewer.get('reviewerName'),reviewer.get('status'),reviewer.get('weight'),reviewer.get('score'),reviewer.get('weightedContribution')])
    if not person_id:
        pending=workbook.create_sheet('待关联认证材料');pending.append(['姓名','申报职级','职类','申请日期','材料状态','材料编号'])
        for row in data['unlinked']:
            if not level and (not query or query.lower() in (row['name']+' '+(row['category'] or '')).lower()):
                pending.append([row['name'],row['level'],row['category'],row['applied_at'],row['status'],row['id']])
    for sheet in workbook:
        sheet.freeze_panes='A2';sheet.auto_filter.ref=sheet.dimensions
        for cell in sheet[1]:cell.font=Font(bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='1B5961')
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                if isinstance(cell.value,str):cell.data_type='s'
                cell.alignment=Alignment(vertical='top',wrap_text=True)
        for index in range(1,sheet.max_column+1):sheet.column_dimensions[get_column_letter(index)].width=22
    instructions=workbook.create_sheet('说明');instructions.append(['口径','说明'])
    instructions.append(['职级认定','申报职级与评分不自动代表晋级；以管理员确认的认证结论、生效职级和生效日期为准。'])
    instructions.append(['次数','认证次数按关联材料及历史补录记录统计；待确认、未通过均保留，晋级次数仅计入已生效的晋级记录。'])
    instructions.append(['身份','同名不自动合并；待关联材料尚未纳入具体人员的职级与次数统计。'])
    instructions.append(['评分','待确认记录显示当前评审数据，已确认记录保留确认时的评分快照。'])
    instructions.column_dimensions['A'].width=20;instructions.column_dimensions['B'].width=100
    output=io.BytesIO();workbook.save(output);workbook.close()
    return output.getvalue(),f'人员职级档案-{now_text()[:10]}.xlsx'


def personnel_http_get(handler,path,query):
    if not path.startswith('/api/admin/personnel'):return False
    if not handler.require_admin():return True
    from urllib.parse import parse_qs
    params=parse_qs(query)
    try:
        match=re.fullmatch(r'/api/admin/personnel/([a-f0-9]+)(/export)?',path)
        if path=='/api/admin/personnel':
            data=personnel_data()
            for person in data['items']:person.pop('records',None)
            handler.json_response(data)
        elif path=='/api/admin/personnel/export' or (match and match[2]):
            content,filename=export_personnel(match[1] if match else None,params.get('q',[''])[0],params.get('level',[''])[0])
            handler.send_download(content,'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',filename)
        elif match:handler.json_response(personnel_data(match[1]))
        else:handler.json_response({'error':'接口不存在'},404)
    except ValueError as exc:handler.json_response({'error':str(exc)},400)
    return True


def personnel_http_post(handler,path,payload):
    if not path.startswith('/api/admin/personnel'):return False
    actor=handler.require_admin()
    if not actor:return True
    match=re.fullmatch(r'/api/admin/personnel/([a-f0-9]+)(?:/(save|link|records)(?:/([a-f0-9]+))?)?',path)
    if path=='/api/admin/personnel':handler.json_response(save_personnel(payload,actor))
    elif match and match[2]=='save':handler.json_response(save_personnel(payload,actor,match[1]))
    elif match and match[2]=='link':handler.json_response(link_personnel_submission(match[1],clean(payload.get('submissionId')),actor))
    elif match and match[2]=='records':handler.json_response(save_personnel_certification(match[1],payload,actor,match[3]))
    else:handler.json_response({'error':'接口不存在'},404)
    return True


class CertificationHTTPServer(ThreadingHTTPServer):
    def handle_error(self, request, client_address):
        APP_LOG.exception("Unhandled connection error client=%s", client_address[0] if client_address else "unknown")


if __name__ == "__main__":
    enable_exception_tracing()
    if "--parse-one" in sys.argv:
        try:
            argument_index = sys.argv.index("--parse-one")
            parse_one_workbook_to_file(Path(sys.argv[argument_index + 1]), sys.argv[argument_index + 2], Path(sys.argv[argument_index + 3]))
        except Exception as exc:
            print(f"{type(exc).__name__}: {exc}", file=sys.stderr)
            raise SystemExit(1)
        raise SystemExit(0)
    if "--backup" in sys.argv:
        init_storage()
        print(create_runtime_backup())
        raise SystemExit(0)
    if "--restore" in sys.argv:
        if "--confirm" not in sys.argv:
            raise SystemExit("恢复操作需要 --confirm，并且必须先停止服务")
        try:
            backup_index = sys.argv.index("--restore") + 1
            restore_runtime_backup(Path(sys.argv[backup_index]))
        except (IndexError, ValueError) as exc:
            raise SystemExit(str(exc)) from exc
        print("恢复完成")
        raise SystemExit(0)
    init_storage()
    port = int(os.environ.get("CERT_PORT", "8787"))
    host = os.environ.get("CERT_HOST", "127.0.0.1")
    APP_LOG.info("service_start version=0.8.0 host=%s port=%s root=%s", host, port, ROOT)
    print(f"任职资格认证系统已启动：http://{host}:{port}")
    try:
        CertificationHTTPServer((host, port), Handler).serve_forever()
    except Exception:
        APP_LOG.exception("service_stopped_unexpectedly")
        raise
